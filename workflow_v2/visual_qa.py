"""比對 PDF 文字位置、預覽圖片與 PPTX 圖形，修復版面問題並記錄警示。"""
import re
import time
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu

from .template_styles import contrast, INK

EMU_PT = 12700
SCALE = 2.0
MAX_WORDS = 90            # contrast samples per page
LOW_CONTRAST = 3.0
FIXABLE = ('text_overflow', 'low_contrast', 'off_slide')
MESSAGES = {
    'blank_page': '這頁幾乎是空白的',
    'off_slide': '有文字超出投影片邊緣',
    'text_overflow': '有文字超出所在的方框',
    'text_overlap': '有文字互相重疊',
    'low_contrast': '有文字顏色和背景太接近，不容易看清楚',
}


def _issue(page, code, detail='', shape_id=None, fixed=False):
    return {'page': page, 'code': code, 'message': MESSAGES[code] + (('：' + detail) if detail else ''),
            'shape_id': shape_id, 'fixable': code in FIXABLE}


def _hex(rgb):
    return '%02X%02X%02X' % tuple(rgb)


def text_shapes(slide):
    """Top-level shapes that carry text: (shape, (left, top, right, bottom) in pt, character count)."""
    found = []
    for shape in slide.shapes:
        if not getattr(shape, 'has_text_frame', False) or shape.left is None or shape.rotation:
            continue
        count = len(re.sub(r'\s', '', shape.text_frame.text))
        if count:
            left, top = shape.left / EMU_PT, shape.top / EMU_PT
            found.append((shape, (left, top, left + shape.width / EMU_PT, top + shape.height / EMU_PT), count))
    return found


def shape_at(shapes, cx, cy, slack=3):
    best = None
    for shape, (l, t, r, b), _ in shapes:
        if l - slack <= cx <= r + slack and t - slack <= cy <= b + slack:
            area = (r - l) * (b - t)
            if best is None or area < best[0]:
                best = (area, shape)
    return best[1] if best else None


def sample_colors(image, box):
    """(background, text) colours of a word box from the page raster, or None when too small."""
    x0, y0, x1, y1 = [int(v * SCALE) for v in box]
    x0, y0 = max(0, x0), max(0, y0)
    x1, y1 = min(image.width, x1), min(image.height, y1)
    if x1 - x0 < 4 or y1 - y0 < 4:
        return None
    colors = image.crop((x0, y0, x1, y1)).getcolors(maxcolors=1 << 20) or []
    if not colors:
        return None
    total = sum(n for n, _ in colors)
    bg = max(colors)[1]
    ranked = sorted(colors, key=lambda c: -abs(_lum(c[1]) - _lum(bg)))
    need = max(3, int(total * .06))
    picked = []
    count = 0
    for n, rgb in ranked:
        picked.append((n, rgb))
        count += n
        if count >= need:
            break
    fg = tuple(round(sum(n * c[i] for n, c in picked) / count) for i in range(3))
    return _hex(bg), _hex(fg)


def _lum(rgb):
    return 0.2126 * rgb[0] + 0.7152 * rgb[1] + 0.0722 * rgb[2]


def analyze(pptx_path, pdf_path):
    """Return {'pages': n, 'issues': [...]} for a rendered deck and its PDF."""
    import pymupdf
    from PIL import Image
    deck = Presentation(str(pptx_path))
    issues = []
    with pymupdf.open(str(pdf_path)) as pdf:
        pages = min(len(pdf), len(deck.slides))
        for index in range(pages):
            number = index + 1
            page = pdf[index]
            slide = deck.slides[index]
            width_pt, height_pt = page.rect.width, page.rect.height
            words = [(w[0], w[1], w[2], w[3], w[4], (w[5], w[6])) for w in page.get_text('words')]
            shapes = text_shapes(slide)
            seen = set()

            def add(code, detail='', shape=None):
                key = (code, shape.shape_id if shape is not None else detail)
                if key not in seen:
                    seen.add(key)
                    issues.append(_issue(number, code, detail, shape.shape_id if shape is not None else None))

            # 1. text beyond the slide edge
            for x0, y0, x1, y1, text, _ in words:
                if x0 < -1 or y0 < -1 or x1 > width_pt + 1 or y1 > height_pt + 1:
                    add('off_slide', text[:12], shape_at(shapes, (x0 + x1) / 2, (y0 + y1) / 2, 40))
            # 2. text that no longer fits the box it was written for
            for shape, (l, t, r, b), count in shapes:
                if count < 4:
                    continue
                inside = sum(len(w[4]) for w in words
                             if l - 4 <= (w[0] + w[2]) / 2 <= r + 4 and w[1] >= t - 6 and w[3] <= b + 6)
                if inside < count * .88:
                    add('text_overflow', shape.text_frame.text.strip()[:14], shape)
            # 3. words drawn on top of each other
            ordered = sorted(words, key=lambda w: w[0])
            for i, a in enumerate(ordered):
                for b in ordered[i + 1:]:
                    if b[0] >= a[2]:
                        break
                    if a[5] == b[5]:
                        continue
                    w = min(a[2], b[2]) - max(a[0], b[0])
                    h = min(a[3], b[3]) - max(a[1], b[1])
                    smaller = min((a[2] - a[0]) * (a[3] - a[1]), (b[2] - b[0]) * (b[3] - b[1]))
                    if w > 2 and h > 2 and smaller > 0 and w * h > .3 * smaller:
                        add('text_overlap', a[4][:8] + '／' + b[4][:8])
            # 4. colours: sampled from the raster
            pix = page.get_pixmap(matrix=pymupdf.Matrix(SCALE, SCALE), alpha=False)
            image = Image.frombytes('RGB', (pix.width, pix.height), pix.samples)
            step = max(1, len(words) // MAX_WORDS)
            worst = None
            sample = [w for w in words if len(re.sub(r'[\W_]', '', w[4])) >= 2]   # punctuation alone has no ink to measure
            for x0, y0, x1, y1, text, _ in sample[::step]:
                sampled = sample_colors(image, (x0, y0, x1, y1))
                if sampled is None:
                    continue
                ratio = contrast(*sampled)
                if ratio < LOW_CONTRAST and (worst is None or ratio < worst[0]):
                    worst = (ratio, text, sampled, (x0 + x1) / 2, (y0 + y1) / 2)
            if worst:
                shape = shape_at(shapes, worst[3], worst[4])
                add('low_contrast', f'「{worst[1][:8]}」對比 {worst[0]:.1f}', shape)
                if shape is not None:
                    issues[-1]['background'] = worst[2][0]
            # 5. nothing to look at
            content = image.resize((image.width // 8, image.height // 8))
            colors = content.getcolors(maxcolors=1 << 20) or [(1, (0, 0, 0))]
            total = sum(n for n, _ in colors)
            dominant = max(colors)[0]
            body_words = [w for w in words if .5 * 72 < w[1] < 6.9 * 72]
            if not body_words and (total - dominant) / total < .02:
                add('blank_page')
    return {'pages': pages, 'issues': issues}


def _runs(shape):
    for paragraph in shape.text_frame.paragraphs:
        yield from paragraph.runs


def apply_fixes(pptx_path, issues):
    """Repair what is safely repairable (smaller type, readable colour, back inside the slide). Returns changes."""
    deck = Presentation(str(pptx_path))
    changed = 0
    for issue in issues:
        if not issue.get('fixable') or issue.get('shape_id') is None:
            continue
        slide = deck.slides[issue['page'] - 1]
        shape = next((s for s in slide.shapes if s.shape_id == issue['shape_id']), None)
        if shape is None:
            continue
        if issue['code'] == 'text_overflow':
            for run in _runs(shape):
                size = run.font.size.pt if run.font.size else None
                if size and size > 14:
                    from pptx.util import Pt
                    run.font.size = Pt(max(14, round(size * .9)))
                    changed += 1
        elif issue['code'] == 'low_contrast':
            background = issue.get('background') or 'FFFFFF'
            best = max((INK, 'FFFFFF'), key=lambda c: contrast(c, background))
            for run in _runs(shape):
                try:
                    current = str(run.font.color.rgb) if run.font.color and run.font.color.type else None
                except AttributeError:
                    current = None
                if current is None or contrast(current, background) < 4.5:
                    run.font.color.rgb = RGBColor.from_string(best)
                    changed += 1
        elif issue['code'] == 'off_slide':
            sw, sh = deck.slide_width, deck.slide_height
            margin = Emu(int(.2 * 914400))
            shape.width = min(shape.width, sw - 2 * margin)
            shape.height = min(shape.height, sh - 2 * margin)
            shape.left = max(margin, min(shape.left, sw - margin - shape.width))
            shape.top = max(Emu(0), min(shape.top, sh - margin - shape.height))
            changed += 1
    if changed:
        partial = Path(pptx_path).with_suffix('.qa.pptx')
        deck.save(str(partial))
        partial.replace(pptx_path)
    return changed


def run(out, reconvert, max_rounds=2, budget_seconds=150, log=None):
    """Check, repair and re-check the job's deck. `reconvert()` must rebuild 教材.pdf from 教材.pptx."""
    out = Path(out)
    pptx, pdf = out / '教材.pptx', out / '教材.pdf'
    started = time.time()
    first = analyze(pptx, pdf)
    current = first
    rounds = 0
    while rounds < max_rounds and time.time() - started < budget_seconds:
        fixable = [i for i in current['issues'] if i['fixable']]
        if not fixable:
            break
        backup = out / 'visual-qa' / 'before.pptx'
        if not backup.exists():
            backup.parent.mkdir(exist_ok=True)
            backup.write_bytes(pptx.read_bytes())
        if not apply_fixes(pptx, fixable):
            break
        reconvert()
        current = analyze(pptx, pdf)
        rounds += 1
    keys = lambda found: {(i['page'], i['code']) for i in found}
    report = {'schema_version': 1, 'checked_pages': first['pages'], 'rounds': rounds,
              'found': len(first['issues']), 'fixed': len(keys(first['issues']) - keys(current['issues'])),
              'remaining': current['issues'], 'remaining_pages': sorted({i['page'] for i in current['issues']}),
              'seconds': round(time.time() - started, 1)}
    from .store import atomic
    atomic(out / 'visual-qa.json', report)
    if log:
        log(report)
    return report
