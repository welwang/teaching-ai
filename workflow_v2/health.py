"""Independent bounded probes; one failed component never hides other results."""
import json,os,shutil,sqlite3,subprocess,tempfile
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor,as_completed
import requests
from filelock import FileLock
from .store import ROOT,DATA,atomic
from .model import endpoint
from .local_images import probe as _image_probe, enabled as image_enabled

def image_probe():
    if not image_enabled():
        return {'ready':False,'reason':'已停用（預設）：圖片改用授權素材搜尋與可編輯圖解，不影響簡報'}
    return _image_probe()

def browser_probe():
    os.environ['PLAYWRIGHT_BROWSERS_PATH']=str(ROOT/'runtime/playwright-v2')
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        try:browser=p.chromium.launch(headless=True,timeout=15000)
        except Exception:browser=p.chromium.launch(channel='msedge',headless=True,timeout=15000)
        try:
            page=browser.new_page();page.set_content('<h1>Local renderer ready</h1>')
            data=page.screenshot();return {'screenshot_bytes':len(data)}
        finally:browser.close()

def database_probe():
    DATA.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(dir=DATA) as tmp:
        c=sqlite3.connect(Path(tmp)/'probe.sqlite')
        try:
            c.execute('create table probe(value integer)');c.execute('insert into probe values(1)');c.commit()
            assert c.execute('select value from probe').fetchone()[0]==1
        finally:c.close()
    return 'SQLite read/write OK'

def office_probe():
    from pptx import Presentation
    with tempfile.TemporaryDirectory(dir=DATA) as tmp:
        folder=Path(tmp);p=Presentation();s=p.slides.add_slide(p.slide_layouts[0]);s.shapes.title.text='Local export test';p.save(folder/'probe.pptx')
        from .office import convert_pdf
        r=convert_pdf(folder/'probe.pptx',folder,timeout=60)
        if r.returncode or not (folder/'probe.pdf').exists():raise RuntimeError('PDF 轉檔測試失敗')
        return 'PPTX → PDF OK'

def models_probe():
    r=requests.get(endpoint()+'/api/tags',timeout=8);r.raise_for_status()
    names=[x['name'] for x in r.json().get('models',[])]
    from .model_policy import configured_model
    from .strata import is_strata
    model=configured_model()
    if not is_strata(model) and model not in names:
        raise RuntimeError('模型服務已啟動，但尚未安裝系統指定模型：'+model)
    return names

def strata_probe():
    from . import strata
    if not strata.installed():
        raise RuntimeError('Strata 尚未安裝完成；下載與安裝進度：logs/strata-install.log')
    return strata.probe()

def gpu_probe():
    r=subprocess.run(['nvidia-smi','--query-gpu=name,memory.total,memory.free,driver_version','--format=csv,noheader'],capture_output=True,text=True,timeout=10,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
    if r.returncode:raise RuntimeError(r.stderr)
    return r.stdout.strip()

def run(deep=False):
    DATA.mkdir(parents=True,exist_ok=True)
    probes={'database':database_probe,'models':models_probe,'gpu':gpu_probe,
        'disk':lambda: {'free_gb':round(shutil.disk_usage(ROOT).free/2**30,1)},'browser':browser_probe,'office':office_probe,
        'local_image':image_probe}
    from .model_policy import configured_model
    from .strata import is_strata
    if is_strata(configured_model()):probes['strata']=strata_probe
    report={}
    with ThreadPoolExecutor(max_workers=4) as pool:
        pending={pool.submit(fn):name for name,fn in probes.items()}
        for future in as_completed(pending):
            name=pending[future]
            try:
                detail=future.result()
                report[name]={'ready':detail['ready'] if name=='local_image' else True,'detail':detail}
            except Exception as exc:report[name]={'ready':False,'detail':str(exc)}
            if name in ('local_image','browser'):
                report[name].update(optional=True,used_by_current_pipeline=False)
    atomic(DATA/'health.json',report);return report

if __name__=='__main__':print(json.dumps(run(),ensure_ascii=False,indent=2))
