"""在指定區域內，以 python-pptx 原生圖形繪製教學圖表。"""
import math
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from .textfit import fit_size,font_size,Overflow

FONT='Microsoft JhengHei'
SLIDE_W=13.333;SLIDE_H=7.5
ICON_SHAPES={'circle':MSO_SHAPE.OVAL,'square':MSO_SHAPE.ROUNDED_RECTANGLE,'star':MSO_SHAPE.STAR_5_POINT,
             'heart':MSO_SHAPE.HEART,'triangle':MSO_SHAPE.ISOSCELES_TRIANGLE,'diamond':MSO_SHAPE.DIAMOND,
             'cup':MSO_SHAPE.FLOWCHART_MANUAL_OPERATION,'drop':MSO_SHAPE.TEAR,'block':MSO_SHAPE.RECTANGLE}

def blend(first,second,weight):
    return ''.join(f'{round(int(first[i:i+2],16)*(1-weight)+int(second[i:i+2],16)*weight):02X}' for i in (0,2,4))

def palette(style):
    p,a=style['primary'],style['accent']
    return [p,a,blend(p,a,.5) if blend(p,a,.5) not in (p,a) else '6C757D','5B6C8F']

def _clamp(box):
    x,y,w,h=box
    x=max(0.0,min(x,SLIDE_W-.1));y=max(0.0,min(y,SLIDE_H-.1))
    return x,y,max(.1,min(w,SLIDE_W-x)),max(.1,min(h,SLIDE_H-y))

def rect(slide,kind,x,y,w,h,fill=None,line=None,line_w=1.5,dash=False,radius=None):
    x,y,w,h=_clamp((x,y,w,h))
    shape=slide.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h))
    if radius is not None and kind==MSO_SHAPE.ROUNDED_RECTANGLE:shape.adjustments[0]=radius
    if fill:
        shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(fill)
    else:shape.fill.background()
    if line:
        shape.line.color.rgb=RGBColor.from_string(line);shape.line.width=Pt(line_w)
        if dash:
            ln=shape.line._get_or_add_ln()
            for old in ln.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}prstDash'):ln.remove(old)
            d=OxmlElement('a:prstDash');d.set('val','dash');ln.append(d)
    else:shape.line.fill.background()
    shape.shadow.inherit=False
    return shape

def write(shape,text,size,color='1A202C',bold=False,align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE,floor=14,pad=.06):
    """Write text into a shape at the largest size (<= size) that fits; returns that size."""
    w=shape.width/914400;h=shape.height/914400
    used=fit_size(text,w,h,size,floor,pad)
    tf=shape.text_frame;tf.clear();tf.word_wrap=True;tf.vertical_anchor=anchor
    tf.margin_left=tf.margin_right=Inches(pad);tf.margin_top=tf.margin_bottom=Inches(.03)
    for i,line in enumerate(str(text).split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align
        run=p.add_run();run.text=line
        run.font.name=FONT;run.font.size=Pt(used);run.font.bold=bold
        run.font.color.rgb=RGBColor.from_string(color)
    return used

def label(slide,x,y,w,h,text,size,color='1A202C',bold=False,align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE,floor=14):
    x,y,w,h=_clamp((x,y,w,h))
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    return write(box,text,size,color,bold,align,anchor,floor),box

def _title(slide,vm,box,style,sizes):
    """Optional caption strip; returns the y offset used."""
    x,y,w,h=box;title=vm.get('title') or ''
    if not title:return 0.0
    used,_=label(slide,x,y,w,.5,title,24,style['primary'],True,PP_ALIGN.LEFT,floor=18)
    sizes.append(used);return .55

def draw_icon_groups(slide,vm,box,style):
    x,y,w,h=box;sizes=[];colors=palette(style)
    top=_title(slide,vm,box,style,sizes);groups=vm['groups']
    row_h=min(1.35,(h-top)/len(groups));lw=min(2.1,w*.22);cw=1.25
    icons_w=max(1.0,w-lw-cw-.2)
    for gi,group in enumerate(groups):
        ry=y+top+gi*row_h;color=colors[gi%len(colors)]
        rect(slide,MSO_SHAPE.ROUNDED_RECTANGLE,x,ry+.04,w,row_h-.08,fill=blend(style['background'],'FFFFFF',.7),
             line=blend(color,'FFFFFF',.6),line_w=1,radius=.12)
        used,_=label(slide,x+.08,ry+.06,lw-.1,row_h-.12,group['label'],28,color,True,PP_ALIGN.CENTER,floor=18)
        sizes.append(used)
        count=group['count'];per=1 if count<=12 else math.ceil(count/12);shown=math.ceil(count/per)
        size=min(row_h*.62,icons_w/shown*.82,.8)
        gap=(icons_w-shown*size)/(shown+1) if shown>1 else (icons_w-size)/2
        gap=min(gap,size*.35) if shown>1 else gap
        total=shown*size+(shown-1)*gap if shown>1 else size
        start=x+lw+(icons_w-total)/2
        for k in range(shown):
            shape=rect(slide,ICON_SHAPES[group['icon']],start+k*(size+gap),ry+(row_h-size)/2-(.08 if per>1 else 0),
                       size,size,fill=color,line='FFFFFF',line_w=1.5,
                       radius=.25 if group['icon']=='square' else None)
        if per>1:
            used,_=label(slide,x+lw,ry+row_h-.34,icons_w,.3,f'1 個圖形代表 {per} 個',16,'666666',False,floor=12)
        used,_=label(slide,x+w-cw-.05,ry+.06,cw,row_h-.12,str(count),36,color,True,floor=28)
        sizes.append(used)
    return sizes

def _num(value):
    return str(int(value)) if float(value)==int(value) else f'{value:g}'

def draw_bar_model(slide,vm,box,style):
    x,y,w,h=box;sizes=[];colors=palette(style)
    top=_title(slide,vm,box,style,sizes);bars=vm['bars']
    row_h=min(1.6,(h-top)/len(bars));lw=2.0;rw=2.3
    bar_w=max(2.0,w-lw-rw-.1);unit_w=bar_w/max(b['parts'] for b in bars)
    for bi,bar in enumerate(bars):
        ry=y+top+bi*row_h;color=colors[bi%len(colors)];bh=min(.95,row_h*.62)
        by=ry+(row_h-bh)/2
        text=bar['label']+'\n'+f"{bar['parts']} 份"
        used,_=label(slide,x,ry+.03,lw,row_h-.06,text,26,color,True,floor=16)
        sizes.append(used)
        for k in range(bar['parts']):
            unknown=bar.get('unknown')
            seg=rect(slide,MSO_SHAPE.RECTANGLE,x+lw+k*unit_w,by,unit_w,bh,
                     fill=blend(color,'FFFFFF',.72) if unknown else color,
                     line=color if unknown else 'FFFFFF',line_w=2.25 if unknown else 2.5,dash=bool(unknown))
            if 'unit_value' in bar and bar['parts']<=8 and not unknown:
                sizes.append(write(seg,_num(bar['unit_value']),28,'FFFFFF',True,floor=14))
            elif unknown and bar['parts']<=8:
                sizes.append(write(seg,'？',28,color,True,floor=14))
        unit=bar.get('unit') or ''
        if 'total' in bar:tot=f"= {_num(bar['total'])}{unit}"
        elif bar.get('unknown'):tot=f"= ？{unit}"
        elif 'unit_value' in bar:tot=f"每份 {_num(bar['unit_value'])}{unit}"
        else:tot=''
        if tot:
            used,_=label(slide,x+lw+bar['parts']*unit_w+.08,by,w-lw-bar['parts']*unit_w-.1,bh,tot,32,
                         style['accent'] if bar.get('unknown') else color,True,PP_ALIGN.LEFT,floor=20)
            sizes.append(used)
    return sizes

def draw_ratio_arrows(slide,vm,box,style):
    x,y,w,h=box;sizes=[];colors=palette(style)
    top=_title(slide,vm,box,style,sizes)
    labels=vm.get('labels') or ['','']
    units=vm.get('units') or ['','']
    labels=(labels+['',''])[:2];units=(units+['',''])[:2]
    header_h=.55 if any(labels) else 0
    avail=h-top-header_h
    row_h=max(.7,min(1.15,avail*.3));gap_h=max(.6,avail-2*row_h)
    bw=min(2.3,w*.3);cx=[x+w*.3,x+w*.7]
    for i in (0,1):
        if labels[i]:
            text=labels[i]+(f'（{units[i]}）' if units[i] else '')
            used,_=label(slide,cx[i]-bw/2-.3,y+top,bw+.6,header_h,text,26,colors[i],True,floor=16)
            sizes.append(used)
    y1=y+top+header_h;y2=y1+row_h+gap_h
    for row,(pair,ry) in enumerate(((vm['from'],y1),(vm['to'],y2))):
        for i in (0,1):
            box_shape=rect(slide,MSO_SHAPE.ROUNDED_RECTANGLE,cx[i]-bw/2,ry,bw,row_h,
                           fill=colors[i] if row==0 else blend(colors[i],'FFFFFF',.78),
                           line=colors[i],line_w=2.25,radius=.2)
            sizes.append(write(box_shape,_num(pair[i]),44,'FFFFFF' if row==0 else colors[i],True,floor=28))
        used,_=label(slide,x+w*.5-.4,ry,.8,row_h,'：',44,style['ink'],True,floor=28)
    for i in (0,1):
        aw=.7;ah=gap_h-.16
        rect(slide,MSO_SHAPE.DOWN_ARROW,cx[i]-aw/2,y1+row_h+.08,aw,ah,fill=style['accent'])
        used,_=label(slide,cx[i]+aw/2+.05,y1+row_h+.08,1.6,ah,vm['op'],40,style['accent'],True,PP_ALIGN.LEFT,floor=28)
        sizes.append(used)
    return sizes

def draw_compare_ab(slide,vm,box,style):
    x,y,w,h=box;sizes=[];colors=palette(style)
    top=_title(slide,vm,box,style,sizes)
    cw=(w-.9)/2;ch=h-top
    for i,(side,tag) in enumerate(((vm['left'],'A'),(vm['right'],'B'))):
        cx=x+i*(cw+.9);color=colors[i]
        rect(slide,MSO_SHAPE.ROUNDED_RECTANGLE,cx,y+top,cw,ch,fill='FFFFFF',line=color,line_w=3,radius=.06)
        badge=rect(slide,MSO_SHAPE.OVAL,cx+.15,y+top+.15,.8,.8,fill=color)
        sizes.append(write(badge,tag,32,'FFFFFF',True,floor=24))
        used,_=label(slide,cx+1.05,y+top+.12,cw-1.2,.86,side['title'],30,color,True,PP_ALIGN.LEFT,floor=20)
        sizes.append(used)
        items=side.get('items') or []
        value_h=1.0 if side.get('value') else 0
        body_h=ch-1.15-value_h-.1
        if items:
            used,_=label(slide,cx+.2,y+top+1.05,cw-.4,max(.4,body_h),'\n'.join(items),28,style['ink'],False,
                         PP_ALIGN.LEFT,MSO_ANCHOR.TOP,floor=18)
            sizes.append(used)
        if side.get('value'):
            vb=rect(slide,MSO_SHAPE.ROUNDED_RECTANGLE,cx+.2,y+top+ch-value_h-.15,cw-.4,value_h,
                    fill=blend(color,'FFFFFF',.82),line=color,line_w=1.5,radius=.2)
            sizes.append(write(vb,side['value'],40,color,True,floor=28))
    vs=rect(slide,MSO_SHAPE.OVAL,x+cw+.08,y+top+ch/2-.37,.74,.74,fill=style['accent'])
    sizes.append(write(vs,'VS',20,'FFFFFF',True,floor=14,pad=.02))
    return sizes

def draw_concept_map(slide,vm,box,style):
    x,y,w,h=box;sizes=[];colors=palette(style)
    top=_title(slide,vm,box,style,sizes)
    children=vm['children'];k=len(children);avail=h-top
    root_w=min(2.7,w*.24);root_h=min(1.6,avail*.42)
    root=rect(slide,MSO_SHAPE.OVAL,x,y+top+(avail-root_h)/2,root_w,root_h,fill=style['primary'])
    sizes.append(write(root,vm['root'],34,'FFFFFF',True,floor=22))
    gap=.1;row_h=(avail-gap*(k-1))/k
    col_x=x+root_w+1.0;col_w=w-root_w-1.0
    label_w=min(2.2,col_w*.34)
    rx=x+root_w;ry=y+top+avail/2
    for i,child in enumerate(children):
        cy=y+top+i*(row_h+gap);color=colors[i%len(colors)]
        conn=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(rx),Inches(ry),Inches(col_x),Inches(cy+row_h/2))
        conn.line.color.rgb=RGBColor.from_string(style['primary']);conn.line.width=Pt(3)
        pill=rect(slide,MSO_SHAPE.ROUNDED_RECTANGLE,col_x,cy,label_w,row_h,fill=color,radius=.25)
        sizes.append(write(pill,child['label'],28,'FFFFFF',True,floor=18))
        detail=child.get('example','')
        if child.get('note'):detail=(detail+'\n' if detail else '')+child['note']
        card=rect(slide,MSO_SHAPE.ROUNDED_RECTANGLE,col_x+label_w+.08,cy,col_w-label_w-.08,row_h,
                  fill='FFFFFF',line=blend(color,'FFFFFF',.55),line_w=1.5,radius=.15)
        if detail:sizes.append(write(card,detail,26,style['ink'],False,PP_ALIGN.LEFT,floor=16,pad=.12))
    return sizes

def draw_number_line(slide,vm,box,style):
    """Native editable axis; point positions are proportional to coordinates."""
    x,y,w,h=box;sizes=[]
    top=_title(slide,vm,box,style,sizes)
    left,right=x+.4,x+w-.4;axis_y=y+top+(h-top)*.58
    low,high=vm['minimum'],vm['maximum']
    axis=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(left),Inches(axis_y),Inches(right),Inches(axis_y))
    axis.line.color.rgb=RGBColor.from_string(style['ink']);axis.line.width=Pt(2)
    for value in (low,0,high):
        if not low<=value<=high:continue
        px=left+(value-low)/(high-low)*(right-left)
        tick=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(px),Inches(axis_y-.09),Inches(px),Inches(axis_y+.09))
        tick.line.color.rgb=RGBColor.from_string(style['ink'])
        label(slide,px-.4,axis_y+.2,.8,.38,f'{value:g}',18,style['ink'])
    for index,point in enumerate(vm['points']):
        px=left+(point['value']-low)/(high-low)*(right-left)
        rect(slide,MSO_SHAPE.OVAL,px-.085,axis_y-.085,.17,.17,fill=style['primary'])
        py=axis_y-.85-(index%2)*.65
        sizes.append(label(slide,max(x,min(px-.55,x+w-1.1)),max(y+top,py),1.1,.65,
            point['label']+'\n'+f"{point['value']:g}",24,style['primary'],True))
    return [size for size in sizes if isinstance(size,(int,float))]

def draw_venn_diagram(slide,vm,box,style):
    """Two editable overlapping circles, with three distinct semantic regions."""
    from pptx.oxml.xmlchemy import OxmlElement
    x,y,w,h=box;sizes=[];top=_title(slide,vm,box,style,sizes);colors=palette(style)
    diameter=min(w/1.55,h-top-.65)
    if diameter<1.5:raise Overflow('維恩圖需要足夠的圖解空間')
    offset=diameter*.55;left=x+(w-diameter-offset)/2;cy=y+top+.45
    for i,side in enumerate((vm['left'],vm['right'])):
        circle=rect(slide,MSO_SHAPE.OVAL,left+i*offset,cy,diameter,diameter,
                    fill=colors[i],line=colors[i],line_w=2)
        alpha=OxmlElement('a:alpha');alpha.set('val','18000')
        color_node=next(circle._element.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr'))
        color_node.append(alpha)
        sizes.append(label(slide,left+i*offset,cy-.55,diameter,.45,side['title'],24,colors[i],True)[0])
    regions=[(left+.08,vm['left']['items']),
             (left+offset+.08,vm.get('intersection',[])),
             (left+diameter+.04,vm['right']['items'])]
    for index,(rx,items) in enumerate(regions):
        width=offset-.10 if index!=1 else diameter-offset+.16
        if index==1:rx-=.10
        # The right non-overlap region starts at the left circle's outer edge.
        text='\n'.join(items)
        if text:
            used=font_size(text,width,diameter*.46,22,20)
            if index==1 and len(text)<=4:used=20
            sizes.append(label(slide,rx,cy+diameter*.30,width,diameter*.46,text,used,style['ink'],
                False,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE,floor=20)[0])
    sizes.append(label(slide,left+offset,cy+.08,diameter-offset,.36,'共同點',16,style['primary'],True,floor=14)[0])
    return sizes

def draw_spatial_scene(slide,vm,box,style):
    """Fixed geometry for English position words; the model supplies no coordinates."""
    x,y,w,h=box;sizes=[];top=_title(slide,vm,box,style,sizes)
    relation=vm['relation'];cx=x+w/2;cy=y+top+(h-top)*.44
    bw=min(2,w*.48);bh=min(1.4,(h-top)*.30);ball=min(.70,bh*.60)
    bx,by=cx-bw/2,cy-bh/2;px,py=cx-ball/2,cy-ball/2
    if relation=='on':py=by-ball
    elif relation=='under':py=by+bh+.12
    elif relation=='next_to':px=bx-ball-.18
    elif relation=='between':bx=x+.15;px=cx-ball/2;bw=min(1.2,w*.24)
    elif relation=='behind':py=by-.30;px=bx+bw*.58
    elif relation=='in_front_of':py=by+bh-ball*.55
    def object_shape():return rect(slide,MSO_SHAPE.OVAL,px,py,ball,ball,fill=style['accent'],line=style['accent'])
    if relation=='behind':object_shape()
    rect(slide,MSO_SHAPE.RECTANGLE,bx,by,bw,bh,fill=blend(style['primary'],'FFFFFF',.80),line=style['primary'],line_w=2)
    if relation=='between':rect(slide,MSO_SHAPE.RECTANGLE,x+w-bw-.15,by,bw,bh,
        fill=blend(style['primary'],'FFFFFF',.80),line=style['primary'],line_w=2)
    if relation!='behind':object_shape()
    for index,ref in enumerate(vm['references']):
        rx=bx if index==0 else x+w-bw-.15
        ry=by-.56 if relation=='under' else by+bh+.08
        sizes.append(label(slide,rx,ry,bw,.5,ref,22,style['primary'],True,floor=20)[0])
    # Labels and the relation stay below the drawing rather than covering it.
    caption=vm['object_label']+' — '+relation.replace('_',' ')
    sizes.append(label(slide,x,y+h-.64,w,.54,caption,24,style['ink'],True,floor=20)[0])
    return sizes

DRAW={'venn_diagram':draw_venn_diagram,'spatial_scene':draw_spatial_scene,'number_line':draw_number_line,'icon_groups':draw_icon_groups,'bar_model':draw_bar_model,'ratio_arrows':draw_ratio_arrows,
      'compare_ab':draw_compare_ab,'concept_map':draw_concept_map}

def draw_model(slide,vm,box,style):
    """Draw a sanitised visual model inside box=(x,y,w,h) inches; returns font sizes used."""
    return DRAW[vm['type']](slide,vm,_clamp(box),style)
