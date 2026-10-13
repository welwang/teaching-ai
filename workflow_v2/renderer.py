"""Deterministic editable PPTX layouts; never executes model code or truncates text."""
import copy,json,math,re,unicodedata
from copy import deepcopy
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from .contracts import LAYOUTS,LEGACY_LAYOUTS,validate,DECK
from .textfit import Overflow,estimated_lines,font_size
from .visuals import sanitize_slide,model_caption
from .slide_layouts import (apply_role,uses_special,compose_question,compose_steps,compose_diagram,X0)
from .store import ROOT
from .template_styles import theme as get_theme, style_for
from . import beautify
from .template_design import layout_boxes,card_boxes
from .text_safety import xml_text,clean_slide_text,format_teaching_text

ASSETS=ROOT/'assets/workflow-v2'
MEDIA=ASSETS/'media'
FONT='Microsoft JhengHei'
BG='F7F2E8'
INK='1A202C'
TEAL='008080'
PALE='E7F4F1'
ORANGE='DE8B35'
MARGIN=.93
LEGACY_ROUTE=dict(zip(LEGACY_LAYOUTS,('Cover','Two_Column','Text_Visual','Steps','Question','Summary')))

def visual_boxes(style):
    """Text/visual proportions differ by the selected composition family."""
    return {k:v for k,v in layout_boxes('Text_Visual',style).items() if k in (1,2)}


def make_template(template_id=None,style=None):
    """In-memory blank canvas for the automatic style (a legacy template id is ignored). Never a file on disk."""
    from .template_builder import build_master
    return build_master(style)


def deck_style(deck,grade=None):
    """Palette and type scale from the deck's grade and subject, never from a template."""
    meta=deck.get('deck_meta') or {}
    grade=grade or meta.get('grade')
    if not grade:
        audience=str(meta.get('target_audience') or '')
        grade=2 if '低年級' in audience else 4 if '中年級' in audience else 6 if '高年級' in audience else 8 if '國中' in audience else None
    return style_for(grade,meta.get('subject'))


def fill(shape,text,max_size=30,min_size=24,color=INK,bold=False,align=PP_ALIGN.LEFT):
    text=format_teaching_text(text)
    size=font_size(text,shape.width/914400,shape.height/914400,max_size,min_size)
    tf=shape.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.07)
    tf.margin_top=tf.margin_bottom=Inches(.05)
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line.replace('**','')
        p.alignment=align
        p._p.get_or_add_pPr().append(OxmlElement('a:buNone'))
        p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold
        p.font.color.rgb=RGBColor.from_string(color)
        p.space_after=Pt(6);p.line_spacing=1.12
    return size

def slide_text(slide):
    b=slide['body'];layout=LEGACY_ROUTE.get(slide['layout_id'],slide['layout_id'])
    if layout=='Two_Column':
        columns=b['columns']
        if len(columns)!=2:raise ValueError('雙欄版型需要恰好兩欄')
        return tuple(c['header']+'\n\n'+'\n\n'.join(c['points']) for c in columns)
    if layout=='Steps':
        modules=slide.get('modules') or []
        steps=('\n'.join(f"{m['keyword']}：{m['description']}" for m in modules)
               if modules else b['code_snippet']['code'])
        if b.get('question','').strip():steps=b['question'].strip()+'\n'+steps
        return steps, '\n\n'.join(b['bullets'])
    if layout=='Question':
        if slide['layout_id']=='quiz_qa':return '想一想\n'+b['question'],'核對解法\n'+b['answer']
        extra=b['answer'] or slide.get('hint','')
        return b['question'],('解法：'+extra if extra else '')
    if slide.get('visual_model') or slide.get('key_sentence'):
        lead='\n\n'.join(b['bullets']) or slide['subheading']
        return lead,slide.get('key_sentence','')
    return '\n\n'.join(b['bullets']) or slide['subheading'],''


def plain_page(page):
    """A page that fell back to plain text must not keep staged/diagram fields."""
    if page.get('visual_status')=='diagram_fallback':page['visual_status']='text_only'
    for key in ('visual_model','stage','hint','key_sentence','think_seconds','answer_area','stage_seq','stage_step','stage_total','staged_from'):
        page.pop(key,None)
    return page

def add_text(slide,left,top,width,height,text,max_size=22,min_size=20,color=INK,bold=False,center=False):
    shape=slide.shapes.add_textbox(Inches(left),Inches(top),Inches(width),Inches(height))
    fill(shape,text,max_size,min_size,color,bold,PP_ALIGN.CENTER if center else PP_ALIGN.LEFT)
    if center:shape.text_frame.vertical_anchor=MSO_ANCHOR.MIDDLE
    return shape


def render_modules(slide,modules,kind,style=None):
    """Editable, grid-aligned teaching modules with one clear reading order."""
    if len(modules)!=3:raise Overflow('資訊卡與步驟頁需要恰好三個教學模組')
    style=style or get_theme(None)
    for index,(module,box) in enumerate(zip(modules,card_boxes(style))):
        x,y,width,height=box
        card=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x),Inches(y),Inches(width),Inches(height))
        card.adjustments[0]=.02 if style['style']=='diagram' else .055
        card.fill.solid();card.fill.fore_color.rgb=RGBColor.from_string(style['pale'] if style['style']=='practice' else 'FDFDFC')
        card.line.color.rgb=RGBColor.from_string(style.get('hairline',style['pale']));card.line.width=Pt(1)
        keyword=(f'{index+1:02}  ' if kind=='Steps' else '')+module['keyword']
        if style['style']=='focus':
            add_text(slide,x+.24,y+.17,2.35,.92,keyword,28,24,style['primary'],True)
            add_text(slide,x+2.75,y+.17,width-3.0,.92,module['description'],26,22,style['ink'])
        else:
            add_text(slide,x+.24,y+.31,width-.48,.9,keyword,30,24,style['primary'],True)
            add_text(slide,x+.24,y+1.38,width-.48,height-1.65,module['description'],26,22,style['ink'])

def short_card_items(item):
    """Reuse model text as editable cards, without changing its wording."""
    bullets=item['body']['bullets']
    if not bullets:return []
    parts=bullets if len(bullets)>1 else [line.strip() for line in bullets[0].splitlines() if line.strip()]
    return parts if 2<=len(parts)<=9 and all(len(part)<=55 for part in parts) else []

def render_short_cards(slide,item,style):
    parts=short_card_items(item)
    if not parts:return None
    columns=1 if style['style']=='focus' and len(parts)<=4 else 2 if len(parts)<=4 else 3
    rows=math.ceil(len(parts)/columns)
    gap_x=.25;gap_y=.18;width=(11.47-gap_x*(columns-1))/columns
    height=(4.65-gap_y*(rows-1))/rows
    max_size=24 if columns==2 else 20
    min_size=20
    try:
        sizes=[font_size(part,width-.4,height-.35,max_size,min_size) for part in parts]
    except Overflow:return None
    for index,(part,size) in enumerate(zip(parts,sizes)):
        x=MARGIN+(index%columns)*(width+gap_x)
        y=1.82+(index//columns)*(height+gap_y)
        card=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x),Inches(y),Inches(width),Inches(height))
        card.adjustments[0]=.055
        card.fill.solid();card.fill.fore_color.rgb=RGBColor.from_string('FDFDFC')
        card.line.color.rgb=RGBColor.from_string(style.get('hairline',style['pale']));card.line.width=Pt(1)
        fill(card,part,size,size,style['ink'],False,PP_ALIGN.CENTER)
        card.text_frame.vertical_anchor=MSO_ANCHOR.MIDDLE
    return sizes


def add_transition(slide,kind):
    """Use one quiet chapter transition; regular teaching pages stay still."""
    if kind not in ('Cover','Section','Summary'):return
    transition=OxmlElement('p:transition')
    transition.set('spd','med')
    transition.append(OxmlElement('p:fade'))
    slide._element.insert_element_before(transition,'p:timing','p:extLst')

def size_quiz_boxes(question_box,answer_box,question,answer):
    """Give a short question only its needed height; keep the rest for the answer."""
    top=1.8;bottom=6.7;gap=.18;available=bottom-top-gap
    width=question_box.width/914400
    lines=estimated_lines(question,max(.1,width-.25),24)
    needed=lines*24*1.35/72+.30
    question_height=max(1.2,needed)
    if question_height>available-1.4:
        raise Overflow('題目過長，無法同時保留學生可讀的解答區')
    answer_height=available-question_height
    # Check both before mutating the shapes so an overflow is never silently clipped.
    font_size(question,width,question_height,30,24)
    font_size(answer,answer_box.width/914400,answer_height,28,22)
    question_box.top=Inches(top);question_box.height=Inches(question_height)
    answer_box.top=Inches(top+question_height+gap);answer_box.height=Inches(answer_height)

def size_quiz_columns(question_box,answer_box,question,answer):
    """Use the right half for an answer card when both texts remain readable."""
    try:
        font_size(question,5.45,4.85,34,22)
        font_size(answer,5.54,4.85,26,20)
    except Overflow:return False
    for box,geometry in ((question_box,(MARGIN,1.72,5.45,4.85)),
                         (answer_box,(6.86,1.72,5.54,4.85))):
        box.left,box.top,box.width,box.height=[Inches(x) for x in geometry]
    return True

def media_files(folder=MEDIA):
    folder=Path(folder)
    if not folder.is_dir():return []
    from PIL import Image
    files=[]
    for path in sorted(folder.iterdir()):
        if not path.is_file() or path.suffix.lower() not in ('.png','.jpg','.jpeg'):continue
        try:
            with Image.open(path) as picture:picture.verify()
            files.append(path)
        except (OSError,ValueError):continue
    return files

def find_media(item,files):
    if item['layout_id'] not in ('bullet_with_image','Text_Visual'):return None
    hint=item['media_hint']['search_keyword'].strip().casefold()
    words=[w for w in hint.replace('_',' ').replace('-',' ').split() if len(w)>1]
    return next((p for p in files if all(w in p.stem.casefold() for w in words)),None) if words else None

def paginate_deck(deck,folder=MEDIA):
    """Split long bullet pages deterministically, preserving every character."""
    result=copy.deepcopy(deck);result['slides']=[];files=media_files(folder)
    visual_width=visual_boxes(deck_style(deck))[1][2]
    for item in deck['slides']:
        if item['layout_id']=='Steps' and len(item.get('modules') or [])>3:
            modules=item['modules'];size=math.ceil(len(modules)/math.ceil(len(modules)/3))
            total=math.ceil(len(modules)/size)
            for start in range(0,len(modules),size):
                page=copy.deepcopy(item);page['modules']=modules[start:start+size]
                page['step_offset']=start
                page['continuation']={'part':start//size+1,'total':total}
                page['speaker_notes']+=f'\n步驟續頁：{start//size+1}/{total}；題目條件完整保留。'
                if start:page.pop('visual_model',None)
                if start+size<len(modules):page.pop('key_sentence',None)
                result['slides'].append(page)
            continue
        if item['layout_id'] not in ('bullet_with_image','summary_checklist','title_hero','Text_Visual','Summary','Cover') or not item['body']['bullets']:
            result['slides'].append(copy.deepcopy(item));continue
        width,height=(11.47,3.5) if item['layout_id'] in ('summary_checklist','title_hero','Summary','Cover') else (visual_width,4.85) if item['layout_id']=='Text_Visual' else (7.6 if find_media(item,files) else 11.47,4.8)
        minimum=24
        maximum=30 if item['layout_id']=='bullet_with_image' else 28
        def fits(parts):
            try:font_size('\n'.join(parts),width,height,maximum,minimum);return True
            except Overflow:return False
        segments=[]
        for bullet in item['body']['bullets']:
            bullet=re.sub(r'\n\s*\n+', '\n', bullet).strip()
            if not bullet:segments.append('');continue
            rest=bullet
            while rest:
                if fits([rest]):segments.append(rest);break
                low,high=1,len(rest)-1;cut=0
                while low<=high:
                    mid=(low+high)//2
                    if fits([rest[:mid]]):cut=mid;low=mid+1
                    else:high=mid-1
                if cut==0:raise Overflow('單一文字無法以學生可讀字級顯示')
                if cut<len(rest):
                    boundary=max(rest.rfind('\n',cut//2,cut),rest.rfind('。',cut//2,cut),
                                 rest.rfind('；',cut//2,cut),rest.rfind(' ',cut//2,cut))
                    if boundary>0:cut=boundary+1
                segments.append(rest[:cut]);rest=rest[cut:]
        groups=[];group=[]
        for segment in segments:
            if group and (len(group)>=3 or not fits(group+[segment])):
                groups.append(group);group=[]
            group.append(segment)
        if group:groups.append(group)
        for number,parts in enumerate(groups,1):
            page=copy.deepcopy(item);page['body']['bullets']=parts
            if len(groups)>1:page['speaker_notes']+=f'\n內容續頁：{number}/{len(groups)}'
            if number>1:page.pop('visual_model',None);page.pop('key_sentence',None)
            if number>1 and 'visual_asset' in page:
                page['visual_asset'].update(is_required=False,asset_type='none',image_prompt='')
            result['slides'].append(page)
    return result

def split_overflow_page(deck,index):
    """Keep every character while turning one overflowing page into readable pages."""
    item=deck['slides'][index]
    visual_width=visual_boxes(deck_style(deck))[1][2]
    main,secondary=slide_text(item)
    caption=model_caption(item.get('visual_model'))
    if caption and caption not in main+secondary:main+='\n'+caption
    if not main and not secondary:raise Overflow('空白內容無法切頁')
    if secondary and not item.get('visual_asset',{}).get('is_required'):
        combined=main+'\n\n'+secondary
        try:font_size(combined,11.7,4.8,30,24)
        except Overflow:pass
        else:
            page=plain_page(copy.deepcopy(item))
            page['layout_id']='bullet_with_image'
            page['body']['bullets']=[combined]
            page['media_hint']={'search_keyword':'','image_description':''}
            result=copy.deepcopy(deck);result['slides'][index]=page
            return result
    parts=[]
    for label,value in (('題目／內容',main),('解答／補充',secondary)):
        if not value:continue
        rest=re.sub(r'\n\s*\n+', '\n', value)
        while rest:
            visual_first=(not parts and item['layout_id']=='Text_Visual' and
                          item.get('visual_asset',{}).get('is_required'))
            width,height=(visual_width,4.85) if visual_first else (11.7,4.8)
            low,high=1,len(rest);cut=0
            while low<=high:
                mid=(low+high)//2
                try:font_size(rest[:mid],width,height,22 if visual_first else 30,20 if visual_first else 24);cut=mid;low=mid+1
                except Overflow:high=mid-1
            if not cut:raise Overflow('單一文字無法以學生可讀字級顯示')
            if cut<len(rest):
                boundary=max(rest.rfind('\n',0,cut),rest.rfind('。',0,cut),rest.rfind('；',0,cut))
                if boundary>=cut//2:cut=boundary+1
            page=plain_page(copy.deepcopy(item))
            page['layout_id']='Text_Visual' if visual_first else 'bullet_with_image'
            page['body']['bullets']=[rest[:cut]]
            if not visual_first:page['media_hint']={'search_keyword':'','image_description':''}
            if not visual_first and 'visual_asset' in page:
                page['visual_asset'].update(is_required=False,asset_type='none',image_prompt='')
            page['speaker_notes']+='\n切頁欄位：'+label
            parts.append(page)
            rest=rest[cut:]
    if not parts or (len(parts)==1 and parts[0]['layout_id']==item['layout_id']):
        raise Overflow('此頁內容無法再拆分')
    # One full-width page can be enough even when the original layout overflows.
    # In that case replacing the narrow layout is a successful reflow, not a failure.
    for number,page in enumerate(parts,1):
        page['speaker_notes']+=f'；內容續頁 {number}/{len(parts)}'
    result=copy.deepcopy(deck)
    result['slides'][index:index+1]=parts
    return result

def safe_fallback_page(deck,index):
    """Preserve visible teaching text in editable, readable full-width pages."""
    item=deck['slides'][index]
    main,secondary=slide_text(item)
    caption=model_caption(item.get('visual_model'))
    if caption and caption not in main+secondary:main+='\n'+caption
    segments=[part for part in (main,secondary) if part]
    if not segments:segments=[item.get('subheading') or item.get('heading') or '教材內容待補']
    chunks=[]
    for segment in segments:
        rest=format_teaching_text(segment)
        while rest:
            low,high=1,len(rest);cut=0
            while low<=high:
                mid=(low+high)//2
                try:font_size(rest[:mid],11.45,4.6,30,24);cut=mid;low=mid+1
                except Overflow:high=mid-1
            if not cut:raise Overflow('保底版面無法容納單一字元')
            chunks.append(rest[:cut]);rest=rest[cut:]
    pages=[]
    for number,chunk in enumerate(chunks,1):
        page=plain_page(copy.deepcopy(item))
        page['layout_id']='bullet_with_image'
        page['heading']=item.get('heading') or '教材內容'
        page['subheading']=''
        page['body']={'bullets':[chunk],'columns':[],'question':'','answer':'',
                      'code_snippet':{'language':'','code':''}}
        page['modules']=[]
        page['media_hint']={'search_keyword':'','image_description':''}
        if 'visual_asset' in page:
            page['visual_asset'].update(is_required=False,asset_type='none',image_prompt='')
        page['speaker_notes']=((item.get('speaker_notes') or '')+'\n【排版降級】原標題：'+
            item.get('heading','')+f'；內容續頁 {number}/{len(chunks)}。')
        pages.append(page)
    result=copy.deepcopy(deck);result['slides'][index:index+1]=pages
    return result

def local_media(deck,folder=MEDIA):
    """Match slide hints to reviewed local files; never fetch a remote image."""
    files=media_files(folder)
    matches={}
    for number,item in enumerate(deck['slides'],1):
        if item.get('visual_asset') and not item['visual_asset']['is_required']:continue
        image=find_media(item,files)
        if image:matches[number]=str(image)
    return matches


def route_visual_pages(deck):
    """Preserve teaching text while reserving a real image column when requested."""
    result=copy.deepcopy(deck)
    for item in result['slides']:
        visual=item.get('visual_asset') or {}
        if not visual.get('is_required') or item['layout_id'] not in ('Section','Three_Cards'):
            continue
        prior=item['layout_id']
        modules=item.get('modules') or []
        original=item['body']['bullets']
        module_lines=[f"{part['keyword']}：{part['description']}" for part in modules]
        item['body']['bullets']=module_lines+original
        item['modules']=modules[:1]
        item['layout_id']='Text_Visual'
        item['speaker_notes']+='\n原規劃版型：'+prior+'；因插圖教學需求改為左文右圖，原有模組已完整保留。'
    return result

def render(deck,target,media=None,sources=None,grade=None):
    deck=clean_slide_text(copy.deepcopy(deck))
    dropped={}
    for number,item in enumerate(deck['slides'],1):
        reason=sanitize_slide(item)
        if reason:dropped[number]=reason
    validate(DECK,deck)
    style=deck_style(deck,grade)
    primary,ink,pale,accent=style['primary'],style['ink'],style['pale'],style['accent']
    from .template_builder import layout_map
    p=Presentation(make_template(None,style));mapping=layout_map()
    report=[]
    for number,item in enumerate(deck['slides'],1):
        kind=LEGACY_ROUTE.get(item['layout_id'],item['layout_id'])
        route=mapping.get(kind)
        if route is None or not 0<=route.get('layout_index',-1)<len(p.slide_layouts):
            raise ValueError(f'第{number}頁「{item["heading"]}」：母片缺少 {item["layout_id"]} 版型')
        layout=p.slide_layouts[route['layout_index']]
        s=p.slides.add_slide(layout)
        add_transition(s,kind)
        if item.get('page_role'):apply_role(s,item,style)
        ph={x.placeholder_format.idx:x for x in s.placeholders}
        # A damaged/older master must not discard the teaching page. Supply
        # missing editable text boxes at the known grid positions by index.
        defaults={0:(MARGIN,.67,11.47,.78),1:(MARGIN,1.72,5.45,4.85),
                  2:(6.86,1.72,5.54,4.85)}
        for idx in (0,1,2):
            if idx not in ph:
                box=defaults[idx]
                ph[idx]=s.shapes.add_textbox(*[Inches(x) for x in box])
        for placeholder in layout.placeholders:
            dest=ph.get(placeholder.placeholder_format.idx)
            if dest is None:continue
            dest.left,dest.top,dest.width,dest.height=placeholder.left,placeholder.top,placeholder.width,placeholder.height
        legacy=item['layout_id'] in LEGACY_LAYOUTS
        if legacy:
            old={0:(.65,.45,12.0,.95),1:(.8,1.8,7.6,4.8),2:(8.7,1.8,3.8,4.8)}
            if item['layout_id'] in ('title_hero','summary_checklist'):
                old={0:(.8,.75,11.7,1.45),1:(.9,2.55,11.5,3.5),2:(.9,6.2,11.5,.5)}
            elif item['layout_id']=='two_column_comparison':
                old.update({1:(.8,1.85,5.6,4.8),2:(6.9,1.85,5.6,4.8)})
            elif item['layout_id']=='quiz_qa':
                old.update({1:(.8,1.8,11.7,2.1),2:(.8,4.35,11.7,2.15)})
            for index,box in old.items():
                ph[index].left,ph[index].top,ph[index].width,ph[index].height=[Inches(x) for x in box]
        main,secondary=slide_text(item)
        quiz_columns=False
        picture=(media or {}).get(number) if kind in ('Text_Visual','Cover') else None
        special=uses_special(item,kind,picture) and not legacy
        if kind=='Question' and not special:
            try:
                quiz_columns=style['style']!='focus' and size_quiz_columns(ph[1],ph[2],main,secondary)
                if not quiz_columns:size_quiz_boxes(ph[1],ph[2],main,secondary)
            except Overflow as exc:raise Overflow(f'第{number}頁「{item["heading"]}」：{exc}') from exc
        picture=(media or {}).get(number) if kind in ('Text_Visual','Cover') else None
        if kind=='Cover' and picture:
            first,second=visual_boxes(style)[1],visual_boxes(style)[2]
            ph[0].left,ph[0].top,ph[0].width,ph[0].height=[Inches(x) for x in (MARGIN,1.02,first[2],1.45)]
            ph[1].left,ph[1].top,ph[1].width,ph[1].height=[Inches(x) for x in (MARGIN,2.7,first[2],3.5)]
            ph[2].left,ph[2].top,ph[2].width,ph[2].height=[Inches(x) for x in second]
        try:
            if special:
                sizes=[fill(ph[0],item['heading'],beautify.title_pt(style,item['heading']),30,
                            accent if kind=='Question' else ink,True)]
                if kind=='Question' or (kind=='Summary' and item['body'].get('question')):sizes+=compose_question(s,item,style,ph)
                elif kind=='Steps':sizes+=compose_steps(s,item,style,ph)
                else:sizes+=compose_diagram(s,item,style,ph,kind)
            elif legacy:
                if item['layout_id']=='bullet_with_image' and not picture:ph[1].width=Inches(11.7)
                title_max=40 if len(item['heading'])<=12 else 32
                cover_page=kind in ('Cover','Section')
                lead=beautify.cover_colors(style)[0] if cover_page else primary
                body_ink=beautify.cover_colors(style)[1] if cover_page else ink
                sizes=[fill(ph[0],item['heading'],title_max,28,lead,True)]
                cards=render_short_cards(s,item,style) if not picture and item['layout_id']=='bullet_with_image' else None
                sizes.extend(cards if cards is not None else [fill(ph[1],main,30,20,body_ink)])
                if secondary:sizes.append(fill(ph[2],secondary,28,22,body_ink))
            elif kind in ('Cover','Section'):
                centered=style['style'] in ('practice','focus') and not picture
                alignment=PP_ALIGN.CENTER if centered else PP_ALIGN.LEFT
                cover_ink,intro_ink,label_ink=beautify.cover_colors(style)
                sizes=[fill(ph[0],item['heading'],beautify.cover_pt(style,item['heading']),32,cover_ink,True,alignment)]
                intro=item['subheading']
                if main and main!=intro:intro+=(('\n\n' if intro else '')+main)
                if intro:sizes.append(fill(ph[1],intro,28 if style['grade']<=6 else 26,22,intro_ink,align=alignment))
                if kind=='Section' and item.get('modules'):
                    detail=item['modules'][0]['keyword']+('\n' if style['style']=='story' else '：')+item['modules'][0]['description']
                    sizes.append(fill(ph[2],detail,22,20,label_ink,True))
                elif kind=='Cover' and not picture:
                    age=('國小 '+str(style['grade'])+' 年級' if style['grade']<=6 else '國中 '+str(style['grade']-6)+' 年級')
                    label=age+'\n'+style['subject'] if style['style']=='story' else age+'　'+style['subject']
                    sizes.append(fill(ph[2],label,32 if style['style']=='story' else 24,20,label_ink,True,alignment))
            else:
                sizes=[fill(ph[0],item['heading'],beautify.title_pt(style,item['heading']),30,
                            accent if kind=='Question' else ink,True)]
                if kind in ('Three_Cards','Steps') and len(item.get('modules') or [])==3:
                    render_modules(s,item['modules'],kind,style)
                    sizes.extend([30,20])
                elif kind in ('Three_Cards','Steps'):
                    ph[1].width=Inches(11.47)
                    sizes.append(fill(ph[1],main+('\n\n'+secondary if secondary else ''),28,24,ink))
                elif kind=='Question':
                    if style['style']=='practice':
                        ph[1].fill.solid();ph[1].fill.fore_color.rgb=RGBColor.from_string(pale)
                        ph[1].line.color.rgb=RGBColor.from_string(primary)
                    ph[2].fill.solid();ph[2].fill.fore_color.rgb=RGBColor.from_string(pale)
                    ph[2].line.color.rgb=RGBColor.from_string(primary)
                    question_size=(36 if len(main)<=70 else 26) if quiz_columns else (54 if len(main)<=40 else 36 if len(main)<=85 else 26)
                    sizes.append(fill(ph[1],main,question_size,22 if quiz_columns else 24,primary,True))
                    ph[1].text_frame.vertical_anchor=MSO_ANCHOR.MIDDLE
                    sizes.append(fill(ph[2],secondary,26 if quiz_columns else 24,20,ink))
                    if secondary and not item.get('stage'):
                        from .animations import appear_answers
                        appear_answers(s,[ph[2]])
                elif kind=='Summary':
                    bullets=item['body']['bullets']
                    sizes.append(fill(ph[1],bullets[0],44,32,primary,True))
                    ph[1].text_frame.vertical_anchor=MSO_ANCHOR.MIDDLE
                    if len(bullets)>1:
                        sizes.append(fill(ph[2],'\n'.join(bullets[1:]),22,20,ink))
                elif kind=='Text_Visual':
                    if not picture:ph[1].width=Inches(11.47)
                    cards=render_short_cards(s,item,style) if not picture else None
                    if cards is None and beautify.is_key_point(item,main,secondary,picture):
                        sizes.append(beautify.key_point(s,ph,main,style))
                    else:sizes.extend(cards if cards is not None else [fill(ph[1],main,28,24,ink)])
                elif kind=='Two_Column' and style['style']=='focus':
                    for index,column in enumerate(item['body']['columns'],1):
                        box=ph[index];x,y,w,h=[v/914400 for v in (box.left,box.top,box.width,box.height)]
                        box.text_frame.clear()
                        add_text(s,x+.18,y+.05,w-.36,.73,column['header'],30,24,primary,True)
                        add_text(s,x+.18,y+.85,w-.36,h-.94,'\n'.join(column['points']),26,24,ink)
                    sizes.extend([30,24])
                else:
                    sizes.append(fill(ph[1],main,28,24,ink))
                    if secondary:
                        focus=(kind=='Two_Column' and ('注意' in item['heading'] or '重點' in item['heading'])
                               and len(secondary)<=45)
                        if focus:
                            ph[2].fill.solid();ph[2].fill.fore_color.rgb=RGBColor.from_string(pale)
                        sizes.append(fill(ph[2],secondary,60 if focus else 28,28 if focus else 24,
                                          primary if focus else ink,focus,
                                          PP_ALIGN.CENTER if focus else PP_ALIGN.LEFT))
            if picture:
                box=ph[2]
                beautify.place_picture(s,picture,(box.left,box.top,box.width,box.height),style)
        except Overflow as exc:raise Overflow(f'第{number}頁「{item["heading"]}」：{exc}') from exc
        beautify.decorate(s,item,style,kind,number,len(deck['slides']),ph[0])
        if '未經 AI 審核' in item['speaker_notes']:
            badge=s.shapes.add_textbox(Inches(.93),Inches(6.82),Inches(5.5),Inches(.5))
            fill(badge,'本頁需教師確認',16,14,accent,True)
        credit=(sources or {}).get(str(number))
        if credit and picture and not credit.get('generated'):
            from .images import credit_line
            beautify.credit_line(s,credit_line(credit),style)
        continuation=re.search(r'內容續頁[： ]*(\d+)\s*/\s*(\d+)',item['speaker_notes'])
        if continuation:
            progress=s.shapes.add_textbox(Inches(10.25),Inches(6.94),Inches(1.25),Inches(.46))
            fill(progress,continuation.group(1)+' / '+continuation.group(2),14,12,primary,True,PP_ALIGN.RIGHT)
        notes=item['speaker_notes']+'\n主教材頁碼：'+','.join(map(str,item['source_pages']))
        source=(sources or {}).get(str(number))
        if source:
            notes+='\n圖片來源：'+('本機 AI 生成' if source.get('generated') else source['description_url'])
            notes+='\n授權：'+source['license']+' '+source['license_url']
            if source.get('teaching_purpose'):notes+='\n插圖教學用途：'+source['teaching_purpose']
        s.notes_slide.notes_text_frame.text=format_teaching_text(notes)
        beautify.finish(s)
        # PowerPoint's stock shape theme adds a shadow even to a flat card.
        # All fills, lines and text are explicit; discard that inherited effect.
        for shape in s.shapes:
            inherited=shape._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
            if inherited is not None:shape._element.remove(inherited)
        report.append({'slide':number,'layout':item['layout_id'],'template_id':style['id'],'font_sizes':sizes,
                       'local_image':str(picture) if picture else '','special':bool(special),
                       'stage':item.get('stage',''),'visual_model':(item.get('visual_model') or {}).get('type',''),
                       'visual_model_dropped':dropped.get(number,'')})
    target=Path(target);target.parent.mkdir(parents=True,exist_ok=True)
    tmp=target.with_suffix('.partial.pptx');p.save(tmp);tmp.replace(target)
    return report
