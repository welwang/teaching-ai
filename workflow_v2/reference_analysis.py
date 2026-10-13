"""Inspect reference presentation design metadata without reading teaching text."""
import hashlib
import json
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from filelock import FileLock

from .store import ROOT, DATA, atomic

REFERENCE = Path(r'C:\Users\wel30\Downloads\優良教材 (1)\優良教材')
OUTPUT = ROOT / 'assets/workflow-v2/reference-design-analysis.json'


def subject_for(path):
    name = str(path)
    return '數學' if '數學' in name else '英語' if '英文' in name else '國語' if '國文' in name else '未標示'


def presentation_stats(path):
    from pptx import Presentation
    deck=Presentation(path)
    colors, fonts, sizes=Counter(),Counter(),Counter()
    pictures=0
    for slide in list(deck.slides)[:min(8,len(deck.slides))]:
        for shape in slide.shapes:
            if shape.shape_type==13:pictures+=1
            if not shape.has_text_frame:continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.name:fonts[run.font.name]+=1
                    if run.font.size:sizes[round(run.font.size.pt)]+=1
                    try:
                        color=run.font.color.rgb if run.font.color else None
                        if color:colors[str(color)]+=1
                    except (AttributeError,TypeError):pass
    return {'slides':len(deck.slides),'aspect':round(deck.slide_width/deck.slide_height,3),
            'pictures_first_8':pictures,'colors':colors.most_common(5),
            'fonts':fonts.most_common(3),'font_sizes':sizes.most_common(4)}


def legacy_stats(path):
    with tempfile.TemporaryDirectory(dir=DATA) as folder:
        folder=Path(folder)
        from .office import convert
        result=convert(path,folder,format='pptx',timeout=50)
        target=folder/(path.stem+'.pptx')
        if result.returncode or not target.is_file():
            raise ValueError('LibreOffice conversion failed')
        return presentation_stats(target)


def other_stats(path):
    if path.suffix.lower()=='.pdf':
        import pymupdf
        with pymupdf.open(path) as pdf:
            if not len(pdf):raise ValueError('Empty PDF')
            return {'slides':len(pdf),'aspect':round(pdf[0].rect.width/pdf[0].rect.height,3)}
    if path.suffix.lower()=='.odp':
        with zipfile.ZipFile(path) as odp:
            import xml.etree.ElementTree as ET
            xml=ET.fromstring(odp.read('content.xml'))
            pages=sum(node.tag.endswith('}page') for node in xml.iter())
            return {'slides':pages}
    return {}


def analyze(root=REFERENCE, output=OUTPUT):
    from pptx import Presentation
    root = Path(root)
    if not root.is_dir():
        raise FileNotFoundError(root)
    entries = []
    for path in sorted(root.rglob('*')):
        if not path.is_file():
            continue
        relative = path.relative_to(root)
        item = {'id': hashlib.sha256(str(relative).encode('utf-8')).hexdigest()[:12],
                'format': path.suffix.lower(), 'bytes': path.stat().st_size,
                'subject': subject_for(relative), 'status': 'cataloged'}
        suffix=path.suffix.lower()
        if suffix in ('.pptx','.ppt','.pdf','.odp'):
            try:
                if suffix=='.pptx' and path.stat().st_size>80_000_000:
                    with zipfile.ZipFile(path) as ppt:
                        slides=sum(name.startswith('ppt/slides/slide') and name.endswith('.xml')
                                   for name in ppt.namelist())
                    item.update(status='inspected_package',slides=slides)
                else:
                    details=(presentation_stats(path) if suffix=='.pptx' else
                             legacy_stats(path) if suffix=='.ppt' else other_stats(path))
                    item.update(status='inspected',**details)
            except Exception as exc:
                item.update(status='unreadable', error=type(exc).__name__)
        entries.append(item)
    report = {'source_kind': 'local_design_reference_only',
              'file_count': len(entries),
              'formats': dict(Counter(item['format'] for item in entries)),
              'inspection': dict(Counter(item['status'] for item in entries)),
              'entries': entries}
    atomic(output, report)
    return report


if __name__ == '__main__':
    print(json.dumps({k: v for k, v in analyze().items() if k != 'entries'}, ensure_ascii=False))
