"""Exercise the eight student-facing layouts without calling a model."""
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation

from workflow_v2.contracts import LAYOUTS,VISUAL_STYLE,slide_schema,validate
from workflow_v2.renderer import render


def sample_deck():
    layouts=[
        ('Cover','角度探險','一起認識各種角',[]),
        ('Section','第一站：認識角','從生活中的轉彎開始',[('觀察','找出兩條邊的共同端點')]),
        ('Text_Visual','角的大小','看開口，不看邊長',[('關鍵','角的大小取決於兩邊張開的程度')]),
        ('Two_Column','比一比','銳角和鈍角',[]),
        ('Three_Cards','三種角','用直角當標準',[('銳角','比直角小'),('直角','恰好九十度'),('鈍角','比直角大')]),
        ('Steps','量角三步驟','拿起量角器',[('對中心','中心對準頂點'),('對零線','零線貼齊一邊'),('讀刻度','讀另一邊指到的刻度')]),
        ('Question','想一想','先自己作答',[]),
        ('Summary','今天學會了','用直角來比較角的大小',[]),
    ]
    slides=[]
    for name,heading,subtitle,items in layouts:
        body={'bullets':['角有一個頂點和兩條邊。','比較角時，觀察兩邊張開的程度。'],
              'columns':[{'header':'銳角','points':['比直角小']},{'header':'鈍角','points':['比直角大']}]
              if name=='Two_Column' else [],
              'question':'一個角是 30°。它比直角大還是小？',
              'answer':'30° 小於 90°，所以比直角小。',
              'code_snippet':{'language':'text','code':''}}
        if name in ('Cover','Section','Two_Column','Three_Cards','Steps','Question'):body['bullets']=[]
        slides.append({'layout_id':name,'heading':heading,'subheading':subtitle,
            'body':body,'modules':[{'keyword':k,'description':d,'visual_hint':''} for k,d in items],
            'visual_asset':{'is_required':False,'reasoning':'本頁以文字與版型呈現即可',
                'asset_type':'none','image_prompt':'','style_tags':VISUAL_STYLE},
            'media_hint':{'search_keyword':'','image_description':''},
            'speaker_notes':'教師提問：你看過哪些角？預估講解 30 秒。','source_pages':[1]})
    return {'schema_version':1,'deck_meta':{'title':'角度探險','subject':'數學',
            'target_audience':'國小學童','primary_color':'#008080'},'slides':slides}


class DesignSystemTest(unittest.TestCase):
    def test_short_points_become_cards_and_quiz_uses_two_columns(self):
        deck=sample_deck()
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'cards.pptx'
            render(deck,path)
            prs=Presentation(path)
            cards=[shape for shape in prs.slides[2].shapes if shape.has_text_frame and
                   shape.text in deck['slides'][2]['body']['bullets']]
            self.assertEqual(len(cards),2)
            quiz=prs.slides[6]
            question=next(shape for shape in quiz.shapes if shape.has_text_frame and
                          deck['slides'][6]['body']['question'] in shape.text)
            answer=next(shape for shape in quiz.shapes if shape.has_text_frame and
                        deck['slides'][6]['body']['answer'] in shape.text)
            self.assertLess(question.left+question.width,answer.left)
    def test_new_layout_schema_and_render(self):
        deck=sample_deck()
        self.assertEqual([slide['layout_id'] for slide in deck['slides']],LAYOUTS)
        for slide in deck['slides']:
            validate(slide_schema({1},slide['layout_id']),slide)
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'design.pptx'
            report=render(deck,path)
            prs=Presentation(path)
            self.assertEqual(len(prs.slides),8)
            self.assertEqual(len(report),8)
            self.assertEqual([prs.slide_layouts[i].name for i in range(8)],LAYOUTS)
            self.assertTrue(all(len(prs.slides[i].shapes)>3 for i in (4,5)))
            self.assertEqual([i for i,slide in enumerate(prs.slides)
                if slide._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}transition') is not None],
                [0,1,7])

    def test_card_text_is_bounded(self):
        slide=sample_deck()['slides'][4]
        slide['modules'][0]['description']='很長的描述'*30
        with self.assertRaises(ValueError):validate(slide_schema({1},'Three_Cards'),slide)


if __name__=='__main__':unittest.main()
