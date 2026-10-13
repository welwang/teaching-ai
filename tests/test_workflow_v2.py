import json,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
from workflow_v2.store import Store,STAGES,atomic
from workflow_v2.contracts import validate_slide,validate,slide_schema,LAYOUTS,LEGACY_LAYOUTS
from workflow_v2.renderer import render,Overflow
from workflow_v2.service import regenerate
from workflow_v2.model_policy import DEFAULT_MODEL

REQUEST=dict(schema_version=1,grade=6,subject='數學',semester='上學期',unit='測試單元',publisher='測試',difficulty='全部',slide_count=8,source_path='test.md',model=DEFAULT_MODEL,locale='zh-TW')
def publish_fixture(request,out):

    from workflow_v2.pipeline import publish
    def completed(*args,**kwargs):
        report={'status':'completed','changed':False}
        atomic(out/'final-optimization.json',report);return report
    with patch('workflow_v2.final_optimizer.optimize',side_effect=completed):publish(request,out)

def slide(layout='bullet_with_image'):
    return dict(layout_id=layout,heading='版型測試',subheading='學生閱讀測試',speaker_notes='講解30秒，請學生觀察。',source_pages=[1],
        media_hint=dict(search_keyword='',image_description=''),body=dict(bullets=['完整的一句話'],columns=[dict(header='甲',points=['第一欄']),dict(header='乙',points=['第二欄'])],
        code_snippet=dict(language='text',code='步驟一\n步驟二'),question='這是什麼？',answer='完整的答案。'))

class WorkflowTest(unittest.TestCase):
    def test_math_text_repairs_escaped_arrows_without_changing_plain_words(self):
        from workflow_v2.text_safety import format_teaching_text
        broken='6k = 18'+chr(13)+'ightarrow k = 3'
        self.assertEqual(format_teaching_text(broken),'6k = 18 →  k = 3')
        self.assertIn('→',format_teaching_text('18_x000D_ightarrow 3'))
        self.assertIn('→',format_teaching_text('18_x000D_rightarrow 3'))
        self.assertEqual(format_teaching_text(r'\(a:b:c = 1:2:3\)'), 'a:b:c = 1:2:3')
        self.assertEqual(format_teaching_text(r'1 \times 3 \rightarrow 3'),'1 × 3 → 3')
        self.assertEqual(format_teaching_text('three times; $100 和 $200'),'three times; $100 和 $200')
    def test_old_slide_math_is_normalized_when_rendered(self):
        from pptx import Presentation
        from workflow_v2.pipeline import normalize_slide
        original=slide()
        original['body']['bullets']=[r'\(a:b:c = 1:2:3\)',
            '6k = 18_x000D_ightarrow k = 3',r'1 \times 3 = 3']
        original['speaker_notes']=r'教師說明：\(1 \times 3\)'
        normalized=normalize_slide(json.loads(json.dumps(original)))
        self.assertIn('×',normalized['body']['bullets'][2])
        self.assertNotIn('_x000D_',normalized['body']['bullets'][1])
        deck=dict(schema_version=1,deck_meta=dict(title='測試',subject='數學',
            target_audience='六年級',primary_color='#087F78'),slides=[original]+[slide() for _ in range(7)])
        with tempfile.TemporaryDirectory() as t:
            path=Path(t)/'math.pptx';render(deck,path)
            page=Presentation(path).slides[0]
            visible=' '.join(shape.text for shape in page.shapes if shape.has_text_frame)
            self.assertIn('a:b:c = 1:2:3',visible)
            self.assertIn('→',visible)
            self.assertIn('×',visible)
            self.assertNotIn('_x000D_',visible)
            self.assertNotIn(r'\times',visible)
            self.assertNotIn(r'\(',visible)
            self.assertIn('×',page.notes_slide.notes_text_frame.text)
    def test_age_guidance_follows_grade_and_subject(self):
        from workflow_v2.pedagogy import developmental_guidance,system_prompt,audience_for
        from workflow_v2.pipeline import page_prompt
        younger=dict(REQUEST,grade=4,audience='大學專題')
        older=dict(REQUEST,grade=9,subject='英語')
        self.assertIn('國小四年級',system_prompt(younger))
        self.assertEqual(audience_for(younger),'國小學童')
        self.assertNotIn('精確定義術語',system_prompt(younger))
        self.assertIn('年級要求優先',developmental_guidance(younger))
        self.assertIn('國中三年級',system_prompt(older))
        self.assertIn('生詞需要例句',system_prompt(older))
        item={'phase':'concept','difficulty':'低難度','source_pages':[1]}
        source={'chunks':[{'page':1,'text':'教學來源'}]}
        with tempfile.TemporaryDirectory() as t:
            prompt=page_prompt(younger,{'objectives':['理解']},item,source,Path(t))
        self.assertIn('國小四年級',prompt)
        self.assertIn('先給具體例子或圖像',prompt)
        self.assertIn('先辨認與模仿',prompt)
    def test_engine_prefers_workspace_runtime_over_stale_cached_path(self):
        from workflow_v2 import engine
        with tempfile.TemporaryDirectory() as t:
            root=Path(t)
            local=root/'runtime/ollama/ollama.exe'
            local.parent.mkdir(parents=True)
            local.touch()
            cache=root/'ollama-binary.txt'
            cache.write_text('Z:/unavailable/ollama.exe',encoding='utf-8')
            with patch.object(engine,'ROOT',root),patch.object(engine,'BINARY_FILE',cache):
                self.assertEqual(engine._binary(),local)
    def test_shutdown_distinguishes_control_ui_from_worker(self):
        import shutdown
        class Process:
            pid=999999
            def __init__(self,command):self.command=command
            def cmdline(self):return self.command
            def cwd(self):return str(shutdown.ROOT)
        ui=Process(['python','-m','streamlit','run','app_v2.py'])
        worker=Process(['python','-m','workflow_v2.worker'])
        self.assertTrue(shutdown._owned_ui(ui))
        self.assertFalse(shutdown._owned_worker(ui))
        self.assertTrue(shutdown._owned_worker(worker))
        self.assertFalse(shutdown._owned_ui(worker))
    def test_output_filenames_include_conditions_and_are_windows_safe(self):
        from workflow_v2.pipeline import output_stem
        request=dict(REQUEST,publisher='康軒:版',unit='一／二? 三*',difficulty='中難度')
        stem=output_stem(request)
        self.assertIn('康軒_版_國小六年級_上學期_數學_',stem)
        self.assertTrue(stem.endswith('_中難度'))
        self.assertFalse(any(char in stem for char in '<>:"/\\|?*'))
        self.assertNotIn('.pdf',output_stem(dict(request,unit='角的大小.pdf')))
    def test_publish_exposes_condition_named_artifacts(self):
        from workflow_v2.pipeline import publish
        with tempfile.TemporaryDirectory() as t:
            out=Path(t)
            atomic(out/'rendered-deck.json',{'slides':[slide()]})
            (out/'教材.pptx').write_bytes(b'pptx')
            (out/'教材.pdf').write_bytes(b'pdf')
            publish_fixture(REQUEST,out)
            result=json.loads((out/'result.json').read_text(encoding='utf-8'))
            self.assertEqual(result['quality'],'ai_reviewed_draft')
            self.assertTrue(all(Path(path).is_file() for path in result['artifacts'].values()))
            self.assertTrue(Path(result['artifacts']['pptx']).name.startswith('測試_國小六年級_上學期_數學_測試單元_全部難度'))
    def test_publish_strips_xml_control_characters_without_losing_teaching_text(self):
        from docx import Document
        from workflow_v2.pipeline import publish
        with tempfile.TemporaryDirectory() as t:
            out=Path(t)
            value=slide();value['heading']='分數\x00加法';value['speaker_notes']='示範\x0b通分\n請學生練習'
            atomic(out/'rendered-deck.json',{'slides':[value]})
            (out/'教材.pptx').write_bytes(b'pptx')
            (out/'教材.pdf').write_bytes(b'pdf')
            publish_fixture(REQUEST,out)
            content='\n'.join(p.text for p in Document(out/'教師備課稿.docx').paragraphs)
            self.assertIn('分數加法',content)
            self.assertIn('示範通分\n請學生練習',content)
            self.assertNotIn('\x00',content)
            self.assertNotIn('\x0b',content)
    def test_local_worksheet_produces_student_copy_and_separate_answers(self):
        from docx import Document
        from workflow_v2.worksheet import generate
        from workflow_v2.pipeline import publish
        request=dict(REQUEST,include_worksheet=True)
        questions=[]
        for difficulty in ('低難度','中難度','高難度'):
            for number in range(3):
                questions.append({'difficulty':difficulty,'question':f'{difficulty}題目 {number+1}：2 + {number+1} = ？',
                    'answer':str(number+3),'explanation':f'把 2 與 {number+1} 相加。','source_pages':[1]})
        worksheet={'instructions':'請先獨立作答，再和同學討論。','questions':questions}
        with tempfile.TemporaryDirectory() as t:
            out=Path(t)
            atomic(out/'source.json',{'pages':[{'page':1,'text':'加法與解題。'}]})
            atomic(out/'plan.json',{'slides':[{'topic':'加法','difficulty':'共用',
                'phase':'example','source_pages':[1]}]})
            atomic(out/'reviewed-deck.json',{'slides':[slide()]})
            atomic(out/'rendered-deck.json',{'slides':[slide()]})
            (out/'教材.pptx').write_bytes(b'pptx')
            (out/'教材.pdf').write_bytes(b'pdf')
            stems={'instructions':worksheet['instructions'],'questions':[
                {k:v for k,v in q.items() if k not in ('answer','explanation')} for q in questions]}
            answer_groups=[{'answers':[{'id':i+1,'answer':q['answer'],'explanation':q['explanation']}
                for i,q in enumerate(questions) if start<=i<start+3]} for start in range(0,9,3)]
            with patch('workflow_v2.worksheet.call',side_effect=[stems,*answer_groups,{'passed':True,'issues':[]}]) as model:
                generate(request,out)
            self.assertEqual(model.call_count,5)
            publish_fixture(request,out)
            result=json.loads((out/'result.json').read_text(encoding='utf-8'))
            self.assertEqual(result['worksheet_status'],'done')
            student='\n'.join(p.text for p in Document(result['artifacts']['worksheet_docx']).paragraphs)
            teacher='\n'.join(p.text for p in Document(result['artifacts']['worksheet_answers_docx']).paragraphs)
            self.assertIn('低難度題目 1',student)
            self.assertNotIn('答案：',student)
            self.assertIn('答案：3',teacher)
            self.assertIn('主教材頁碼：1',teacher)
    def test_worksheet_failure_keeps_slide_artifacts(self):
        from workflow_v2.pipeline import publish
        request=dict(REQUEST,include_worksheet=True)
        with tempfile.TemporaryDirectory() as t:
            out=Path(t)
            atomic(out/'rendered-deck.json',{'slides':[slide()]})
            (out/'教材.pptx').write_bytes(b'pptx')
            (out/'教材.pdf').write_bytes(b'pdf')
            publish_fixture(request,out)
            result=json.loads((out/'result.json').read_text(encoding='utf-8'))
            self.assertEqual(result['worksheet_status'],'failed')
            self.assertIn('pptx',result['artifacts'])
            self.assertNotIn('worksheet_docx',result['artifacts'])
    def test_heavy_context_cap_and_transient_http_retry(self):
        from contextlib import nullcontext
        from workflow_v2.model import call
        class Response:
            def __init__(self,status):
                self.status_code=status;self.ok=status==200
                self.text='temporary backend error' if status!=200 else ''
            def json(self):
                return {'message':{'content':'{"answer":"ready"}'},'done_reason':'stop'}
        schema={'type':'object','properties':{'answer':{'type':'string'}},
                'required':['answer'],'additionalProperties':False}
        with patch('workflow_v2.model.gpu_slot',return_value=nullcontext()), \
             patch('workflow_v2.model_policy.settings',return_value={'generation_context':8192}), \
             patch('workflow_v2.model.release_image_checkpoint'), \
             patch('workflow_v2.model.requests.post',side_effect=[Response(500),Response(200)]) as post, \
             patch('workflow_v2.model.unload_model') as unload, \
             patch('workflow_v2.model.time.sleep') as sleep, \
             patch('workflow_v2.model.log_event') as log:
            self.assertEqual(call(DEFAULT_MODEL,'文字'*4000,schema,attempts=2),{'answer':'ready'})
        self.assertEqual([c.kwargs['json']['options']['num_ctx'] for c in post.call_args_list],[8192,8192])
        sleep.assert_called_once_with(5)
        unload.assert_called_once_with(DEFAULT_MODEL)
        self.assertTrue(any(c.args[0]=='model_retry_scheduled' and c.kwargs['retry_count']==1
                            for c in log.call_args_list))
    def test_only_missing_json_closers_are_repaired(self):
        from workflow_v2.model import parse_complete_json
        self.assertEqual(parse_complete_json('{"a":[1,2'),{'a':[1,2]})
        with self.assertRaises(ValueError):parse_complete_json('{"a":"unfinished')
        with self.assertRaises(ValueError):parse_complete_json('{"a":1]')
    def test_failed_page_is_a_teacher_prompt_without_textbook_text_and_has_quality_marker(self):
        from workflow_v2.pipeline import fallback_page
        with tempfile.TemporaryDirectory() as t:
            out=Path(t)
            atomic(out/'plan.json',{'slides':[{'topic':'來源主題','source_pages':[2]}]})
            atomic(out/'source.json',{'pages':[{'page':2,'text':'原文 12 × 4 = 48。'}]})
            page=fallback_page(REQUEST,out,0,'模型退件')
            shown=json.dumps(page,ensure_ascii=False)
            self.assertNotIn('12 × 4 = 48',shown)
            self.assertIn('第 2 頁',page['body']['bullets'][1])
            self.assertIn('未經 AI 審核',page['speaker_notes'])
            warning=json.loads((out/'quality-warnings.json').read_text(encoding='utf-8'))['1']
            self.assertEqual(warning['mode'],'teacher_scaffold_no_source');self.assertFalse(warning['full_source_in_notes'])
    def test_bounded_write_failure_renders_readable_source_fallback(self):
        from workflow_v2.pipeline import write_one,deck
        from workflow_v2.renderer import paginate_deck
        with tempfile.TemporaryDirectory() as t:
            out=Path(t);source='主教材原文：一個角有一個頂點與兩條邊。'
            plan={'objectives':['認識角'],'slides':[{'topic':'角的認識','phase':'concept',
                'difficulty':'低難度','layout_id':'Text_Visual','source_pages':[1]}]*8}
            atomic(out/'plan.json',plan)
            atomic(out/'source.json',{'pages':[{'page':1,'text':source}],'chunks':[]})
            with patch('workflow_v2.pipeline.author_page',side_effect=ValueError('內容無法驗證')) as author:
                write_one(REQUEST,out,0,plan['slides'][0])
            self.assertEqual(author.call_count,2)
            self.assertNotIn(source,json.dumps(json.loads((out/'slides/001.json').read_text(encoding='utf-8')),ensure_ascii=False))
            for i in range(1,8):atomic(out/'slides'/f'{i+1:03}.json',slide())
            target=out/'fallback.pptx'
            render(paginate_deck(deck(REQUEST,out)),target)
            prs=__import__('pptx').Presentation(target)
            self.assertNotIn(source,' '.join(shape.text for shape in prs.slides[0].shapes if shape.has_text_frame))
            self.assertIn('本頁需教師確認',' '.join(shape.text for shape in prs.slides[0].shapes if shape.has_text_frame))
    def test_long_fallback_never_copies_the_source_into_slide_or_notes(self):
        from workflow_v2.pipeline import fallback_page
        from workflow_v2.renderer import paginate_deck
        with tempfile.TemporaryDirectory() as t:
            out=Path(t);source='教材完整原文。'*600
            atomic(out/'plan.json',{'slides':[{'topic':'來源主題','source_pages':[1]}]})
            atomic(out/'source.json',{'pages':[{'page':1,'text':source}]})
            page=fallback_page(REQUEST,out,0,'模型退件')
            self.assertNotIn('教材完整原文。教材完整原文。',json.dumps(page,ensure_ascii=False))
            self.assertLess(len(page['speaker_notes']),800)
            self.assertLess(len(page['body']['bullets'][0]),250)
            deck={'schema_version':1,'deck_meta':{'title':'來源主題','subject':'數學',
                'target_audience':'國小學童','primary_color':'#008080'},'slides':[page]}
            self.assertLessEqual(len(paginate_deck(deck)['slides']),2)
    def test_runtime_log_tracks_job_without_textbook_or_prompt(self):
        from workflow_v2.runlog import log_event
        with tempfile.TemporaryDirectory() as t:
            s=Store(t);jid=s.submit(REQUEST)
            log_event('model_call_finished',jid,root=s.root,stage='write',model=DEFAULT_MODEL,
                      duration_seconds=2.5,prompt_tokens=100,output_tokens=50)
            records=[json.loads(line) for line in (s.folder(jid)/'run.jsonl').read_text(encoding='utf-8').splitlines()]
            self.assertEqual(records[-1]['event'],'model_call_finished')
            self.assertEqual(records[-1]['output_tokens'],50)
            self.assertNotIn('prompt',records[-1])
    def test_model_router_uses_draft_model_then_request_model_for_repair(self):
        from workflow_v2.pipeline import author_page
        request=dict(REQUEST,model='example:repair',page_model='example:draft')
        with tempfile.TemporaryDirectory() as t:
            out=Path(t)
            atomic(out/'plan.json',{'objectives':['目標'],'slides':[{'layout_id':'bullet_with_image','source_pages':[1]}]})
            atomic(out/'source.json',{'chunks':[]})
            with patch('workflow_v2.pipeline.page_prompt',return_value='提示'), \
                 patch('workflow_v2.pipeline.call',return_value=slide()) as model_call:
                author_page(request,out,0)
                self.assertEqual(model_call.call_args.args[0],'example:draft')
                author_page(request,out,0,feedback='修正')
                self.assertEqual(model_call.call_args.args[0],'example:repair')
                self.assertEqual(model_call.call_args.kwargs['keep_alive'],'30m')
    def test_single_model_page_authoring_never_escalates_and_may_share_the_two_slots(self):
        from workflow_v2.pipeline import author_page
        request=dict(REQUEST,page_model=DEFAULT_MODEL)
        with tempfile.TemporaryDirectory() as t:
            out=Path(t)
            atomic(out/'plan.json',{'objectives':['目標'],'slides':[{'layout_id':'bullet_with_image','source_pages':[1]}]})
            atomic(out/'source.json',{'chunks':[]})
            with patch('workflow_v2.pipeline.page_prompt',return_value='提示'), \
                 patch('workflow_v2.pipeline.log_event') as logged, \
                 patch('workflow_v2.pipeline.call',return_value=slide()) as model_call:
                author_page(request,out,0,parallel=True)
                author_page(request,out,0,feedback='修正',parallel=True)
            self.assertEqual([c.args[0] for c in model_call.call_args_list],[DEFAULT_MODEL,DEFAULT_MODEL])
            self.assertEqual([c.kwargs['parallel'] for c in model_call.call_args_list],[True,True])
            self.assertFalse(any(c.args[0]=='model_escalation' for c in logged.call_args_list))
    def test_job_controls_pause_resume_and_remove_independently(self):
        with tempfile.TemporaryDirectory() as t:
            s=Store(t);first=s.submit(REQUEST);second=s.submit(REQUEST)
            self.assertEqual(s.pause_job(first),'paused')
            self.assertEqual(s.claim()['id'],second)
            s.resume_job(first)
            self.assertIsNone(s.claim())
            self.assertEqual(s.pause_job(second),'pausing')
            s.success(second)
            self.assertEqual(s.get(second)['status'],'paused')
            self.assertEqual(s.claim()['id'],first)
            self.assertEqual(s.pause_job(first),'pausing')
            s.success(first)
            self.assertEqual(s.get(first)['status'],'paused')
            target=s.remove_job(second)
            self.assertTrue(target.is_dir())
            self.assertEqual(s.get(first)['status'],'paused')
            with self.assertRaises(KeyError):s.get(second)
    def test_queue_finishes_current_deck_before_starting_next(self):
        with tempfile.TemporaryDirectory() as t:
            s=Store(t);first=s.submit(REQUEST);second=s.submit(REQUEST)
            self.assertEqual(s.claim()['id'],first)
            self.assertIsNone(Store(t).claim())
            for stage in STAGES[:-1]:
                s.success(first)
                self.assertEqual(s.claim()['id'],first,stage)
            s.success(first)
            self.assertEqual(s.get(first)['status'],'done')
            self.assertEqual(s.claim()['id'],second)
    def test_resumed_older_deck_does_not_interrupt_current_deck(self):
        with tempfile.TemporaryDirectory() as t:
            s=Store(t);older=s.submit(REQUEST);current=s.submit(REQUEST)
            s.pause_job(older)
            self.assertEqual(s.claim()['id'],current)
            s.resume_job(older)
            s.success(current)
            self.assertEqual(s.claim()['id'],current)
            s.pause_job(current)
            s.success(current)
            self.assertEqual(s.claim()['id'],older)
    def test_quiz_schema_rejects_misplaced_answers_after_normalization(self):
        from workflow_v2.pipeline import normalize_slide
        v=slide('quiz_qa')
        v['body']['question']='';v['body']['answer']=''
        v['body']['columns']=[{'header':'題目','points':['兩個數是多少？','答案是 3。']}]
        with self.assertRaises(ValueError):validate(slide_schema({1},'quiz_qa'),v)
        validate(slide_schema({1},'quiz_qa'),normalize_slide(v))
        self.assertEqual(v['body']['question'],'兩個數是多少？')
        self.assertEqual(v['body']['answer'],'答案是 3。')
    def test_long_source_uses_cached_local_summary(self):
        from workflow_v2.pipeline import source_context,summary_batches,digest
        with tempfile.TemporaryDirectory() as t:
            out=Path(t)
            chunks=[{'page':i,'text':'教材'*1500} for i in range(1,5)]
            src={'chunks':chunks}
            for batch in summary_batches(chunks):
                atomic(out/'source-summary'/('batch-'+digest(batch)+'.json'),
                       {'summary':'本地模型整理的摘要'})
            context=source_context(src,{1,2,3,4},out,limit=1000)
            self.assertEqual({p for b in context for p in b['pages']},{1,2,3,4})
            self.assertTrue(all(b['text']=='本地模型整理的摘要' for b in context))
    def test_quiz_columns_are_mapped_without_losing_extra_questions(self):
        from workflow_v2.pipeline import normalize_slide
        v=slide('quiz_qa')
        v['body']['question']='';v['body']['answer']=''
        v['body']['columns']=[
            {'header':'題目','points':['第一題？','第二題？']},
            {'header':'解答','points':['第一題答案','第二題答案']}]
        result=normalize_slide(v)
        validate_slide(result,{1})
        self.assertEqual(result['body']['question'],'第一題？')
        self.assertEqual(result['body']['answer'],'第一題答案')
        self.assertIn('第二題？',result['speaker_notes'])
        self.assertIn('第二題答案',result['speaker_notes'])
    def test_monitor_uses_recorded_tokens_and_work_time(self):
        from workflow_v2.monitoring import snapshot
        with tempfile.TemporaryDirectory() as t, patch('workflow_v2.model_policy.settings',return_value={}):
            s=Store(t);jid=s.submit(REQUEST);folder=s.folder(jid)
            atomic(folder/'calls/page-1-attempt-0.json',
                   {'prompt_eval_count':100,'eval_count':50,'total_duration':10_000_000_000})
            s.claim();s.event(jid,'source','開始執行')
            with patch('workflow_v2.monitoring.gpu',return_value={'utilization':20,'memory_used_mb':1000,'memory_total_mb':32000}):
                state=snapshot(s)
            self.assertEqual(state['job']['model'],DEFAULT_MODEL)
            self.assertEqual(state['job']['tokens']['total_tokens'],150)
            self.assertIsNone(state['job']['eta_seconds'])
    def test_monitor_reports_saved_pages_and_next_stage(self):
        from workflow_v2.monitoring import snapshot
        with tempfile.TemporaryDirectory() as t:
            s=Store(t);jid=s.submit(REQUEST);folder=s.folder(jid)
            atomic(folder/'plan.json',{'slides':[{} for _ in range(8)]})
            for number in range(1,4):atomic(folder/'slides'/f'{number:03}.json',{'heading':'測試'})
            with s.db() as connection:
                connection.execute("UPDATE jobs SET stage=2,status='running' WHERE id=?",(jid,))
            with patch('workflow_v2.monitoring.gpu',return_value={}):
                job=snapshot(s)['job']
            self.assertEqual((job['stage_number'],job['stage_count']),(3,7))
            self.assertEqual((job['stage_completed'],job['stage_total']),(3,8))
            self.assertEqual(job['next_stage'],'本地模型審核')
    def test_monitor_uses_review_checkpoint_not_rejected_attempts(self):
        from workflow_v2.monitoring import snapshot
        with tempfile.TemporaryDirectory() as t:
            s=Store(t);jid=s.submit(REQUEST);folder=s.folder(jid)
            atomic(folder/'review-progress.json',{'completed_pages':2,'total_pages':8})
            atomic(folder/'calls/review-3-attempt-0.json',{'prompt_eval_count':100})
            with s.db() as connection:
                connection.execute("UPDATE jobs SET stage=3,status='running' WHERE id=?",(jid,))
            with patch('workflow_v2.monitoring.gpu',return_value={}):
                job=snapshot(s)['job']
            self.assertEqual((job['stage_completed'],job['stage_total']),(2,8))
            self.assertIn('第 3 頁',job['progress'])
    def test_monitor_recovers_progress_for_review_started_before_checkpoint(self):
        from workflow_v2.monitoring import snapshot
        import time
        with tempfile.TemporaryDirectory() as t:
            s=Store(t);jid=s.submit(REQUEST);folder=s.folder(jid)
            with s.db() as connection:
                connection.execute("UPDATE jobs SET stage=3,status='running' WHERE id=?",(jid,))
            atomic(folder/'activity.json',{'stage':'review','status':'inference',
                'label':'審核第 5 頁','updated':time.time()})
            with patch('workflow_v2.monitoring.gpu',return_value={}):
                job=snapshot(s)['job']
            self.assertEqual(job['stage_completed'],4)
            self.assertEqual(job['progress'],'審核第 5 頁')
            atomic(folder/'activity.json',{'stage':'write','status':'inference',
                'label':'編寫第 5 頁','updated':time.time()})
            with patch('workflow_v2.monitoring.gpu',return_value={}):
                job=snapshot(s)['job']
            self.assertEqual(job['stage_completed'],4)
            self.assertEqual(job['progress'],'修復第 5 頁')
    def test_summary_batches_preserve_source(self):
        from workflow_v2.pipeline import summary_batches
        chunks=[dict(page=i,text='字'*2000) for i in range(1,10)]
        batches=list(summary_batches(chunks))
        self.assertEqual(len(batches),3)
        self.assertEqual([c for b in batches for c in b],chunks)
    def test_plan_uses_drafting_model_for_long_source_summaries(self):
        from workflow_v2.pipeline import plan
        with tempfile.TemporaryDirectory() as t:
            out=Path(t)
            atomic(out/'source.json',{'pages':[{'page':1,'text':'教材'*4000}],
                'chunks':[{'page':1,'text':'教材'*2000},{'page':1,'text':'例題'*2000}]})
            calls=[]
            def fake_call(model,prompt,schema,budget,record,**kwargs):
                calls.append((model,budget,prompt))
                return {'summary':'頁碼 1：核心概念及例題'} if record.parent.name=='source-summary' else {'slides':[]}
            with patch('workflow_v2.pipeline.call',side_effect=fake_call), \
                 patch('workflow_v2.pipeline.log_event') as logged, \
                 patch('workflow_v2.pipeline.outline_ready'):
                plan(dict(REQUEST,page_model='example:draft'),out)

            self.assertEqual([model for model,_,_ in calls],
                             ['example:draft','example:draft','example:draft',DEFAULT_MODEL])
            self.assertTrue(all(budget==1200 for _,budget,_ in calls[:2]))
            self.assertTrue((out/'analysis.json').exists())
            self.assertIn('必要例題',calls[0][2])
            self.assertLess(logged.call_args.kwargs['prompt_chars'],30000)
    def test_unpublished_artifacts_are_hidden(self):
        from workflow_v2.service import result
        with tempfile.TemporaryDirectory() as t:
            s=Store(t);jid=s.submit(REQUEST)
            atomic(s.folder(jid)/'result.json',{'artifacts':{'pptx':'draft.pptx'}})
            self.assertEqual(result(jid,s)['artifacts'],{})
    def test_layout_image_never_overlays_quiz_answer(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as t:
            d=dict(schema_version=1,deck_meta=dict(title='測試',subject='數學',target_audience='六年級',primary_color='#087F78'),slides=[slide('quiz_qa') for _ in range(8)])
            path=Path(t)/'test.pptx'
            render(d,path,media={1:'nonexistent-image.png'})
            self.assertIn('完整的答案',''.join(s.text for s in Presentation(path).slides[0].shapes if s.has_text_frame))
    def test_quiz_geometry_gives_longer_question_more_room_without_clipping_answer(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as t:
            short=slide('quiz_qa');long=slide('quiz_qa')
            long['body']['question']='請仔細讀題並找出題目中已知的數量與條件。'*5
            deck=dict(schema_version=1,deck_meta=dict(title='測試',subject='數學',
                target_audience='六年級',primary_color='#087F78'),slides=[short,long]+[slide() for _ in range(6)])
            path=Path(t)/'quiz.pptx';render(deck,path)
            slides=Presentation(path).slides
            for page in (slides[0],slides[1]):
                question,answer=page.placeholders[1],page.placeholders[2]
                self.assertLess(question.left+question.width,answer.left)
                self.assertIn('完整的答案',answer.text)
    def test_pagination_keeps_every_bullet_and_renders_continuations(self):
        from workflow_v2.renderer import paginate_deck
        with tempfile.TemporaryDirectory() as t:
            original=slide();original['body']['bullets']=['重點'+str(i) for i in range(6)]+['長段落'*100]
            deck=dict(schema_version=1,deck_meta=dict(title='測試',subject='數學',
                target_audience='六年級',primary_color='#087F78'),slides=[original]+[slide() for _ in range(7)])
            paged=paginate_deck(deck)
            self.assertGreater(len(paged['slides']),8)
            rendered=[p for p in paged['slides'] if p['heading']==original['heading']][:len(paged['slides'])-7]
            self.assertEqual(''.join(x for p in rendered for x in p['body']['bullets']),
                             ''.join(original['body']['bullets']))
            path=Path(t)/'paged.pptx';render(paged,path)
            from pptx import Presentation
            self.assertEqual(len(Presentation(path).slides),len(paged['slides']))
    def test_overflow_split_preserves_quiz_question_and_answer(self):
        from workflow_v2.renderer import split_overflow_page,slide_text
        original=slide('quiz_qa')
        original['body']['question']='題目條件。'*90
        original['body']['answer']='完整解法。'*120
        deck=dict(schema_version=1,deck_meta=dict(title='測試',subject='數學',
            target_audience='六年級',primary_color='#087F78'),slides=[original]+[slide() for _ in range(7)])
        main,secondary=slide_text(original)
        paged=split_overflow_page(deck,0)
        self.assertEqual(''.join(p['body']['bullets'][0] for p in paged['slides'][0:-7]),main+secondary)
        self.assertTrue(all(p['layout_id']=='bullet_with_image' for p in paged['slides'][0:-7]))
        with tempfile.TemporaryDirectory() as t:
            render(paged,Path(t)/'split.pptx')
    def test_fallback_layout_preserves_an_unsplittable_page_and_renders(self):
        from workflow_v2.renderer import safe_fallback_page,slide_text
        from pptx import Presentation
        original=slide('quiz_qa')
        original['body']['question']='請找出正確的數字。'*140
        original['body']['answer']='解答與理由完整保留。'*140
        deck=dict(schema_version=1,deck_meta=dict(title='測試',subject='數學',
            target_audience='六年級',primary_color='#087F78'),slides=[original]+[slide() for _ in range(7)])
        main,answer=slide_text(original)
        fallback=safe_fallback_page(deck,0)
        pages=fallback['slides'][:-7]
        self.assertGreater(len(pages),1)
        self.assertEqual(''.join(p['body']['bullets'][0] for p in pages),main+answer)
        self.assertTrue(all(p['layout_id']=='bullet_with_image' for p in pages))
        with tempfile.TemporaryDirectory() as t:
            path=Path(t)/'fallback.pptx';render(fallback,path)
            self.assertEqual(len(Presentation(path).slides),len(fallback['slides']))
    def test_render_stage_downgrades_only_problem_page(self):
        from workflow_v2 import pipeline
        original=slide('quiz_qa')
        original['body']['question']='長題目。'*100
        original['body']['answer']='完整解答。'*100
        deck=dict(schema_version=1,deck_meta=dict(title='測試',subject='數學',
            target_audience='六年級',primary_color='#087F78'),slides=[original]+[slide() for _ in range(7)])
        real_render=render
        calls=[]
        def render_after_first_overflow(value,target,**kwargs):
            calls.append(value)
            if len(calls)==1:raise Overflow('第1頁「版型測試」：文字過長')
            return real_render(value,target,**kwargs)
        with tempfile.TemporaryDirectory() as t, \
             patch.object(pipeline,'deck',return_value=deck), \
             patch.object(pipeline,'split_overflow_page',side_effect=Overflow('此頁內容無法再拆分')), \
             patch.object(pipeline,'generate_assets',return_value={}), \
             patch.object(pipeline,'render',side_effect=render_after_first_overflow):
            out=Path(t);pipeline.rendering(REQUEST,out)
            warnings=json.loads((out/'layout-warnings.json').read_text(encoding='utf-8'))
            self.assertEqual(warnings[0]['original_page'],1)
            self.assertTrue((out/'教材.pptx').exists())
            saved=json.loads((out/'rendered-deck.json').read_text(encoding='utf-8'))
            self.assertEqual(len(saved['slides']),8)
            self.assertIn(original['body']['answer'],saved['slides'][0]['speaker_notes'])
            self.assertFalse(warnings[0]['complete_visible'])
    def test_worker_interrupts_active_stage_on_pause_without_advancing(self):
        import threading,time
        from workflow_v2 import worker
        class FakeProcess:
            pid=12345
            stopped=False
            returncode=None
            def poll(self):return 75 if self.stopped else None
        process=FakeProcess()
        with tempfile.TemporaryDirectory() as t:
            store=Store(t);jid=store.submit(REQUEST)
            def stop(proc):proc.stopped=True;proc.returncode=75
            with patch.object(worker.subprocess,'Popen',return_value=process), \
                 patch.object(worker,'stop_tree',side_effect=stop):
                thread=threading.Thread(target=worker.step,args=(store,));thread.start()
                time.sleep(.08);store.pause_job(jid);thread.join(timeout=3)
            self.assertFalse(thread.is_alive())
            self.assertEqual(store.get(jid)['status'],'paused')
            self.assertEqual(store.get(jid)['stage'],0)
            self.assertEqual(store.get(jid)['attempts'],0)
    def test_worker_interrupts_active_stage_on_remove(self):
        import threading,time
        from workflow_v2 import worker
        class FakeProcess:
            pid=12345
            stopped=False
            returncode=None
            def poll(self):return -9 if self.stopped else None
        process=FakeProcess()
        with tempfile.TemporaryDirectory() as t:
            store=Store(t);jid=store.submit(REQUEST)
            def stop(proc):proc.stopped=True;proc.returncode=-9
            with patch.object(worker.subprocess,'Popen',return_value=process), \
                 patch.object(worker,'stop_tree',side_effect=stop):
                thread=threading.Thread(target=worker.step,args=(store,));thread.start()
                time.sleep(.08);store.remove_job(jid);thread.join(timeout=3)
            self.assertFalse(thread.is_alive())
            with self.assertRaises(KeyError):store.get(jid)
            self.assertTrue(list((Path(t)/'trash').glob(jid+'-*')))
    def test_resume_during_interruption_requeues_same_stage(self):
        with tempfile.TemporaryDirectory() as t:
            store=Store(t);jid=store.submit(REQUEST)
            store.claim();store.pause_job(jid);store.resume_job(jid)
            self.assertEqual(store.complete_interruption(jid),'pending')
            self.assertEqual(store.get(jid)['stage'],0)
            self.assertEqual(store.claim()['id'],jid)
    def test_local_media_matches_only_valid_local_image(self):
        from PIL import Image
        from workflow_v2.renderer import local_media
        with tempfile.TemporaryDirectory() as t:
            folder=Path(t)
            Image.new('RGB',(32,32),'blue').save(folder/'fraction-chart.png')
            (folder/'fraction-broken.png').write_text('invalid',encoding='utf-8')
            concept=slide();concept['media_hint']['search_keyword']='fraction chart'
            deck={'slides':[concept,slide()]}
            self.assertEqual(local_media(deck,folder),{1:str(folder/'fraction-chart.png')})
    def test_online_image_accepts_only_cc0_and_keeps_provenance(self):
        import io
        from PIL import Image
        from workflow_v2.images import cc0_candidate,online_images
        page={'title':'File:Fractions.png','imageinfo':[{'mime':'image/png',
            'thumburl':'https://upload.wikimedia.org/example.png',
            'descriptionurl':'https://commons.wikimedia.org/wiki/File:Fractions.png',
            'extmetadata':{'LicenseShortName':{'value':'CC0 1.0'},
                           'LicenseUrl':{'value':'https://creativecommons.org/publicdomain/zero/1.0/'}}}]}
        self.assertIsNotNone(cc0_candidate(page))
        restricted=json.loads(json.dumps(page))
        restricted['imageinfo'][0]['extmetadata']['LicenseShortName']['value']='CC BY-SA 4.0'
        self.assertIsNone(cc0_candidate(restricted))
        buffer=io.BytesIO();Image.new('RGB',(24,24),'green').save(buffer,format='PNG')
        class Response:
            status_code=200
            def __init__(self,data):self.data=data
            def raise_for_status(self):pass
            def json(self):return {'query':{'pages':[restricted,page]}}
            def iter_content(self,n):yield self.data
        class Session:
            def get(self,url,**kwargs):
                return Response(buffer.getvalue() if url.startswith('https://upload.') else b'')
        concept=slide();concept['media_hint']['search_keyword']='fractions'
        with tempfile.TemporaryDirectory() as t:
            media=online_images({'slides':[concept]},Path(t),session=Session())
            self.assertIn(1,media)
            saved=json.loads((Path(t)/'image-sources.json').read_text(encoding='utf-8'))
            self.assertEqual(saved['1']['license'],'CC0 1.0')
            self.assertTrue(Path(media[1]).is_file())
    def test_retry_preserves_stage_and_other_jobs(self):
        with tempfile.TemporaryDirectory() as t:
            s=Store(t);a=s.submit(REQUEST);b=s.submit(REQUEST)
            for _ in range(3):s.fail(a,'offline')
            self.assertEqual(s.get(a)['status'],'failed');self.assertEqual(s.claim()['id'],b)
            s.retry(a);self.assertEqual(s.get(a)['stage'],0)
            for _ in STAGES:s.success(a)
            self.assertEqual(s.get(a)['status'],'done')
    def test_delete_done_moves_artifacts_to_trash_and_rejects_active(self):
        with tempfile.TemporaryDirectory() as t:
            s=Store(t);jid=s.submit(REQUEST);out=s.folder(jid);atomic(out/'教材.pptx',{'x':1})
            with s.db() as c:c.execute("UPDATE jobs SET status='running' WHERE id=?",(jid,))
            with self.assertRaises(ValueError):s.delete_done(jid)
            with s.db() as c:c.execute("UPDATE jobs SET status='done' WHERE id=?",(jid,))
            target=s.delete_done(jid)
            self.assertTrue((target/'教材.pptx').exists())
            with self.assertRaises(KeyError):s.get(jid)
    def test_session_usage_counts_only_calls_after_session_marker(self):
        from workflow_v2.monitoring import session_usage
        import os,time
        with tempfile.TemporaryDirectory() as t:
            s=Store(t);jid=s.submit(REQUEST);out=s.folder(jid)
            old=out/'calls/page-1-attempt-0.json';atomic(old,{'prompt_eval_count':9,'eval_count':1})
            now=time.time();os.utime(old,(now-100,now-100))
            fresh=out/'calls/page-2-attempt-0.json';atomic(fresh,{'prompt_eval_count':20,'eval_count':5})
            with self.assertRaises(KeyError):s.get('missing')

            usage=session_usage(s,now-10)
            self.assertEqual(usage['total_tokens'],25)
    def test_outline_approval_validates_then_resumes_at_write(self):
        from workflow_v2.service import approve_outline
        with tempfile.TemporaryDirectory() as t:
            s=Store(t);jid=s.submit(REQUEST);out=s.folder(jid)
            atomic(out/'source.json',{'pages':[{'page':1,'text':'來源'}]})
            phases=['hook','concept','example','recall','recap','concept','example','recall']
            layouts=['title_hero','bullet_with_image','code_walkthrough','quiz_qa',
                     'summary_checklist','bullet_with_image','code_walkthrough','quiz_qa']
            plan={'objectives':['能說出重點'],
                  'slides':[{'topic':str(i),'phase':phase,'layout_id':layouts[i],
                      'difficulty':'共用','source_pages':[1]} for i,phase in enumerate(phases)]}
            with s.db() as c:c.execute("UPDATE jobs SET status='awaiting_outline',stage=2 WHERE id=?",(jid,))
            bad=json.loads(json.dumps(plan));bad['slides'][0]['source_pages']=[99]
            with self.assertRaises(ValueError):approve_outline(jid,bad,s)
            self.assertEqual(s.get(jid)['status'],'awaiting_outline')
            approve_outline(jid,plan,s)
            self.assertEqual(s.get(jid)['status'],'pending')
            self.assertEqual(s.get(jid)['stage'],2)
    def test_outline_citations_are_bounded_and_unique(self):
        from workflow_v2.contracts import outline_schema,validate
        phases=['hook','concept','example','recall','recap','concept','example','recall']
        slides=[{'layout_id':'quiz_qa' if phase=='recall' else 'bullet_with_image',
                 'topic':'重點','phase':phase,'difficulty':'共用','source_pages':[1]}
                for phase in phases]
        value={'objectives':['理解重點'],'slides':slides}
        validate(outline_schema(8,{1,2,3,4}),value)
        value['slides'][0]['source_pages']=[1,2,3,4]
        with self.assertRaises(ValueError):validate(outline_schema(8,{1,2,3,4}),value)
        value['slides'][0]['source_pages']=[1,1]
        with self.assertRaises(ValueError):validate(outline_schema(8,{1,2,3,4}),value)
    def test_outline_repairs_duplicate_citations_and_missing_recall_before_validation(self):
        from workflow_v2.pipeline import normalize_outline
        from workflow_v2.pedagogy import validate_teaching_cycle
        phases=['hook','concept','example','concept','example','concept','example','recap']
        value={'objectives':['理解重點'],'slides':[
            {'layout_id':'Text_Visual','topic':'重點','phase':phase,
             'difficulty':'低難度','source_pages':[1,1]}
            for phase in phases]}
        normalize_outline(value)
        validate_teaching_cycle(value)
        self.assertEqual(value['slides'][3]['source_pages'],[1])
        self.assertEqual(value['slides'][-2]['layout_id'],'Question')
    def test_parallel_writing_waits_for_earlier_pages_before_recap(self):
        import time
        from workflow_v2.pipeline import write
        with tempfile.TemporaryDirectory() as t:
            out=Path(t);atomic(out/'plan.json',{'slides':[
                {'phase':'concept'},{'phase':'example'},{'phase':'recap'}]})
            completed=[];started=[]
            def fake_write(request,folder,index,item,parallel=False):
                started.append((index,parallel))
                if index<2:time.sleep(.05)
                if index==2:self.assertEqual(set(completed),{0,1})
                completed.append(index)
            with patch.dict('os.environ',{'MAX_PPT_CONCURRENCY':'2'}), \
                 patch('workflow_v2.pipeline.unload_model'), \
                 patch('workflow_v2.pipeline.write_one',side_effect=fake_write):
                write({'model':'heavy','page_model':'fast'},out)
            self.assertEqual({index for index,_ in started},{0,1,2})
            self.assertFalse(dict(started)[2])
    def test_parallel_draft_finishes_before_large_model_repair(self):
        from workflow_v2.pipeline import write
        with tempfile.TemporaryDirectory() as t:
            out=Path(t)
            atomic(out/'source.json',{'pages':[]})
            atomic(out/'plan.json',{'slides':[{'phase':'concept'},{'phase':'example'}]})
            order=[]
            def fake_author(request,folder,index,feedback='',parallel=False,final=False):
                if index==0 and not feedback:
                    order.append('fast-failed');raise ValueError('需修正')
                order.append('repair' if feedback else 'fast-done')
            with patch.dict('os.environ',{'MAX_PPT_CONCURRENCY':'2'},clear=True), \
                 patch('workflow_v2.pipeline.unload_model'), \
                 patch('workflow_v2.pipeline.recover_page',return_value=False), \
                 patch('workflow_v2.pipeline.author_page',side_effect=fake_author):
                write(dict(REQUEST,model='example:repair',page_model='example:draft'),out)
            self.assertEqual(order[-1],'repair')
            self.assertIn('fast-done',order[:-1])
    def test_serial_drafts_finish_before_loading_rescue_model(self):
        from workflow_v2.pipeline import write
        with tempfile.TemporaryDirectory() as t:
            out=Path(t)
            atomic(out/'source.json',{'pages':[]})
            atomic(out/'plan.json',{'slides':[{'phase':'concept'},{'phase':'example'}]})
            order=[]
            def fake_author(request,folder,index,feedback='',parallel=False,final=False):
                if index==0 and not feedback:
                    order.append('fast-failed');raise ValueError('需修正')
                order.append('repair' if feedback else 'fast-done')
            with patch.dict('os.environ',{'MAX_PPT_CONCURRENCY':'1'}), \
                 patch('workflow_v2.pipeline.unload_model'), \
                 patch('workflow_v2.pipeline.recover_page',return_value=False), \
                 patch('workflow_v2.pipeline.author_page',side_effect=fake_author):
                write(dict(REQUEST,model='example:repair',page_model='example:draft'),out)
            self.assertEqual(order,['fast-failed','fast-done','repair'])
    def test_model_context_avoids_runner_switches_for_adjacent_pages(self):
        from workflow_v2.model import context_size
        self.assertEqual(context_size('example:local','x'*4000),8192)
        self.assertEqual(context_size('example:local','x'*9000),16384)
        with patch('workflow_v2.model_policy.settings',return_value={'generation_context':32768}):
            self.assertEqual(context_size(DEFAULT_MODEL,'x'*7000,
                                          Path('review-4.json')),32768)
    def test_visual_only_conflict_is_normalized_without_rewriting_teaching_text(self):
        from workflow_v2.pipeline import normalize_slide
        from workflow_v2.contracts import VISUAL_STYLE
        value=slide('Question')
        value['visual_asset']={'is_required':True,'reasoning':'圖片輔助',
            'asset_type':'illustration','image_prompt':'A friendly animal',
            'style_tags':VISUAL_STYLE}
        original=value['body'].copy()
        normalize_slide(value)
        self.assertEqual(value['body'],original)
        self.assertFalse(value['visual_asset']['is_required'])
        self.assertEqual(value['visual_asset']['image_prompt'],'')
    def test_visual_prompt_is_cleaned_without_rejecting_valid_slide(self):
        from workflow_v2.pipeline import normalize_slide
        from workflow_v2.contracts import VISUAL_STYLE
        value=slide('Text_Visual');value['body']['bullets']=['算式：12 + 3 = 15']
        value['visual_asset']={'is_required':True,'reasoning':'概念圖',
            'asset_type':'illustration','image_prompt':'A friendly cat 123 教材, beside a tree',
            'style_tags':VISUAL_STYLE}
        normalized=normalize_slide(value)
        self.assertEqual(normalized['body']['bullets'],['算式：12 + 3 = 15'])
        self.assertEqual(normalized['visual_asset']['image_prompt'],'A friendly cat , beside a tree')
    def test_single_model_write_and_review_never_unload_or_defer(self):
        from workflow_v2.pipeline import write,review
        request=dict(REQUEST,page_model=DEFAULT_MODEL)
        with tempfile.TemporaryDirectory() as t:
            out=Path(t)
            atomic(out/'source.json',{'pages':[],'sha256':'source-hash','chunks':[]})
            atomic(out/'plan.json',{'slides':[{'phase':'concept','source_pages':[1],'layout_id':'Question'},
                                              {'phase':'example','source_pages':[1],'layout_id':'Question'}]})
            atomic(out/'slides/001.json',slide('Question'))
            calls=[]
            def fake_author(request,folder,index,feedback='',parallel=False,final=False):
                calls.append((index,feedback,parallel))
            with patch.dict('os.environ',{'MAX_PPT_CONCURRENCY':'1'},clear=True), \
                 patch('workflow_v2.pipeline.unload_model') as unload, \
                 patch('workflow_v2.pipeline.recover_page',return_value=False), \
                 patch('workflow_v2.pipeline.author_page',side_effect=fake_author):
                write(request,out)
            self.assertEqual(calls,[(0,'',False),(1,'',False)])
            unload.assert_not_called()
            plan=json.loads((out/'plan.json').read_text(encoding='utf-8'))
            with patch('workflow_v2.pipeline.unload_model') as unload, \
                 patch('workflow_v2.pipeline.source_context',return_value=[]), \
                 patch('workflow_v2.pipeline.call',return_value={'passed':True,'issues':[]}), \
                 patch('workflow_v2.pipeline.deck',return_value={'slides':[]}):
                atomic(out/'plan.json',{'slides':plan['slides'][:1]})
                atomic(out/'slides/001.json',slide('Question'))
                review(request,out)
            unload.assert_not_called()
    def test_review_allows_only_one_model_rewrite_before_source_fallback(self):
        from workflow_v2.pipeline import review
        request=dict(REQUEST,model='example:repair',page_model='example:draft')
        with tempfile.TemporaryDirectory() as t:
            out=Path(t)
            atomic(out/'plan.json',{'slides':[{'source_pages':[1],'layout_id':'Question'}]})
            atomic(out/'source.json',{'sha256':'source-hash','chunks':[]})
            atomic(out/'slides/001.json',slide('Question'))
            revised=slide('Question');revised['heading']='修正後的標題'
            qa={'passed':False,'issues':['答案計算錯誤']}
            with patch('workflow_v2.pipeline.unload_model'), \
                 patch('workflow_v2.pipeline.source_context',return_value=[]), \
                 patch('workflow_v2.pipeline.call',return_value=qa) as model_call, \
                 patch('workflow_v2.pipeline.author_page',return_value=revised) as rewrite, \
                 patch('workflow_v2.pipeline.fallback_page') as fallback, \
                 patch('workflow_v2.pipeline.deck',return_value={'slides':[]}):
                review(request,out)
            self.assertEqual(model_call.call_count,2)
            rewrite.assert_called_once()
            fallback.assert_called_once()
    def test_gpu_slots_allow_two_bounded_page_calls(self):
        import threading,time
        from workflow_v2.model import gpu_slot
        barrier=threading.Barrier(2)
        def occupy():
            with gpu_slot(parallel=True):
                barrier.wait(timeout=2)
                time.sleep(.02)
        with tempfile.TemporaryDirectory() as t,patch.dict('os.environ',{'MAX_PPT_CONCURRENCY':'2'}),patch('workflow_v2.model.DATA',Path(t)):
            a=threading.Thread(target=occupy);b=threading.Thread(target=occupy)
            a.start();b.start();a.join(timeout=3);b.join(timeout=3)
        self.assertFalse(a.is_alive() or b.is_alive())
    def test_exclusive_gpu_call_waits_for_second_parallel_slot(self):
        import threading,time
        from filelock import FileLock
        from workflow_v2.model import gpu_slot
        entered=threading.Event()
        def gpu_slot_task():
            with gpu_slot(parallel=False):entered.set()
        with tempfile.TemporaryDirectory() as t,patch('workflow_v2.model.DATA',Path(t)):
            with FileLock(str(Path(t)/'gpu-slot-2.lock')):
                thread=threading.Thread(target=gpu_slot_task)
                thread.start()
                time.sleep(.1)
                self.assertFalse(entered.is_set())
            thread.join(timeout=3)
        self.assertTrue(entered.is_set())
        self.assertFalse(thread.is_alive())
    def test_does_not_accept_incomplete_arithmetic(self):
        v=slide();v['body']['bullets']=['48 =']
        with self.assertRaises(ValueError):validate_slide(v,{1})
    def test_all_layouts_and_no_silent_truncation(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as t:
            d=dict(schema_version=1,deck_meta=dict(title='測試',subject='數學',target_audience='六年級',primary_color='#087F78'),slides=[slide(x) for x in LEGACY_LAYOUTS]+[slide(),slide()])
            path=Path(t)/'test.pptx';render(d,path)
            p=Presentation(path);self.assertEqual(len(p.slides),len(d['slides']))
            self.assertIn('講解30秒',p.slides[0].notes_slide.notes_text_frame.text)
            self.assertIn('完整的一句話',''.join(s.text for s in p.slides[2].shapes if s.has_text_frame))
            d['slides'][0]['body']['bullets']=['很長'*2000]
            d['slides'][0]['subheading']=''
            with self.assertRaises(Overflow):render(d,path)
    def test_regenerate_forks_only_selected_page(self):
        with tempfile.TemporaryDirectory() as t:
            s=Store(t);a=s.submit(REQUEST);out=s.folder(a)
            atomic(out/'source.json',{'pages':[]});atomic(out/'plan.json',{'slides':[{},{}]})
            atomic(out/'slides/001.json',slide());atomic(out/'slides/002.json',slide())
            with s.db() as c:c.execute("update jobs set status='done' where id=?",(a,))
            b=regenerate(a,1,'改寫',s)
            self.assertTrue((s.folder(b)/'slides/001.json').exists())
            self.assertFalse((s.folder(b)/'slides/002.json').exists())
            self.assertTrue((out/'slides/002.json').exists())
            self.assertEqual(s.get(b)['stage'],2)
            self.assertEqual(s.get(a)['status'],'done')
    def test_placeholder_regeneration_keeps_other_fields(self):
        from workflow_v2.pipeline import apply_regenerate_scope
        with tempfile.TemporaryDirectory() as t:
            s=Store(t);parent=s.submit(REQUEST);child=s.submit(REQUEST)
            original=slide();atomic(s.folder(parent)/'slides/001.json',original)
            atomic(s.folder(child)/'regenerate.json',{'page':0,'field':'heading',
                'parent_job':parent,'instruction':'改標題'})
            revised=slide();revised['heading']='新標題';revised['body']['bullets']=['不應改變']
            result=apply_regenerate_scope(s.folder(child),0,revised)
            self.assertEqual(result['heading'],'新標題')
            self.assertEqual(result['body'],original['body'])

if __name__=='__main__':unittest.main()
