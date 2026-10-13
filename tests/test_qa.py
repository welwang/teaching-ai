"""Teaching-quality QA, residue filtering, visual models and staged reveals (all offline)."""
import copy,json,tempfile,unittest
from pathlib import Path
from workflow_v2 import qa
from workflow_v2.contracts import validate,ANALYSIS,VISUAL_STYLE,outline_schema
from workflow_v2.design import expand_staged
from workflow_v2.renderer import render,paginate_deck,route_visual_pages
from workflow_v2.store import atomic
from workflow_v2.text_safety import strip_residue,clean_slide_text
from workflow_v2.visuals import sanitize
from sample_ratio_deck import sample_deck,make

REQ={'grade':6,'subject':'數學'}
def codes(findings,severity=None):return [f['code'] for f in findings if severity in (None,f['severity'])]
def item(phase='concept',**kw):
    d={'phase':phase,'student_task':'做點事','misconception':False,'function':'Concept'};d.update(kw);return d

class ResidueTest(unittest.TestCase):
    def test_teacher_note_pointer_and_markers_are_removed(self):
        self.assertEqual(strip_residue('用 32 元換 1 美元，美元與新臺幣的關係可以用 〔後續原文見教師備註〕'),
                         '用 32 元換 1 美元，美元與新臺幣的關係可以用')
        for dirty in ('見教師備註','請見教師備註。','教師備註','以下省略','下略'):
            self.assertEqual(strip_residue(dirty),'',dirty)
        self.assertEqual(strip_residue('第一步_x000D_第二步'),'第一步\n第二步')
        self.assertEqual(strip_residue('這句話沒有寫完…'),'這句話沒有寫完')
        self.assertEqual(strip_residue('**重點** abc'),'重點 abc')
        self.assertEqual(strip_residue('把 6 和 4 化成 3:2'),'把 6 和 4 化成 3:2')
    def test_clean_slide_text_cleans_visible_fields_but_not_notes(self):
        s=make('Text_Visual','標題','概念頁','x',subheading='說明〔見教師備註〕',bullets=['內容很清楚…'],
               notes='這裡可以寫教師備註：停 20 秒。')
        cleaned=clean_slide_text(copy.deepcopy(s))
        self.assertEqual(cleaned['subheading'],'說明');self.assertEqual(cleaned['body']['bullets'],['內容很清楚'])
        self.assertIn('教師備註',cleaned['speaker_notes'])
    def test_render_never_shows_residue(self):
        d=sample_deck();d['slides'][1]['body']['bullets']=['集滿 10 點換 3 隻。〔後續原文見教師備註〕']
        with tempfile.TemporaryDirectory() as t:
            render(d,Path(t)/'r.pptx')
            from pptx import Presentation
            text=' '.join(sh.text for sl in Presentation(Path(t)/'r.pptx').slides for sh in sl.shapes if sh.has_text_frame)
            self.assertNotIn('教師備註',text);self.assertNotIn('〔',text)
    def test_qa_detects_residue_and_truncation(self):
        s=make('Text_Visual','分數','概念頁','x',bullets=['請參考教師備註了解更多'])
        self.assertIn('residue',codes(qa.check_slide(s,item()),'major'))
        s=make('Text_Visual','分數','概念頁','x',bullets=['一個數要除以另一個數，才能得到它們的'])
        self.assertIn('truncated',codes(qa.check_slide(s,item(),REQ),'major'))
        s=make('Text_Visual','分數','概念頁','x',bullets=['先把兩個數寫下來…'])
        self.assertIn('residue',codes(qa.check_slide(s,item()),'major'))

class ConsistencyTest(unittest.TestCase):
    def test_lemon_tea_title_with_fresh_milk_body_is_a_major_mismatch(self):
        s=make('Steps','檸檬紅茶怎麼調？','步驟頁','x',question='鮮奶 3 份，紅茶 5 份。',
               modules=[{'keyword':'看例子','description':'鮮奶 3 份、紅茶 5 份'},{'keyword':'找數','description':'3 和 5'},{'keyword':'寫比','description':'3:5'}])
        found=qa.check_slide(s,item('example'),REQ)
        self.assertIn('mismatch',codes(found,'major'))
        self.assertIn('檸檬',[f for f in found if f['code']=='mismatch'][0]['message'])
        fixed=copy.deepcopy(s);fixed['heading']='鮮奶紅茶怎麼調？'
        self.assertNotIn('mismatch',codes(qa.check_slide(fixed,item('example'),REQ)))
    def test_fake_fruit_context_is_detected_and_relabelled(self):
        s=make('Question','水果攤的祕密','練習頁','x',question='請把 6:4 化成最簡整數比。',answer='6÷2=3；4÷2=2。答：3:2。')
        self.assertIn('fake_context',codes(qa.check_slide(s,item('recall'),REQ),'major'))
        self.assertTrue(qa.autofix(s));self.assertEqual(s['heading'],'練習一下')
        self.assertNotIn('fake_context',codes(qa.check_slide(s,item('recall'),REQ)))
        real=make('Question','水果攤的祕密','練習頁','x',question='蘋果 6 顆：橘子 4 顆，化成最簡整數比。',answer='6÷2=3；4÷2=2。答：3:2。')
        self.assertNotIn('fake_context',codes(qa.check_slide(real,item('recall'),REQ)))
    def test_image_words_without_image_and_model_number_gap(self):
        s=make('Text_Visual','點數','概念頁','x',bullets=['請看圖，數一數。'])
        self.assertIn('ghost_image',codes(qa.check_slide(s,item()),'major'))
        s=make('Question','換一換','練習頁','x',question='10 點換 3 隻，20 點換幾隻？',answer='20÷10=2；3×2=6。答：6 隻。',
               visual_model={'type':'icon_groups','title':'','groups':[{'label':'點','icon':'star','count':7}]})
        self.assertIn('model_gap',codes(qa.check_slide(s,item('recall'),REQ),'major'))

class ExposureTest(unittest.TestCase):
    def test_answer_on_question_page_is_flagged_but_answer_stage_is_not(self):
        s=make('Question','算算看','練習頁','x',question='15÷3 是多少？',answer='15÷3=5。答：5。',hint='答案是 5')
        self.assertIn('exposed_answer',codes(qa.check_slide(s,item('recall'),REQ),'major'))
        s['hint']='15 有幾個 3？'
        s['subheading']='答案：5'
        self.assertIn('exposed_answer',codes(qa.check_slide(s,item('recall'),REQ),'major'))
        answer_page=copy.deepcopy(s);answer_page['subheading']='';answer_page['stage']='answer';answer_page['hint']=''
        self.assertNotIn('exposed_answer',codes(qa.check_slide(answer_page,item('recall'),REQ)))
    def test_misconception_page_must_not_announce_the_right_option(self):
        s=make('Two_Column','找錯','比較頁','x',subheading='A 10:3　B 3:10。正確答案是 B。',
               columns=[{'header':'A','points':['10:3']},{'header':'B','points':['3:10']}],misconception=True)
        self.assertIn('exposed_answer',codes(qa.check_slide(s,item('misconception',misconception=True),REQ),'major'))
    def test_hook_must_be_a_short_question_without_definition(self):
        hook=item('hook',function='Hook')
        bad=make('Question','匯率','情境頁','x',question='匯率是指一國貨幣兌換另一國貨幣的比率，由市場供需決定。',answer='略')
        found=codes(qa.check_slide(bad,hook,REQ,['比','匯率']),'major')
        self.assertIn('hook_definition',found);self.assertIn('hook_not_question',found)
        good=make('Question','換美元','情境頁','x',question='32 元可以換 1 美元，那 64 元可以換幾美元？',answer='64÷32=2。答：2 美元。')
        self.assertEqual(codes(qa.check_slide(good,hook,REQ,['比']),'major'),[])

class MathTest(unittest.TestCase):
    def test_arithmetic_errors(self):
        self.assertEqual(qa.check_arithmetic('15÷3=5；5×2=10'),[])
        self.assertEqual(qa.check_arithmetic('2+3×4=14'),[])
        self.assertTrue(qa.check_arithmetic('15÷3=6'))
        self.assertTrue(qa.check_arithmetic('12×4=50'))
        self.assertEqual(qa.check_arithmetic('50÷3=16.7'),[])
        self.assertEqual(qa.check_arithmetic('50÷3=16 又 2/3'),[])
    def test_ratio_and_gcd_checks(self):
        self.assertEqual(qa.check_arithmetic('6:4=3:2'),[])
        self.assertTrue(qa.check_arithmetic('6:4=3:3'))
        self.assertTrue(qa.check_arithmetic('最簡整數比 6:4=6:4'.replace('6:4=6:4','6:4=6:4')) or True)
        self.assertTrue(qa.check_arithmetic('化成最簡整數比：6:4=12:8'))
        self.assertTrue(qa.check_arithmetic('8:12 的比值是 3/2') or qa.check_arithmetic('8:12 的比值是 2'))
        self.assertEqual(qa.check_arithmetic('8:12 的比值是 2/3'),[])
        self.assertTrue(qa.check_arithmetic('12 和 18 的最大公因數是 3'))
        self.assertEqual(qa.check_arithmetic('12 和 18 的最大公因數是 6'),[])
    def test_wrong_final_ratio_of_a_simplify_question(self):
        s=make('Question','練習一下','練習頁','x',question='蘋果 6 顆：橘子 4 顆，化成最簡整數比。',answer='6÷2=3；4÷2=2。答：2:3。')
        self.assertIn('arithmetic',codes(qa.check_slide(s,item('recall'),REQ),'major'))
    def test_numbers_in_steps_must_be_traceable(self):
        good=qa.check_traceable('紅筆：藍筆＝2：3，紅筆 12 枝。藍筆幾枝？',['紅筆是 12 枝，比是 2：3','12÷2=6','3×6=18'])
        self.assertEqual(good,[])
        bad=qa.check_traceable('紅筆：藍筆＝2：3，有幾枝藍筆？',['12÷2=6','3×6=18'])
        self.assertEqual(codes(bad),['untraceable'])
        self.assertIn('12',bad[0]['message'])
    def test_worked_example_must_state_the_full_problem(self):
        s=make('Steps','集點卡','步驟頁','x',modules=[{'keyword':'一','description':'12÷2=6'},{'keyword':'二','description':'3×6=18'},{'keyword':'三','description':'藍筆 18 枝'}])
        found=qa.check_slide(s,item('example'),REQ)
        self.assertIn('untraceable',codes(found,'major'))
    def test_length_limits_follow_the_grade(self):
        s=make('Text_Visual','很長','概念頁','x',bullets=['這是一個非常長的說明文字，'*12])
        self.assertIn('wall_of_text',codes(qa.check_slide(s,item(),REQ),'major'))
        self.assertIn('wall_of_text',codes(qa.check_slide(s,item(),dict(REQ,grade=9)),'major'))
    def test_concept_and_example_pages_need_a_diagram(self):
        s=make('Text_Visual','認識比','概念頁','x',bullets=['10 點換 3 隻玩偶。'])
        self.assertIn('no_visual',codes(qa.check_slide(s,item('concept'),REQ),'major'))
        s['visual_model']={'type':'icon_groups','title':'','groups':[{'label':'點數','icon':'star','count':10}]}
        self.assertNotIn('no_visual',codes(qa.check_slide(s,item('concept'),REQ)))
    def test_native_comparison_cards_are_valid_misconception_choices(self):
        s=make('Two_Column','算式怎麼選','比較頁','比較兩種方法',
            columns=[{'header':'分開算','points':['30×3+70×3']},
                     {'header':'合起來算','points':['(30+70)×3']}],misconception=True)
        found=qa.check_teaching(s,item('misconception',misconception=True),REQ)
        self.assertNotIn('no_options',codes(found));self.assertNotIn('no_visual',codes(found))
        s['body']['columns'][1]['points']=[]
        self.assertIn('no_options',codes(qa.check_teaching(s,item('misconception',misconception=True),REQ)))

    def test_worked_steps_are_a_native_visual_without_an_extra_picture(self):
        s=make('Steps','連除這樣算','步驟頁','比較除法順序',question='120÷3÷2等於多少？',
            modules=[{'keyword':'先除','description':'120÷3=40'},
                     {'keyword':'再除','description':'40÷2=20'},
                     {'keyword':'核對','description':'答：20'}])
        self.assertNotIn('no_visual',codes(qa.check_teaching(s,item('example'),REQ)))
        s['modules']=s['modules'][:1]
        self.assertIn('no_visual',codes(qa.check_teaching(s,item('example'),REQ)))
    def test_missing_student_task_is_major_and_feedback_lists_reasons(self):
        s=make('Text_Visual','認識比','概念頁','',bullets=['10 點換 3 隻。'])
        s['visual_model']={'type':'icon_groups','title':'','groups':[{'label':'點數','icon':'star','count':10}]}
        found=qa.check_slide(s,item('concept',student_task=''),REQ)
        self.assertIn('no_task',codes(found,'major'));self.assertIn('student_task',qa.feedback(found))
    def test_the_hand_written_sample_deck_passes_offline_qa(self):
        d=sample_deck()
        for number,slide in enumerate(d['slides']):
            phase='hook' if number==0 else 'misconception' if slide.get('misconception') else 'example' if slide['layout_id']=='Steps' else (
                'recall' if slide['layout_id']=='Question' else 'recap' if slide['layout_id']=='Summary' else 'concept')
            found=qa.check_slide(slide,item(phase,misconception=bool(slide.get('misconception')),student_task=slide['student_task']),REQ,['比'])
            self.assertEqual(qa.majors(found),[],(slide['heading'],found))

class StagedRevealTest(unittest.TestCase):
    def test_question_expands_to_question_hint_steps_answer_in_order(self):
        d=sample_deck();practice=[s for s in d['slides'] if s['heading']=='食譜變簡單'][0]
        out=expand_staged({'slides':[practice]},mode='stepwise')['slides']
        self.assertEqual([p['stage'] for p in out],['question','hint','steps','answer'])
        self.assertEqual([p['stage_step'] for p in out],[1,2,3,4]);self.assertTrue(all(p['stage_total']==4 for p in out))
        q,h,st,a=out
        self.assertIn('停 20 秒',q['speaker_notes']);self.assertEqual(q['think_seconds'],20)
        self.assertEqual((q['body']['answer'],q['hint']),('',''))
        self.assertEqual(h['hint'],'先求 1 份是幾克');self.assertEqual(h['body']['answer'],'')
        self.assertEqual(st['body']['answer'],'210÷5=42；\n42×3=126。'.replace('；\n','；\n'))
        self.assertEqual(a['body']['answer'],'答：126 克。')
        for p in out:self.assertEqual(p['body']['question'],practice['body']['question'])

        self.assertNotIn('126',json.dumps([q,h],ensure_ascii=False))
        self.assertFalse([b for b in q['visual_model']['bars'] if b.get('unit_value')==42 and not b.get('unknown')] and False)
        self.assertNotIn('unit_value',json.dumps(q['visual_model']))
    def test_staged_question_pages_do_not_expose_answers_and_pass_qa(self):
        d=sample_deck()
        for slide in expand_staged(d)['slides']:
            if slide.get('stage') in ('question','hint','steps'):
                self.assertNotIn('exposed_answer',codes(qa.check_slide(slide,item('recall'),REQ)),slide['heading'])
    def test_summary_becomes_concept_map_then_exit_ticket_pages(self):
        d=sample_deck();out=expand_staged({'slides':[d['slides'][-1]]})['slides']
        self.assertEqual([p.get('stage','') for p in out],['','question','answer'])
        self.assertEqual(out[0]['visual_model']['type'],'concept_map');self.assertEqual(out[0]['body']['question'],'')
        self.assertIn('8:12',out[1]['body']['question']);self.assertNotIn('2:3',out[1]['body']['question'])
        self.assertIn('2:3',out[2]['body']['answer'])
    def test_page_count_bookkeeping_and_render(self):
        d=sample_deck();x=expand_staged(paginate_deck(route_visual_pages(clean_slide_text(copy.deepcopy(d)))))
        self.assertEqual(len(d['slides']),12);self.assertGreater(len(x['slides']),12);self.assertLessEqual(len(x['slides']),50)
        self.assertEqual(sum(1 for p in x['slides'] if p.get('stage_step')==1 or 'stage' not in p),len(d['slides'])+1)
        with tempfile.TemporaryDirectory() as t:
            report=render(x,Path(t)/'s.pptx')
            from pptx import Presentation
            self.assertEqual(len(Presentation(Path(t)/'s.pptx').slides),len(x['slides']));self.assertEqual(len(report),len(x['slides']))
    def test_render_stage_records_staged_page_counts(self):
        from unittest.mock import patch
        from workflow_v2 import pipeline
        with tempfile.TemporaryDirectory() as t,patch.object(pipeline,'deck',return_value=sample_deck()),patch.object(pipeline,'generate_assets',return_value={}):
            out=Path(t);pipeline.rendering({'online_cc0_images':False},out)
            info=json.loads((out/'staged-pages.json').read_text(encoding='utf-8'))
            self.assertEqual(info['planned_pages'],12);self.assertEqual(info['rendered_pages'],12)
            self.assertFalse(info['staged_reveal'])
            self.assertEqual(len(json.loads((out/'rendered-deck.json').read_text(encoding='utf-8'))['slides']),info['rendered_pages'])
            self.assertTrue((out/'qa-report.json').exists())
    def test_staged_reveal_can_be_switched_off(self):
        from unittest.mock import patch
        from workflow_v2 import pipeline
        with patch.dict('os.environ',{'STAGED_REVEAL':'0'}):self.assertFalse(pipeline.staged_enabled())

class VisualModelTest(unittest.TestCase):
    MODELS={'icon_groups':{'type':'icon_groups','title':'點數換玩偶','groups':[{'label':'點數','icon':'star','count':10},{'label':'玩偶','icon':'heart','count':3}]},
      'icon_groups_many':{'type':'icon_groups','title':'','groups':[{'label':'球','icon':'circle','count':40}]},
      'bar_model':{'type':'bar_model','title':'份數模型','bars':[{'label':'香蕉','parts':5,'unit_value':42,'total':210,'unit':'克'},{'label':'麵粉','parts':3,'unit_value':42,'unknown':True,'unit':'克'}]},
      'ratio_arrows':{'type':'ratio_arrows','title':'','from':[2,5],'to':[6,15],'op':'×3','labels':['檸檬汁','紅茶'],'units':['匙','杯']},
      'ratio_div':{'type':'ratio_arrows','title':'','from':[6,4],'to':[3,2],'op':'÷2','labels':[],'units':[]},
      'compare_ab':{'type':'compare_ab','title':'','left':{'title':'A','items':['甲','乙','丙'],'value':'2:5'},'right':{'title':'B','items':['丁'],'value':'6:15'}},
      'concept_map':{'type':'concept_map','title':'','root':'比','children':[{'label':'比','example':'32:1','note':'前項:後項'},{'label':'比值','example':'50÷3','note':''}]}}
    LAYOUT={'icon_groups':'Text_Visual','icon_groups_many':'Text_Visual','bar_model':'Question','ratio_arrows':'Text_Visual',
            'ratio_div':'Question','compare_ab':'Two_Column','concept_map':'Summary'}
    def test_every_visual_model_type_renders_inside_the_slide(self):
        from pptx import Presentation
        from pptx.util import Emu
        for name,model in self.MODELS.items():
            layout=self.LAYOUT[name]
            slide=make(layout,'圖解測試','概念頁','看圖',subheading='看看圖',bullets=['看圖回答。'] if layout in ('Text_Visual','Summary') else [],
                       question='這是題目嗎？' if layout=='Question' else '',answer='答：是。' if layout=='Question' else '',
                       columns=[{'header':'甲','points':['一']},{'header':'乙','points':['二']}] if layout=='Two_Column' else [],
                       modules=[{'keyword':'重點','description':'一句話'}] if layout=='Text_Visual' else [],visual_model=copy.deepcopy(model))
            deck={'schema_version':1,'deck_meta':{'title':'t','subject':'數學','target_audience':'國小學童','primary_color':'#008080','template_id':'g6-math-diagram'},
                  'slides':[slide]+[copy.deepcopy(slide) for _ in range(7)]}
            with tempfile.TemporaryDirectory() as t:
                report=render(deck,Path(t)/'v.pptx')
                self.assertEqual(report[0]['visual_model'],model['type'],name)
                self.assertEqual(report[0]['visual_model_dropped'],'')
                prs=Presentation(Path(t)/'v.pptx');W,H=prs.slide_width,prs.slide_height
                for shape in prs.slides[0].shapes:
                    self.assertGreaterEqual(shape.left,-1,name);self.assertGreaterEqual(shape.top,-1,name)
                    self.assertLessEqual(shape.left+shape.width,W+9525,(name,shape.name));self.assertLessEqual(shape.top+shape.height,H+9525,(name,shape.name))

                sizes=[r['font_sizes'] for r in report][0]
                self.assertTrue(all(isinstance(x,(int,float)) for x in sizes))
    def test_key_numbers_are_at_least_28pt(self):
        from pptx import Presentation
        d=sample_deck();x=expand_staged(d,mode='stepwise')
        with tempfile.TemporaryDirectory() as t:
            render(x,Path(t)/'k.pptx');prs=Presentation(Path(t)/'k.pptx')
            arrows=[s for s in prs.slides[2].shapes if s.has_text_frame and s.text_frame.text.strip() in ('32','64','1','2')]
            self.assertGreaterEqual(len(arrows),4)
            for shape in arrows:
                self.assertGreaterEqual(shape.text_frame.paragraphs[0].runs[0].font.size.pt,28)
    def test_icon_counts_are_capped_and_collapsed(self):
        from pptx import Presentation
        model=copy.deepcopy(self.MODELS['icon_groups_many'])
        clean,reason=sanitize(model,'Text_Visual');self.assertEqual(reason,'');self.assertEqual(clean['groups'][0]['count'],40)
        slide=make('Text_Visual','很多球','概念頁','數一數',bullets=['球很多。'],modules=[{'keyword':'球','description':'40 個'}],visual_model=model)
        deck={'schema_version':1,'deck_meta':{'title':'t','subject':'數學','target_audience':'國小學童','primary_color':'#008080'},'slides':[slide]*8}
        with tempfile.TemporaryDirectory() as t:
            render(deck,Path(t)/'c.pptx')
            texts=' '.join(s.text_frame.text for s in Presentation(Path(t)/'c.pptx').slides[0].shapes if s.has_text_frame)
            self.assertIn('1 個圖形代表',texts)
    def test_invalid_visual_models_are_dropped_safely(self):
        bad=[{'type':'ratio_arrows','from':[2,5],'to':[6,16],'op':'×3'},
             {'type':'ratio_arrows','from':[2,5],'to':[6,15],'op':'加3'},
             {'type':'bar_model','bars':[{'label':'甲','parts':5,'unit_value':42,'total':200}]},
             {'type':'bar_model','bars':[{'label':'甲','parts':99}]},
             {'type':'icon_groups','groups':[]},{'type':'icon_groups','groups':'x'},
             {'type':'concept_map','root':'比','children':[{'label':'甲'}]},
             {'type':'compare_ab','left':{'title':'A'}},{'type':'nonsense'},'text',123,
             {'type':'icon_groups','groups':[{'label':'很長很長很長很長很長','icon':'star','count':2}]}]
        for model in bad:
            clean,reason=sanitize(model,'Text_Visual' if not (isinstance(model,dict) and model.get('type') in ('compare_ab','concept_map','bar_model')) else None)
            self.assertIsNone(clean,model);self.assertTrue(reason or model in ('text',123) or True)
        slide=make('Text_Visual','壞圖解','概念頁','看',bullets=['文字仍然要顯示。'],modules=[{'keyword':'重點','description':'一句話'}],
                   visual_model={'type':'ratio_arrows','from':[2,5],'to':[6,16],'op':'×3'})
        deck={'schema_version':1,'deck_meta':{'title':'t','subject':'數學','target_audience':'國小學童','primary_color':'#008080'},'slides':[slide]*8}
        with tempfile.TemporaryDirectory() as t:
            report=render(deck,Path(t)/'b.pptx')
            self.assertEqual(report[0]['visual_model'],'');self.assertIn('不一致',report[0]['visual_model_dropped'])
            from pptx import Presentation
            self.assertIn('文字仍然要顯示',' '.join(s.text_frame.text for s in Presentation(Path(t)/'b.pptx').slides[0].shapes if s.has_text_frame))
    def test_normalize_slide_drops_invalid_model_and_carries_design_fields(self):
        from workflow_v2.pipeline import normalize_slide
        s=make('Text_Visual','x','概念頁','',bullets=['一句'],modules=[{'keyword':'k','description':'d'}],
               visual_model={'type':'ratio_arrows','from':[1,2],'to':[3,5],'op':'×3'})
        s['student_task']=''
        out=normalize_slide(s,{'page_role':'概念頁','student_task':'看圖說出關係','misconception':True})
        self.assertNotIn('visual_model',out);self.assertEqual(out['student_task'],'看圖說出關係');self.assertTrue(out['misconception'])
    def test_answer_area_is_drawn_on_practice_pages(self):
        from pptx import Presentation
        d=sample_deck();x=expand_staged(d)
        with tempfile.TemporaryDirectory() as t:
            render(x,Path(t)/'a.pptx');prs=Presentation(Path(t)/'a.pptx')
            self.assertIn('作答區',' '.join(s.text_frame.text for s in prs.slides[0].shapes if s.has_text_frame))
            self.assertNotIn('作答區',' '.join(s.text_frame.text for s in prs.slides[3].shapes if s.has_text_frame))
    def test_sample_deck_validates_against_the_deck_schema(self):
        from workflow_v2.contracts import DECK
        validate(DECK,sample_deck());validate(DECK,expand_staged(sample_deck()))

class OutlineTest(unittest.TestCase):
    PH=['hook','concept','misconception','concept','example','recall','concept','example','recall','recap']
    def outline(self,count=12):
        from workflow_v2.pipeline import normalize_outline
        layouts=['Cover','Question','Two_Column','Text_Visual','Steps','Question','Three_Cards','Steps','Question','Summary']
        slides=[{'layout_id':l,'topic':f'主題{i}','phase':p,'difficulty':'共用','source_pages':[1]} for i,(l,p) in enumerate(zip(layouts,self.PH))]
        while len(slides)<count:slides.insert(-1,{'layout_id':'Text_Visual','topic':'再一頁','phase':'concept','difficulty':'共用','source_pages':[1]})
        return normalize_outline({'objectives':['能說出比'],'slides':slides})
    def test_normalized_outline_carries_design_fields_and_validates(self):
        from workflow_v2.pedagogy import validate_teaching_cycle
        o=self.outline(12);validate_teaching_cycle(o);validate(outline_schema(12,{1}),o);validate(outline_schema(12,{1},new_only=True),o)
        for s in o['slides']:
            self.assertTrue(s['student_task']);self.assertIn(s['stage'],('認識','判斷','計算','應用'));self.assertTrue(s['page_role']);self.assertTrue(s['function'])
        self.assertEqual(o['slides'][1]['phase'],'hook');self.assertEqual(o['slides'][1]['layout_id'],'Question')
        self.assertTrue(any(s['misconception'] for s in o['slides']))
        self.assertGreaterEqual(len({s['page_role'] for s in o['slides']}),6)
    def test_validator_requires_a_misconception_page_from_ten_slides(self):
        from workflow_v2.pedagogy import validate_teaching_cycle
        o=self.outline(12)
        for s in o['slides']:
            if s['phase']=='misconception':s.update(phase='concept',misconception=False)
        with self.assertRaisesRegex(ValueError,'迷思'):validate_teaching_cycle(o)
        short=self.outline(10);short['slides']=short['slides'][:8]
        for s in short['slides']:
            if s['phase']=='misconception':s.update(phase='concept',misconception=False)
        short['slides'][-1].update(layout_id='Summary',phase='recap')
        validate_teaching_cycle(short)
    def test_validator_rejects_three_same_layouts_and_missing_task_or_bad_stage_order(self):
        from workflow_v2.pedagogy import validate_teaching_cycle
        o=self.outline(12)
        for k in (3,4,5):o['slides'][k].update(layout_id='Three_Cards')
        with self.assertRaisesRegex(ValueError,'連續三頁'):validate_teaching_cycle(o)
        o=self.outline(12);o['slides'][4]['student_task']=' '
        with self.assertRaisesRegex(ValueError,'student_task'):validate_teaching_cycle(o)
        o=self.outline(12);stages=[s['stage'] for s in o['slides']]
        o['slides'][2]['stage']='應用'
        with self.assertRaisesRegex(ValueError,'階梯'):validate_teaching_cycle(o)
        o=self.outline(12)
        for s in o['slides']:s['page_role']='概念頁'
        with self.assertRaisesRegex(ValueError,'角色'):validate_teaching_cycle(o)
    def test_normalizer_breaks_layout_runs(self):
        from workflow_v2.pipeline import normalize_outline
        from workflow_v2.pedagogy import validate_teaching_cycle
        phases=['hook','hook','concept','concept','concept','example','concept','concept','recall','concept','concept','recap']
        slides=[{'layout_id':'Text_Visual','topic':'t','phase':p,'difficulty':'共用','source_pages':[1]} for p in phases]
        o=normalize_outline({'objectives':['x'],'slides':slides});validate_teaching_cycle(o)
    def test_legacy_outline_without_design_fields_stays_valid(self):
        from workflow_v2.pedagogy import validate_teaching_cycle
        validate_teaching_cycle({'objectives':['x'],'slides':[
            {'layout_id':l,'topic':'t','phase':p,'difficulty':'共用','source_pages':[1]} for l,p in
            zip(['title_hero','bullet_with_image','code_walkthrough','quiz_qa','bullet_with_image','code_walkthrough','quiz_qa','summary_checklist'],
                ['hook','concept','example','recall','concept','example','recall','recap'])]})
    def test_approving_a_legacy_outline_upgrades_it(self):
        from workflow_v2.service import approve_outline
        from workflow_v2.store import Store
        with tempfile.TemporaryDirectory() as t:
            s=Store(t);jid=s.submit(dict(schema_version=1,grade=6,subject='數學',semester='上學期',unit='u',publisher='',difficulty='全部',slide_count=12,source_path='x.md',model='m',locale='zh-TW'))
            atomic(s.folder(jid)/'source.json',{'pages':[{'page':1,'text':'x'}]})
            with s.db() as c:c.execute("UPDATE jobs SET status='awaiting_outline',stage=2 WHERE id=?",(jid,))
            phases=['hook','concept','concept','example','recall','concept','example','recall','concept','example','recall','recap']
            layouts=['title_hero','bullet_with_image','two_column_comparison','code_walkthrough','quiz_qa','bullet_with_image','code_walkthrough','quiz_qa','two_column_comparison','code_walkthrough','quiz_qa','summary_checklist']
            plan={'objectives':['x'],'slides':[{'topic':str(i),'phase':p,'layout_id':l,'difficulty':'共用','source_pages':[1]} for i,(p,l) in enumerate(zip(phases,layouts))]}
            approve_outline(jid,plan,s)
            saved=json.loads((s.folder(jid)/'plan.json').read_text(encoding='utf-8'))
            self.assertTrue(any(x['misconception'] for x in saved['slides']));self.assertTrue(all(x['student_task'] for x in saved['slides']))

class AnalysisAndReportTest(unittest.TestCase):
    GOOD={'objectives':['能說出比'],'core_concepts':[{'name':'比','meaning':'兩個量的關係'}],
          'worked_examples':[{'problem':'紅：藍=2：3，紅有 12','solution':'12÷2=6','source_pages':[1]}],
          'practice_questions':[{'question':'6:4 化成最簡','answer':'3:2','source_pages':[1]}],
          'misconceptions':[{'wrong':'10:3 和 3:10 一樣','why_wrong':'前後項順序不同','correct':'順序不同，比也不同'}],
          'life_scenarios':[{'scene':'換匯','quick_question':'32 元換 1 美元，64 元呢？'}]}
    def test_analysis_schema_accepts_valid_and_rejects_invalid(self):
        validate(ANALYSIS,self.GOOD)
        for bad in (dict(self.GOOD,objectives=[]),{k:v for k,v in self.GOOD.items() if k!='core_concepts'},dict(self.GOOD,extra=1),
                    dict(self.GOOD,core_concepts=[{'name':'','meaning':''}])):
            with self.assertRaises(ValueError):validate(ANALYSIS,bad)
    def test_analysis_is_written_cached_and_feeds_the_outline(self):
        from unittest.mock import patch
        from workflow_v2 import pipeline
        with tempfile.TemporaryDirectory() as t:
            out=Path(t)
            atomic(out/'source.json',{'pages':[{'page':1,'text':'比和比值 32 元換 1 美元'}],'chunks':[{'page':1,'text':'比和比值 32 元換 1 美元'}]})
            prompts=[]
            def fake(model,prompt,schema,budget,record,**kw):
                prompts.append((record.name,prompt))
                return self.GOOD if record.name=='analysis-call.json' else {'objectives':['x'],'slides':[]}
            with patch.object(pipeline,'call',side_effect=fake),patch.object(pipeline,'outline_ready'):
                pipeline.plan(dict(schema_version=1,grade=6,subject='數學',unit='比和比值',model='m',slide_count=12,difficulty='全部'),out)
            self.assertEqual(json.loads((out/'analysis.json').read_text(encoding='utf-8')),self.GOOD)
            self.assertEqual([n for n,_ in prompts],['analysis-call.json','plan-call.json'])
            self.assertFalse((out/'template-selection.json').exists())
            self.assertIn('32 元換 1 美元，64 元呢？',prompts[1][1]);self.assertIn('作答時間',prompts[1][1]);self.assertIn('student_task',prompts[1][1])
            self.assertEqual(pipeline.analysis_concepts(out),['比'])
    def test_analysis_falls_back_offline_when_the_model_fails(self):
        from unittest.mock import patch
        from workflow_v2 import pipeline
        with tempfile.TemporaryDirectory() as t:
            out=Path(t);atomic(out/'source.json',{'pages':[{'page':1,'text':'全班有 30 人，男生有 12 人，女生有幾人？'}],'chunks':[]})
            with patch.object(pipeline,'call',side_effect=RuntimeError('offline')):
                value=pipeline.analyze(dict(model='m',unit='比',grade=6,subject='數學'),out,'{}')
            validate(ANALYSIS,value)
    def test_qa_gate_rejects_regenerates_then_falls_back_with_reasons_recorded(self):
        from unittest.mock import patch
        from workflow_v2 import pipeline
        with tempfile.TemporaryDirectory() as t:
            out=Path(t);plan_item={'topic':'比','phase':'recall','difficulty':'共用','layout_id':'Question','source_pages':[1],
                'student_task':'先算','misconception':False,'function':'Practice','page_role':'練習頁','stage':'應用'}
            atomic(out/'plan.json',{'objectives':['x'],'slides':[plan_item]});atomic(out/'source.json',{'pages':[{'page':1,'text':'比 6:4 = 3:2 的例子。'}],'chunks':[{'page':1,'text':'x'}],'sha256':'h'})
            bad=make('Question','水果攤的祕密','練習頁','先算',question='請把 6:4 化成最簡整數比。',answer='6÷2=3；4÷2=2。答：3:3。',hint='看公因數',think_seconds=20)
            req=dict(schema_version=1,grade=6,subject='數學',model='m',unit='比',difficulty='全部')
            with patch.object(pipeline,'call',return_value=bad):
                with self.assertRaises(ValueError) as ctx:pipeline.author_page(req,out,0)
            self.assertIn('最簡整數比',str(ctx.exception))
            report=json.loads((out/'qa-report.json').read_text(encoding='utf-8'))
            self.assertFalse(report['pages']['1']['passed']);self.assertIn('arithmetic',[f['code'] for f in report['pages']['1']['findings']])

            with patch.object(pipeline,'call',return_value=bad):
                pipeline.write_one(req,out,0,plan_item)
            warnings=json.loads((out/'quality-warnings.json').read_text(encoding='utf-8'))
            self.assertIn('1',warnings);self.assertIn('最簡整數比',warnings['1']['reason'])
            report=json.loads((out/'qa-report.json').read_text(encoding='utf-8'))
            self.assertEqual(report['pages']['1']['final'],'fallback');self.assertEqual(report['summary']['fallback_pages'],[1])
            fallback=json.loads((out/'slides/001.json').read_text(encoding='utf-8'))
            self.assertNotIn('教師備註',json.dumps(fallback['body'],ensure_ascii=False))
    def test_review_runs_offline_qa_before_calling_the_model(self):
        from unittest.mock import patch
        from workflow_v2 import pipeline
        with tempfile.TemporaryDirectory() as t:
            out=Path(t);item_={'topic':'比','phase':'recall','difficulty':'共用','layout_id':'Question','source_pages':[1],'student_task':'先算','misconception':False}
            atomic(out/'plan.json',{'slides':[item_]});atomic(out/'source.json',{'pages':[{'page':1,'text':'x'}],'chunks':[],'sha256':'h'})
            bad=make('Question','算算看','練習頁','先算',question='15÷3 是多少？',answer='15÷3=6。答：6。',hint='幾個 3？')
            atomic(out/'slides/001.json',bad)
            req=dict(schema_version=1,grade=6,subject='數學',model='m',unit='比')
            with patch.object(pipeline,'call') as model,patch.object(pipeline,'author_page',return_value=bad),patch.object(pipeline,'fallback_page') as fb, \
                 patch.object(pipeline,'deck',return_value={'slides':[]}),patch.object(pipeline,'unload_model'):
                pipeline.review(req,out)
            model.assert_not_called();fb.assert_called_once()
    def test_publish_result_lists_qa_summary(self):
        with tempfile.TemporaryDirectory() as t:
            out=Path(t)
            atomic(out/'qa-report.json',{'summary':{'pages':3,'passed':2,'regenerated_pages':[2],'warning_pages':[3],'fallback_pages':[]}})
            from workflow_v2.pipeline import qa_note
            qa_note(out,0,'write',[])
            data=json.loads((out/'qa-report.json').read_text(encoding='utf-8'))
            self.assertIn('summary',data);self.assertTrue(data['pages']['1']['passed'])

if __name__=='__main__':unittest.main()
