import copy,io,json,tempfile,unittest
from pathlib import Path
from unittest.mock import Mock,patch
from pptx import Presentation
from PIL import Image
from workflow_v2 import final_optimizer as final
from workflow_v2.renderer import render
from workflow_v2.page_policy import prepare_deck,reflow_single_page
from workflow_v2.store import atomic
from workflow_v2 import qa
from tests.test_design_system import sample_deck

class FinalOptimizerTest(unittest.TestCase):
    def fixture(self,out):
        deck=sample_deck();deck['slides'][2]['body']['bullets']=['這是一個關於位置的說明。']
        deck['slides'][2].pop('visual_model',None)
        render(deck,out/'教材.pptx');atomic(out/'rendered-deck.json',deck)
        prs=Presentation(out/'教材.pptx');buffer=io.BytesIO()
        Image.new('RGB',(20,12),'teal').save(buffer,format='PNG');buffer.seek(0)
        prs.slides[0].shapes.add_picture(buffer,7000000,3000000,width=914400)
        prs.save(out/'教材.pptx')
        return {'grade':6,'subject':'數學','unit':'位置','model':'example:model'}

    def test_last_pass_edits_text_without_moving_or_restyling_anything(self):
        with tempfile.TemporaryDirectory() as folder:
            out=Path(folder);request=self.fixture(out)
            before=Presentation(out/'教材.pptx');mapping,pages=final.fields(before)
            key=next(key for key,value in mapping.items() if value[0].text=='這是一個關於位置的說明。')
            model=Mock(side_effect=[{'revisions':[{'field':key,'text':'這頁說明位置。','reason':'刪除贅詞'}],'issues':[]},
                                    {'revisions':[],'issues':[]}])
            result=final.optimize(request,out,model)
            after=Presentation(out/'教材.pptx')
            self.assertEqual(final.geometry(before),final.geometry(after))
            self.assertEqual(final.formatting(before),final.formatting(after))
            self.assertIn('這頁說明位置。','\n'.join(shape.text for shape in after.slides[2].shapes if shape.has_text_frame))
            self.assertTrue(result['changed']);self.assertEqual(result['model_reviewed_pages'],8)
            self.assertEqual(result['steps'][12]['status'],'manual_playback_required')
            self.assertEqual(model.call_count,2)
            self.assertEqual(final.optimize(request,out,model),result);self.assertEqual(model.call_count,2)
            self.assertIn('這頁說明位置。',(out/'final-optimized-deck.json').read_text(encoding='utf-8'))

    def test_numeric_units_entities_and_negation_cannot_change_in_a_rephrase(self):
        prs=Presentation();shape=prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(0,0,9144000,914400)
        for old,new in [('5 公斤共有10元','5 公尺共有10元'),('The panda is under the desk.','The cat is under the desk.'),
                        ('這個答案不是正確的。','這個答案是正確的。'),('共有18元。','共有12元。'),
                        ('找出兩條邊的共同端點','找出兩條線的共同端點')]:
            shape.text=old;run=shape.text_frame.paragraphs[0].runs[0]
            accepted,_=final.safe_edit((run,shape.text_frame,shape.width,shape.height,1),new)
            self.assertFalse(accepted);self.assertEqual(run.text,old)

    def test_model_failure_keeps_original_and_reports_incomplete_review(self):
        with tempfile.TemporaryDirectory() as folder:
            out=Path(folder);request=self.fixture(out);before=final.sha(out/'教材.pptx')
            result=final.optimize(request,out,Mock(side_effect=RuntimeError('離線')))
            self.assertEqual(result['status'],'unavailable');self.assertEqual(result['model_reviewed_pages'],0)
            self.assertFalse(result['changed']);self.assertEqual(final.sha(out/'教材.pptx'),before)
            self.assertTrue(result['requires_teacher_review'])

    def test_punctuation_split_between_runs_is_checked_as_a_complete_frame(self):
        prs=Presentation();shape=prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(0,0,9144000,914400)
        paragraph=shape.text_frame.paragraphs[0]
        paragraph.add_run().text='(1'
        run=paragraph.add_run();run.text=') The ball is under the box. 哪個詞說明位置？'
        accepted,_=final.safe_edit((run,shape.text_frame,shape.width,shape.height,1),
            ') The ball is under the box. 哪詞說明位置？')
        self.assertTrue(accepted)
        self.assertEqual(shape.text.count('('),shape.text.count(')'))

    def test_changed_audience_invalidates_the_final_review_cache(self):
        with tempfile.TemporaryDirectory() as folder:
            out=Path(folder);request=self.fixture(out)
            model=Mock(return_value={'revisions':[],'issues':[]})
            first=final.optimize(request,out,model)
            request['audience']='國小學生'
            second=final.optimize(request,out,model)
            self.assertEqual(model.call_count,4)
            self.assertNotEqual(first['context_sha256'],second['context_sha256'])

    def test_any_format_mutation_rolls_back_all_proposed_changes(self):
        with tempfile.TemporaryDirectory() as folder:
            out=Path(folder);request=self.fixture(out);before=final.sha(out/'教材.pptx')
            def mutating_edit(entry,text):
                entry[0].text=text
                entry[0].font.bold=True
                return True,''
            def propose(model,prompt,schema,*args,**kwargs):
                key=schema['properties']['revisions']['items']['properties']['field']['enum'][0]
                return {'revisions':[{'field':key,'text':'較短文字','reason':'測試意外格式異動'}],'issues':[]}
            with patch.object(final,'safe_edit',side_effect=mutating_edit):
                result=final.optimize(request,out,propose)
            self.assertFalse(result['changed']);self.assertTrue(result['requires_teacher_review'])
            self.assertEqual(final.sha(out/'教材.pptx'),before)
            self.assertTrue(result['rejected_changes'])

    def test_calculator_repairs_affirmed_values_but_preserves_misconceptions(self):
        page=sample_deck()['slides'][6];page['body']['answer']='5:10 的比值為 2。2:3 的比值為 8。24和36的最大公因數為9。'
        original=copy.deepcopy(page);notes=qa.autofix(page)
        self.assertIn('比值為 0.5',page['body']['answer']);self.assertIn('比值為 2/3',page['body']['answer'])
        self.assertIn('最大公因數為12',page['body']['answer']);self.assertFalse(qa.check_arithmetic(page['body']['answer']))
        self.assertGreaterEqual(len(notes),3)
        original['misconception']=True;qa.autofix(original);self.assertIn('比值為 2。',original['body']['answer'])

    def test_title_binds_to_the_question_without_changing_the_conditions(self):
        page=sample_deck()['slides'][6];page['heading']='紅茶的比值';page['body']['question']='檸檬汁5杯：檸檬汁10杯的比值是多少？'
        question=page['body']['question'];qa.autofix(page)
        self.assertEqual(page['heading'],'檸檬汁的比值');self.assertEqual(page['body']['question'],question)

    def test_overflow_reflows_one_page_and_preserves_full_source_in_notes(self):
        deck=sample_deck();page=deck['slides'][6]
        page['body']['question']='這是一個必須保留條件的完整題目。'*100
        page['body']['answer']='這是完整的第一題答案與第二題答案。'*100
        result,complete=reflow_single_page(deck,6)
        self.assertEqual(len(result['slides']),len(deck['slides']));self.assertFalse(complete)
        self.assertIn(page['body']['answer'],result['slides'][6]['speaker_notes'])
        with tempfile.TemporaryDirectory() as folder:render(result,Path(folder)/'single.pptx')

    def test_supplement_moves_as_whole_items_and_not_sentence_fragments(self):
        deck=sample_deck();page=deck['slides'][2];page['body']['bullets']=['第一個概念。'*8,'第二個補充。'*8,'第三個補充。'*8]
        result=prepare_deck(deck,{'grade':4})
        self.assertEqual(len(result['slides']),8)
        self.assertTrue(all(p in page['body']['bullets'] for p in result['slides'][2]['body']['bullets']))
        self.assertIn('第三個補充。'*8,result['slides'][2]['speaker_notes'])

    def test_answer_animation_targets_only_existing_answer_shapes(self):
        deck=sample_deck()
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'appear.pptx';render(deck,path);prs=Presentation(path)
            slide=prs.slides[6];ids={str(s.shape_id) for s in slide.shapes}
            targets=slide._element.xpath('./p:timing//p:spTgt')
            self.assertTrue(targets);self.assertTrue(all(t.get('spid') in ids for t in targets))
            self.assertEqual(len(prs.slides),8)

if __name__=='__main__':unittest.main()
