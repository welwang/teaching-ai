"""國中（7-9 年級）任務控制風格：對比、字級、只影響國中、年齡層規則不變、舊教材可重新排版。"""
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation
from pptx.enum.dml import MSO_FILL
from pptx.util import Emu

from tests.test_beautify import soffice_available
from tests.test_design_system import sample_deck
from workflow_v2 import age_bands, beautify, deck_check, restyle, template_styles as ts, visual_qa
from workflow_v2.renderer import render
from workflow_v2.store import Store, atomic

ROOT = Path(__file__).resolve().parents[1]
SUBJECTS = ('數學', '國語', '英語')


def junior_deck(grade=8, subject='數學'):
    deck = sample_deck()
    deck['deck_meta'].update(subject=subject, grade=grade, target_audience='國中學生')
    return deck


class StyleSpecTest(unittest.TestCase):
    def test_only_grades_7_to_9_get_the_mission_look(self):
        for grade in range(1, 10):
            for subject in SUBJECTS:
                style = ts.style_for(grade, subject)
                self.assertEqual(style.get('look') == 'mission', grade >= 7, (grade, subject))
        self.assertNotIn('look', ts.style_for(None, None))

    def test_primary_school_styles_are_exactly_what_they_were(self):
        expected = {(2, '數學'): ('math-low', '0B6B73', 'AA5009', 'F7F3E9', 46), (4, '數學'): ('math-mid', '0B6B73', 'AA5009', 'F7F3E9', 44),
                    (6, '數學'): ('math-high', '116A73', 'AB5411', 'FAF7F1', 40)}
        for (grade, subject), (sid, primary, accent, paper, title) in expected.items():
            style = ts.style_for(grade, subject)
            self.assertEqual((style['id'], style['primary'], style['accent'], style['background'], style['title_max']),
                             (sid, primary, accent, paper, title))
            self.assertEqual(style['style'], 'story')

    def test_one_restrained_accent_clean_white_and_near_black(self):
        for subject in SUBJECTS:
            style = ts.style_for(8, subject)
            self.assertEqual(style['background'], 'FFFFFF')
            self.assertEqual(style['accent'], '0A5EA8')
            self.assertLess(ts.luminance(style['ink']), .01)
            self.assertLess(ts.luminance(style['dark']), .01)
        self.assertEqual(ts.style_for(7, '數學')['accent'], ts.style_for(9, '英語')['accent'])

    def test_every_text_pair_has_projection_grade_contrast(self):
        for grade in (7, 8, 9):
            for subject in SUBJECTS:
                s = ts.style_for(grade, subject)
                report = beautify.palette_report(s)
                self.assertGreaterEqual(min(report.values()), 4.5, report)
                self.assertGreaterEqual(report['ink-on-paper'], 15)
                pairs = {'muted-on-paper': (s['muted'], s['background']), 'muted-on-pale': (s['muted'], s['pale']),
                         'title-on-dark': (s['dark_ink'], s['dark']), 'intro-on-dark': (s['dark_muted'], s['dark']),
                         'label-on-dark': (s['dark_accent'], s['dark']), 'accent-on-pale': (s['accent'], s['pale'])}
                for name, (a, b) in pairs.items():
                    self.assertGreaterEqual(ts.contrast(a, b), 4.5, (grade, subject, name))
                self.assertGreaterEqual(ts.contrast(s['dark_ink'], s['dark']), 15)
                self.assertGreaterEqual(ts.contrast(s['hairline'], s['background']), 1.1)

    def test_titles_are_big_and_age_band_rules_are_untouched(self):
        self.assertGreaterEqual(ts.style_for(8, '數學')['title_max'], 40)
        junior = age_bands.BANDS['junior']
        self.assertEqual((junior['label'], junior['grades'], junior['body_target'], junior['body_max'], junior['heading_max'],
                          junior['bullets_max']), ('國中', (7, 8, 9), (70, 120), 160, 26, 5))
        sizes = [ts.style_for(g, '數學')['title_max'] for g in (2, 4, 6, 8)]
        self.assertEqual(sizes, sorted(sizes, reverse=True))


class RenderedDeckTest(unittest.TestCase):
    def build(self, grade=8, subject='數學'):
        folder = Path(tempfile.mkdtemp())
        report = render(junior_deck(grade, subject), folder / 'deck.pptx', grade=grade)
        return folder, Presentation(folder / 'deck.pptx'), report

    def test_cover_and_section_are_full_bleed_near_black_and_content_pages_are_white(self):
        _, prs, report = self.build()
        for slide, row in zip(prs.slides, report):
            panels = [s for s in slide.slide_layout.shapes if s.width == Emu(int(13.333 * 914400)) and s.height == Emu(int(7.5 * 914400))]
            if row['layout'] in ('Cover', 'Section'):
                self.assertEqual(len(panels), 1, row['layout'])
                self.assertEqual(str(panels[0].fill.fore_color.rgb), '0B0D10')
            else:
                self.assertEqual(panels, [], row['layout'])

    def test_no_clip_art_icon_or_pill_and_a_plain_page_number(self):
        from workflow_v2.slide_layouts import ROLE_TAG  # noqa: F401 - role chips stay (they are labels, not decoration)
        _, prs, _ = self.build()
        total = len(prs.slides)
        for number, slide in enumerate(prs.slides, 1):
            numbers = [s for s in slide.shapes if s.name == '頁碼']
            self.assertEqual(len(numbers), 1)
            self.assertEqual(numbers[0].text_frame.text, f'{number} / {total}')
            self.assertNotEqual(numbers[0].fill.type, MSO_FILL.SOLID)
            self.assertEqual(numbers[0].text_frame.paragraphs[0].runs[0].font.size.pt, 12)
            decorations = [s for s in slide.shapes if s.shape_type == 1 and s.auto_shape_type in beautify.ROLE_ICON.values()]
            self.assertEqual(decorations, [], number)

    def test_primary_school_decks_keep_their_bars_icons_and_pills(self):
        _, prs, _ = self.build(grade=6)
        slide = prs.slides[2]
        self.assertTrue([s for s in slide.shapes if s.name == '頁碼' and s.fill.type == MSO_FILL.SOLID])

    def test_all_font_sizes_are_projection_sized(self):
        for grade in (7, 8, 9):
            for subject in SUBJECTS:
                _, _, report = self.build(grade, subject)
                for row in report:
                    sizes = [s for s in row['font_sizes'] if s]
                    self.assertGreaterEqual(min(sizes), 18, (grade, subject, row['slide'], sizes))
                    self.assertGreaterEqual(sizes[0], 28, (grade, subject, row['slide'], 'title'))

    def test_every_run_is_at_least_the_footer_minimum_and_text_is_never_light_on_white(self):
        _, prs, _ = self.build()
        for slide in prs.slides:
            dark_page = any(s.width == Emu(int(13.333 * 914400)) for s in slide.slide_layout.shapes)
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        if run.font.size:
                            self.assertGreaterEqual(run.font.size.pt, 11)
                        if not dark_page and run.font.color and run.font.color.type is not None and run.text.strip() \
                                and shape.fill.type is None:
                            self.assertGreaterEqual(ts.contrast(str(run.font.color.rgb), 'FFFFFF'), 3.0, run.text)

    def test_cover_text_is_light_on_the_dark_page(self):
        _, prs, _ = self.build()
        cover = prs.slides[0]
        colours = {str(r.font.color.rgb) for s in cover.shapes if s.has_text_frame for p in s.text_frame.paragraphs for r in p.runs
                   if r.font.color and r.font.color.type is not None and r.text.strip() and s.name != '頁碼'}
        self.assertTrue(colours)
        for colour in colours:
            self.assertGreaterEqual(ts.contrast(colour, '0B0D10'), 7, colour)

    def test_content_checks_do_not_depend_on_the_look(self):
        deck = junior_deck()
        low = deck_check.check_deck(deck, {'grade': 8})
        self.assertEqual(low['band'], '國中')
        self.assertEqual(low['major'], 0)
        self.assertEqual(deck_check.check_deck(deck, {'grade': 8})['findings'], low['findings'])


@unittest.skipUnless(soffice_available(), 'LibreOffice is not installed')
class VisualQaTest(unittest.TestCase):
    def test_real_junior_decks_pass_the_visual_check_in_all_three_grades(self):
        from workflow_v2.office import convert_pdf
        for grade, subject in ((7, '國語'), (8, '數學'), (9, '英語')):
            folder = Path(tempfile.mkdtemp())
            render(junior_deck(grade, subject), folder / '教材.pptx', grade=grade)
            convert_pdf(folder / '教材.pptx', folder)
            report = visual_qa.run(folder, lambda: convert_pdf(folder / '教材.pptx', folder))
            self.assertEqual(report['remaining'], [], (grade, subject, report['remaining']))
            self.assertEqual(report['checked_pages'], 8)


class RestyleTest(unittest.TestCase):
    def fake_job(self, root, grade=8):
        """A finished job folder written by hand (the restyle code only needs the row and the saved pages)."""
        store = Store(root)
        import uuid, time as _t
        jid = uuid.uuid4().hex
        request = {'unit': '角度探險', 'grade': grade, 'semester': '上學期', 'subject': '數學', 'slide_count': 8}
        with store.db() as c:
            c.execute('INSERT INTO jobs(id,request,status,stage,attempts,error,created,updated) VALUES(?,?,?,?,?,?,?,?)',
                      (jid, json.dumps(request, ensure_ascii=False), 'done', 7, 0, '', _t.time(), _t.time()))
        out = store.folder(jid)
        (out / 'slides').mkdir(parents=True)
        deck = junior_deck(grade)
        atomic(out / 'plan.json', {'slides': [{'heading': s['heading']} for s in deck['slides']]})
        for number, slide in enumerate(deck['slides'], 1):
            atomic(out / 'slides' / f'{number:03}.json', slide)
        (out / '教材.pptx').write_bytes(b'old deck')
        (out / 'previews').mkdir()
        (out / 'previews' / '001.png').write_bytes(b'old preview')
        return store, request, jid, out

    def test_old_deck_is_listed_then_restyled_without_a_model_and_marked_current(self):
        from workflow_v2 import pipeline
        with tempfile.TemporaryDirectory() as t:
            store, request, jid, out = self.fake_job(Path(t) / 'data')
            self.assertEqual([j for j, _ in restyle.candidates(store)], [jid])
            self.assertFalse(restyle.is_current(request, out))
            with patch.object(pipeline, 'preview_stage') as preview, patch.object(pipeline, 'log_event'), \
                    patch('workflow_v2.model.call', side_effect=AssertionError('no model call allowed')), \
                    patch.object(pipeline, 'generate_assets', return_value={}):
                result = restyle.restyle(jid, store)
            preview.assert_called_once()
            self.assertEqual(result['pages'], 8)
            self.assertTrue(restyle.is_current(request, out))
            self.assertEqual(restyle.candidates(store), [])
            self.assertEqual(len(Presentation(out / '教材.pptx').slides), 8)
            backup = Path(result['backup'])
            self.assertEqual((backup / '教材.pptx').read_bytes(), b'old deck')
            self.assertEqual((backup / 'previews' / '001.png').read_bytes(), b'old preview')
            self.assertIn('備份', result['message'])

    def test_a_failure_puts_the_old_deck_back_with_a_plain_message(self):
        from workflow_v2 import pipeline
        with tempfile.TemporaryDirectory() as t:
            store, request, jid, out = self.fake_job(Path(t) / 'data')
            with patch.object(pipeline, 'rendering', side_effect=ValueError('LibreOffice 壞了')), patch.object(pipeline, 'log_event'):
                with self.assertRaises(ValueError) as caught:
                    restyle.restyle(jid, store)
            self.assertIn('已還原原本的簡報', str(caught.exception))
            self.assertEqual((out / '教材.pptx').read_bytes(), b'old deck')
            self.assertFalse(restyle.is_current(request, out))

    def test_primary_school_unfinished_and_content_less_decks_are_never_offered(self):
        with tempfile.TemporaryDirectory() as t:
            store, _, jid, out = self.fake_job(Path(t) / 'data', grade=5)
            self.assertEqual(restyle.candidates(store), [])
            with self.assertRaises(ValueError):
                restyle.restyle('0' * 32, store)
            store2, _, jid2, out2 = self.fake_job(Path(t) / 'data2')
            for f in (out2 / 'slides').glob('*.json'):
                f.unlink()
            self.assertEqual(restyle.candidates(store2), [])
            with self.assertRaises(ValueError) as caught:
                restyle.restyle(jid2, store2)
            self.assertIn('重新製作', str(caught.exception))

    def test_new_decks_record_their_design_so_they_are_never_offered(self):
        source = (ROOT / 'workflow_v2' / 'pipeline.py').read_text(encoding='utf-8')
        self.assertIn("atomic(out/'design.json',design_mark(request))", source)
        self.assertEqual(restyle.design_mark({'grade': 8, 'subject': '數學'}), {'look': 'mission', 'rev': ts.MISSION['rev'], 'band': 'junior'})
        self.assertEqual(restyle.design_mark({'grade': 4, 'subject': '數學'})['look'], 'story')

    def test_console_offers_the_button_only_in_the_my_decks_tab_text(self):
        source = (ROOT / 'workflow_v2' / 'panel.py').read_text(encoding='utf-8')
        self.assertIn('套用新版面', source)
        self.assertIn('restyle.restyle(row[\'id\'], store)', source)


if __name__ == '__main__':
    unittest.main()
