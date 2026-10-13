"""Local image decision, embedding and selective runtime-log clearing."""
import base64
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from PIL import Image
from pptx import Presentation

from tests.test_design_system import sample_deck
from workflow_v2.contracts import VISUAL_STYLE, validate_slide
from workflow_v2.local_images import endpoint, generate_assets
from workflow_v2.renderer import render
from workflow_v2 import runlog


class LocalImagesTest(unittest.TestCase):
    def test_visual_decision_must_match_safe_layout(self):
        page=sample_deck()['slides'][6]
        page['visual_asset'].update(is_required=True,asset_type='illustration',
                                    image_prompt='A small compass on a desk')
        with self.assertRaises(ValueError):validate_slide(page,{1})

    def test_endpoint_rejects_remote_service(self):
        with patch('workflow_v2.local_images.yaml.safe_load',return_value={
            'image_generation_url':'https://example.com/sdapi/v1/txt2img'}):
            with self.assertRaises(ValueError):endpoint()

    def test_generation_embeds_image_and_preserves_logical_text(self):
        deck=sample_deck()
        page=deck['slides'][2]
        page['visual_asset']={'is_required':True,'reasoning':'用開口大小的生活情境幫助理解',
                              'asset_type':'illustration','image_prompt':'Two open doors at different angles',
                              'style_tags':VISUAL_STYLE}
        image=Image.new('RGB',(512,512),'#008080')
        buffer=io.BytesIO();image.save(buffer,format='PNG')
        response=Mock();response.json.return_value={'images':[base64.b64encode(buffer.getvalue()).decode()]}
        response.raise_for_status.return_value=None
        session=Mock();session.get.return_value=Mock(status_code=200,
            json=lambda:{'sd_model_checkpoint':'local-test-model'})
        session.post.return_value=response
        with tempfile.TemporaryDirectory() as folder, \
                patch('workflow_v2.local_images.gpu_slot') as slot:
            from contextlib import nullcontext
            slot.return_value=nullcontext()
            root=Path(folder)
            media=generate_assets(deck,root,session=session)
            self.assertIn(3,media)
            self.assertTrue(Path(media[3]).is_file())
            self.assertEqual(session.post.call_count,1)
            pptx=root/'deck.pptx'
            render(deck,pptx,media={**media,1:media[3]})
            prs=Presentation(pptx)
            self.assertTrue(any(shape.shape_type==13 for shape in prs.slides[0].shapes))
            self.assertTrue(any(shape.shape_type==13 for shape in prs.slides[2].shapes))
            self.assertIn('角的大小',prs.slides[2].shapes.title.text)
            manifest=json.loads((root/'visual-assets.json').read_text(encoding='utf-8'))
            self.assertEqual(manifest['3']['status'],'generated')
            media2=generate_assets(deck,root,session=session)
            self.assertEqual(session.post.call_count,1)
            self.assertEqual(media2[3],media[3])

    def test_unavailable_service_keeps_deck_renderable(self):
        deck=sample_deck()
        deck['slides'][0]['visual_asset'].update(is_required=True,asset_type='illustration',
            image_prompt='A cheerful classroom with colorful learning tools')
        session=Mock();session.get.side_effect=__import__('requests').ConnectionError('offline')
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder)
            self.assertEqual(generate_assets(deck,root,session=session),{})
            self.assertEqual(json.loads((root/'visual-assets.json').read_text(encoding='utf-8'))['1']['status'],'unavailable')
            render(deck,root/'deck.pptx')
            self.assertTrue((root/'deck.pptx').is_file())

    def test_clear_selected_log_only(self):
        jid='a'*32
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);job=root/'jobs'/jid;job.mkdir(parents=True)
            system=root/'system.jsonl';system.write_text('system\n',encoding='utf-8')
            log=job/'run.jsonl';log.write_text('job\n',encoding='utf-8')
            with patch.object(runlog,'DATA',root),patch.object(runlog,'SYSTEM_LOG',system):
                self.assertGreater(runlog.clear_log(jid),0)
                self.assertEqual(log.read_text(encoding='utf-8'),'')
                self.assertEqual(system.read_text(encoding='utf-8'),'system\n')
                self.assertGreater(runlog.clear_log(),0)
                self.assertEqual(system.read_text(encoding='utf-8'),'')


if __name__=='__main__':unittest.main()
