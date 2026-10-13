"""Regressions reproduced from the October 2 production mathematics deck."""
import copy
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock,patch
import pymupdf
from pptx import Presentation
from workflow_v2 import pipeline,qa
from workflow_v2.contracts import validate_slide
from workflow_v2.design import expand_staged
from workflow_v2.renderer import paginate_deck,render
from workflow_v2.service import regenerate
from workflow_v2.source_reader import damaged_text,recover_pdf_page
from workflow_v2.store import Store,atomic
from workflow_v2.visuals import sanitize
from workflow_v2.worksheet import _evidence,_make_draft,_schema
from tests.test_workflow_v2 import REQUEST,slide
from tests.test_design_system import sample_deck

class AuditRegressions(unittest.TestCase):
    def test_negative_and_parenthesized_arithmetic_is_checked_exactly(self):
        for text in ('2-5=-3','(-4)-(-7)=3','-5-11=-16','7+(-3)=4','-(250+180)=-430','1,000-1,005=-5'):
            self.assertEqual(qa.check_arithmetic(text),[],text)
        for text in ('2-5=-2','-5-11=-15','(-4)-(-7)=-3','-(250+180)=430'):
            self.assertTrue(qa.check_arithmetic(text),text)
        self.assertEqual(qa.check_arithmetic('x+5=12'),[])
        self.assertEqual(qa.check_arithmetic('6:4=3:2'),[])
        self.assertEqual(qa.check_arithmetic('50÷3=16 又 2/3'),[])
    def test_negative_results_and_short_division_do_not_trigger_rescue(self):
        self.assertEqual(qa.check_traceable('A(-5)、B(11)的距離',['-5-11=-16','距離是16']),[])
        self.assertEqual(qa.check_traceable('48做質因數分解',['48÷2=24','24÷2=12','12÷2=6','6÷2=3']),[])
        self.assertTrue(qa.check_traceable('2:3的比',['未知的12÷2=6']))
    def test_exit_ticket_keeps_both_answers_in_export(self):
        deck=sample_deck();summary=deck['slides'][-1]
        summary['body'].update(question='A=-3、B=7。(1)距離？(2)B左邊5單位的坐標？',
                               answer='(1)距離=10。\n(2)坐標=7-5=2。')
        expanded=expand_staged(deck)
        self.assertEqual(expanded['slides'][-1]['body']['answer'],summary['body']['answer'])
        with tempfile.TemporaryDirectory() as t:
            path=Path(t)/'answers.pptx';render(expanded,path)
            text='\n'.join(s.text for s in Presentation(path).slides[-1].shapes if s.has_text_frame)
            self.assertIn('(1)距離=10',text);self.assertIn('(2)坐標=7-5=2',text)
    def test_worked_solution_is_never_turned_into_a_hint_automatically(self):
        page=sample_deck()['slides'][6];page['body']['answer']='30÷3=10。\n答案：10。'
        expanded=expand_staged({'slides':[page]})['slides']
        self.assertNotIn('hint',[p['stage'] for p in expanded])
        self.assertEqual(expanded[0]['body']['answer'],'')
    def test_four_steps_render_in_order_with_complete_problem(self):
        deck=sample_deck();page=deck['slides'][5]
        page['body']['question']='48做質因數分解，逐次除以2。'
        page['modules']=[{'keyword':'短除','description':text} for text in
                         ('48÷2=24','24÷2=12','12÷2=6','6÷2=3')]
        validate_slide(page,[1]);expanded=paginate_deck(deck)
        parts=[p for p in expanded['slides'] if p['layout_id']=='Steps']
        self.assertEqual([m['description'] for p in parts for m in p['modules']],
                         [m['description'] for m in page['modules']])
        self.assertEqual([p['step_offset'] for p in parts],[0,2])
        with tempfile.TemporaryDirectory() as t:
            path=Path(t)/'steps.pptx';render(expanded,path)
            texts=['\n'.join(s.text for s in p.shapes if s.has_text_frame) for p in Presentation(path).slides]
            self.assertEqual(sum(page['body']['question'] in text for text in texts),2)
            self.assertIn('6÷2=3','\n'.join(texts))
    def test_number_line_is_editable_and_positions_are_proportional(self):
        vm={'type':'number_line','minimum':-6,'maximum':12,'points':[{'label':'A','value':-5},{'label':'B','value':11}]}
        clean,reason=sanitize(vm,'Text_Visual');self.assertEqual(reason,'')
        deck=sample_deck();page=deck['slides'][2];page['visual_model']=clean
        page['body']['bullets']=['A(-5)、B(11)，從左向右讀取坐標。']
        with tempfile.TemporaryDirectory() as t:
            path=Path(t)/'line.pptx';render(deck,path)
            shapes=Presentation(path).slides[2].shapes
            labels=[s for s in shapes if s.has_text_frame and s.text in ('A\n-5','B\n11')]
            self.assertEqual(len(labels),2);self.assertLess(labels[0].left,labels[1].left)
            self.assertFalse(any(s.shape_type==13 for s in shapes))
        self.assertIsNone(sanitize(dict(vm,maximum=-7),'Text_Visual')[0])
    def test_changed_topic_invalidates_page_without_reusing_old_model_reply(self):
        with tempfile.TemporaryDirectory() as t:
            out=Path(t);item={'layout_id':'bullet_with_image','topic':'舊主題','phase':'concept','source_pages':[1]}
            atomic(out/'source.json',{'pages':[{'page':1,'text':'主教材。'}]})
            atomic(out/'plan.json',{'objectives':['舊目標'],'slides':[item]})
            atomic(out/'slides/001.json',slide())
            atomic(out/'page-policies/001.json',{'fingerprint':pipeline.page_fingerprint(REQUEST,out,0,item),'started':0})
            changed=dict(item,topic='新主題')
            atomic(out/'plan.json',{'objectives':['新目標'],'slides':[changed]})
            with patch.object(pipeline,'recover_page',return_value=False),patch.object(pipeline,'author_page') as author:
                pipeline.write_one(REQUEST,out,0,changed)
            author.assert_called_once()
            self.assertEqual(len(list((out/'invalidated-pages').glob('*.json'))),1)
    def test_revision_preserves_template_and_other_page_warnings(self):
        with tempfile.TemporaryDirectory() as t:
            store=Store(t);parent=store.submit(REQUEST);out=store.folder(parent)
            atomic(out/'source.json',{'pages':[]})
            atomic(out/'plan.json',{'slides':[{},{}]})
            atomic(out/'slides/001.json',slide());atomic(out/'slides/002.json',slide())
            atomic(out/'template-selection.json',{'template_id':'g6-math-practice'})
            atomic(out/'quality-warnings.json',{'1':{'reason':'保留警示'},'2':{'reason':'重製此頁'}})
            with store.db() as c:c.execute("UPDATE jobs SET status='done' WHERE id=?",(parent,))
            child=store.folder(regenerate(parent,1,'重製',store))
            self.assertEqual(pipeline.read(child/'template-selection.json'),pipeline.read(out/'template-selection.json'))
            self.assertEqual(set(pipeline.read(child/'quality-warnings.json')),{'1'})
            self.assertTrue((child/'page-policies/001.json').exists())
            self.assertFalse((child/'page-policies/002.json').exists())
    def test_pdf_font_map_corruption_uses_local_vision_and_page_cache(self):
        self.assertTrue(damaged_text('炼ĲĶı ғөġ 參加拔河。'))
        self.assertFalse(damaged_text('距離=16，α是角度，cat是英文。'))
        with tempfile.TemporaryDirectory() as t,pymupdf.open() as pdf:
            page=pdf.new_page();page.insert_text((50,50),'2-5=-3')
            model=Mock(return_value={'text':'2-5=-3','complete':True,'issues':[]})
            out=Path(t)
            self.assertEqual(recover_pdf_page(page,1,REQUEST,out,'hash',model),'2-5=-3')
            self.assertTrue(model.call_args.kwargs['images'])
            recover_pdf_page(page,1,REQUEST,out,'hash',model);model.assert_called_once()
    def test_fallback_does_not_show_corrupt_source_or_cut_math_title(self):
        with tempfile.TemporaryDirectory() as t:
            out=Path(t);heading='例題：求A(-5)、B(11)的距離'
            atomic(out/'plan.json',{'slides':[{'topic':heading,'source_pages':[1]}]})
            atomic(out/'source.json',{'pages':[{'page':1,'text':'炼ĲĶı ғөġ 無法讀取。'}]})
            page=pipeline.fallback_page(REQUEST,out,0,'需確認')
            self.assertEqual(page['heading'],heading)
            self.assertNotIn('Ĳ',page['body']['bullets'][0])
            self.assertIn('無法可靠辨識',page['body']['bullets'][0])
    def test_worksheet_evidence_preserves_conditions_after_character_180(self):
        with tempfile.TemporaryDirectory() as t:
            out=Path(t);page=slide();text='完整必要條件。'*35+'最後條件：共20元。'
            page['body'].update(question=text,answer='解答。'*70+'答案：20元。')
            atomic(out/'source.json',{'pages':[{'page':1,'text':text*3}]})
            atomic(out/'plan.json',{'slides':[{'topic':'完整題幹','difficulty':'低難度','phase':'example','source_pages':[1]}]})
            atomic(out/'reviewed-deck.json',{'slides':[page]})
            value,pages=_evidence(out)
            self.assertEqual(value['slides'][0]['question'],text)
            self.assertEqual(value['slides'][0]['answer'],page['body']['answer'])
            self.assertEqual(value['source_excerpts'][0]['text'],text*3)
    def test_answer_batch_requires_exact_unique_ids(self):
        with tempfile.TemporaryDirectory() as t:
            stems={'instructions':'請獨立作答。','questions':[{'difficulty':'低難度','question':f'{i}加2是多少？','source_pages':[1]} for i in range(6)]}
            duplicates={'answers':[{'id':1,'answer':'2','explanation':'把兩個數相加。'}]*3}
            with patch('workflow_v2.worksheet.call',side_effect=[stems,duplicates]):
                with self.assertRaisesRegex(ValueError,'題號重複或缺漏'):
                    _make_draft(REQUEST,Path(t),'prompt',_schema({'低難度':6},{1}),{'source_excerpts':[]})
    def test_first_pass_count_excludes_repaired_pages(self):
        summary=qa.summarize({'1':{'passed':True,'final':'passed'},'2':{'passed':True,'final':'regenerated'},'3':{'passed':False,'final':'fallback'}})
        self.assertEqual(summary['passed'],2);self.assertEqual(summary['first_passed'],1)

if __name__=='__main__':unittest.main()
