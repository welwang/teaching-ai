"""品牌素材、現行頁首與監控刷新測試。"""
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from workflow_v2 import brand, beautify, console_state as cs
from tests.test_monitoring_report import REQUEST
from workflow_v2.store import Store

ROOT = Path(__file__).resolve().parents[1]


def luminance(code):
    r, g, b = [int(code.lstrip('#')[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    f = lambda c: c / 12.92 if c <= .03928 else ((c + .055) / 1.055) ** 2.4
    return .2126 * f(r) + .7152 * f(g) + .0722 * f(b)


def contrast(a, b):
    hi, lo = sorted([luminance(a), luminance(b)], reverse=True)
    return (hi + .05) / (lo + .05)


class LogoTest(unittest.TestCase):
    def test_every_size_exists_and_is_square_with_rounded_transparent_corners(self):
        for size in (512, 192, 64, 48):
            image = Image.open(brand.logo_path(size))
            self.assertEqual(image.size, (size, size))
            self.assertEqual(image.mode, 'RGBA')
            self.assertEqual(image.getpixel((0, 0))[3], 0, size)
            self.assertEqual(image.getpixel((size // 2, size // 2))[3], 255)
        self.assertTrue(brand.logo_path('original'))
        icon = Image.open(brand.logo_path('ico'))
        self.assertEqual(icon.format, 'ICO')

    def test_console_bar_shows_exactly_one_inline_logo(self):
        uri = brand.data_uri(48)
        self.assertTrue(uri.startswith('data:image/png;base64,'))
        html = self.header_html()
        self.assertEqual(html.count('<img '), 1)
        self.assertIn(uri, html)

    @staticmethod
    def header_html():
        from streamlit.testing.v1 import AppTest

        def app():
            import streamlit as st
            from workflow_v2.panel import _header
            st.session_state['pn-prefs'] = {'theme': '白色', 'size': '中', 'calm': False}
            _header()

        at = AppTest.from_function(app).run()
        if at.exception:
            raise AssertionError([e.value for e in at.exception])
        return next(item.value for item in at.markdown if 'pn-brand' in item.value)

    def test_missing_logo_files_never_break_the_bar_or_the_tab_icon(self):
        with patch.object(brand, 'BRAND', Path(tempfile.gettempdir()) / 'no-such-brand-folder'):
            self.assertEqual(brand.data_uri(48), '')
            self.assertIsNone(brand.page_icon())
            self.assertNotIn('<img ', self.header_html())

    def test_browser_tab_uses_the_logo(self):
        self.assertIn('page_icon=brand.page_icon()', (ROOT / 'app_v2.py').read_text(encoding='utf-8'))
        self.assertEqual(brand.page_icon().size, (192, 192))

    def test_slide_logo_is_off_by_default_and_only_on_the_cover_when_switched_on(self):
        self.assertEqual(brand.slide_logo_mode({}), 'off')
        self.assertEqual(brand.slide_logo_mode({'slide_logo': 'banana'}), 'off')
        self.assertEqual(brand.slide_logo_mode({'slide_logo': 'Cover'}), 'cover')
        self.assertEqual(brand.slide_logo_mode(), 'off')
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self.assertIsNone(beautify.brand_mark(slide))
        with patch.object(brand, 'slide_logo_mode', return_value='cover'):
            picture = beautify.brand_mark(slide)
        self.assertEqual(picture.name, '標誌')
        self.assertLess(picture.left + picture.width, Inches(3))
        self.assertGreater(picture.top, Inches(6.9))
        self.assertEqual(len([s for s in slide.shapes if s.name == '標誌']), 1)

    def test_installer_ships_and_uses_the_logo(self):
        build = (ROOT / 'build_installer.ps1').read_text(encoding='utf-8')
        self.assertIn('assets\\brand', build)
        install = (ROOT / 'installer' / 'install.ps1').read_text(encoding='utf-8')
        self.assertIn('favicon.ico', install)
        self.assertIn('logo-64.png', install)


class RefreshSettingTest(unittest.TestCase):
    def test_default_is_one_second_and_bad_values_are_safe(self):
        self.assertEqual(cs.refresh_seconds({}), 1)
        self.assertEqual(cs.refresh_seconds({'console_refresh_seconds': 3}), 3)
        self.assertEqual(cs.refresh_seconds({'console_refresh_seconds': 0}), 1)
        self.assertEqual(cs.refresh_seconds({'console_refresh_seconds': 0.2}), 1)
        self.assertEqual(cs.refresh_seconds({'console_refresh_seconds': 999}), 60)
        for bad in ('快', None, float('nan'), [2]):
            self.assertEqual(cs.refresh_seconds({'console_refresh_seconds': bad}), 1, bad)

    def test_shipped_config_explains_the_setting_in_chinese_and_uses_one_second(self):
        text = (ROOT / 'config.yaml').read_text(encoding='utf-8')
        self.assertRegex(text, r'(?m)^console_refresh_seconds: 1$')
        before = text.split('console_refresh_seconds:')[0].splitlines()[-2:]
        self.assertTrue(all(line.startswith('#') for line in before), before)
        self.assertRegex(''.join(before), r'[\u4e00-\u9fff]')
        self.assertEqual(cs.refresh_seconds(), 1)

    def test_only_the_live_fragment_runs_every_interval_and_the_page_is_not_rerun(self):
        from workflow_v2 import panel
        source = (ROOT / 'workflow_v2' / 'panel.py').read_text(encoding='utf-8')
        self.assertIn("@st.fragment(run_every=f'{REFRESH}s')", source)
        self.assertEqual(panel.REFRESH, cs.refresh_seconds())
        self.assertEqual(source.count("run_every=f'{REFRESH}s'"), 1)
        self.assertIn("st.rerun(scope='fragment')", source)

    def test_one_refresh_reads_few_files_and_is_cheap_on_a_running_job(self):
        import time
        from workflow_v2.monitoring import snapshot
        with tempfile.TemporaryDirectory() as t:
            store = Store(Path(t) / 'data')
            store.submit(REQUEST)
            with patch('workflow_v2.monitoring.gpu', return_value=None), \
                    patch('workflow_v2.monitoring._worker_alive', return_value=True):
                snapshot(store)
                start = time.perf_counter()
                for _ in range(10):
                    snapshot(store)
                each = (time.perf_counter() - start) / 10
        self.assertLess(each, .25, '每次更新應遠小於 1 秒')


if __name__ == '__main__':
    unittest.main()
