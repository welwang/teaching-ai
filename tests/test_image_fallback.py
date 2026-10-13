import copy,json,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
from PIL import Image
from pptx import Presentation
from workflow_v2 import images
from workflow_v2.image_fallback import apply_missing_visuals
from workflow_v2.renderer import render,safe_fallback_page
from workflow_v2.store import atomic
from tests.test_design_system import sample_deck

class ImageFallbackTest(unittest.TestCase):
    def page(self):
        page=sample_deck()['slides'][2]
        page['heading']='球在哪裡';page['body']['bullets']=['The ball is under the box.']
        page['visual_asset'].update(is_required=True,asset_type='illustration',image_prompt='A ball under a box')
        page['media_hint']['search_keyword']='ball box illustration'
        page['fallback_visual_model']={'type':'spatial_scene','relation':'under','object_label':'ball','references':['box']}
        return page

    def test_queries_keep_subject_and_try_a_known_synonym(self):
        values=images.search_variants('cute hippo illustration')
        self.assertEqual(values,['cute hippo illustration','hippo','hippopotamus'])
        self.assertLessEqual(len(images.search_variants('giant panda illustration')),3)
        self.assertTrue(images.relevant({'title':'Apple fruit','keywords':'red apples'},'apple fruit'))
        self.assertFalse(images.relevant({'title':'Orange fruit','keywords':'fruit'},'apple fruit'))

    def test_failed_first_query_can_download_the_subject_synonym(self):
        page=self.page();page['media_hint']['search_keyword']='hippo'
        candidate={'title':'Hippopotamus','url':'https://upload.wikimedia.org/hippo.png','license':'CC0 1.0',
            'license_url':images.CC0_URL,'description_url':'https://commons.wikimedia.org/wiki/Hippo'}
        with tempfile.TemporaryDirectory() as folder:
            out=Path(folder);picture=out/'test.png';Image.new('RGB',(400,400),'teal').save(picture)
            with patch.object(images,'find_image',side_effect=[None,candidate]) as lookup,patch.object(images,'download',return_value=picture):
                media=images.online_images({'slides':[page]},out)
            self.assertIn(1,media);self.assertEqual([c.args[0] for c in lookup.call_args_list],['hippo','hippopotamus'])
            self.assertEqual(json.loads((out/'image-sources.json').read_text(encoding='utf-8'))['1']['resolved_query'],'hippopotamus')

    def test_deadline_stops_new_network_calls_and_failure_is_remembered(self):
        with tempfile.TemporaryDirectory() as folder:
            failed=set();out=Path(folder)
            with patch.object(images,'find_image') as lookup:
                images.online_images({'slides':[self.page()]},out,failed_keywords=failed,search_budget_seconds=0)
                images.online_images({'slides':[self.page()]},out,failed_keywords=failed)
            lookup.assert_not_called()
            self.assertIn('ball box illustration',failed)
            self.assertEqual(json.loads((out/'visual-assets.json').read_text(encoding='utf-8'))['1']['status'],'unavailable')

    def test_corrupt_cached_file_is_not_embedded(self):
        with tempfile.TemporaryDirectory() as folder:
            out=Path(folder);bad=out/'bad.png';bad.write_bytes(b'not an image')
            atomic(out/'image-sources.json',{'1':{'search_keyword':'ball box illustration','local_path':str(bad)}})
            with patch.object(images,'find_image',return_value=None):
                self.assertEqual(images.online_images({'slides':[self.page()]},out),{})
            self.assertFalse(images.valid_image(bad))

    def test_missing_picture_uses_the_model_authored_editable_diagram(self):
        deck=sample_deck();deck['slides'][2]=self.page()
        with tempfile.TemporaryDirectory() as folder:
            out=Path(folder);apply_missing_visuals(deck,out,{})
            page=deck['slides'][2]
            self.assertEqual(page['visual_status'],'diagram_fallback')
            self.assertFalse(page['visual_asset']['is_required'])
            self.assertEqual(page['visual_model']['relation'],'under')
            path=out/'fallback.pptx';report=render(deck,path)
            self.assertEqual(report[2]['visual_model'],'spatial_scene')
            self.assertFalse(any(s.shape_type==13 for s in Presentation(path).slides[2].shapes))
            self.assertEqual(json.loads((out/'visual-assets.json').read_text(encoding='utf-8'))['3']['status'],'diagram_fallback')

    def test_fallback_does_not_invent_numbers_or_claim_a_missing_photo_exists(self):
        deck=sample_deck();deck['slides'][2]=self.page()
        deck['slides'][2]['fallback_visual_model']={'type':'number_line','minimum':0,'maximum':100,
            'points':[{'label':'A','value':97}]}
        with tempfile.TemporaryDirectory() as folder:
            out=Path(folder);apply_missing_visuals(deck,out,{})
            self.assertEqual(deck['slides'][2]['visual_status'],'text_only')
            self.assertNotIn('visual_model',deck['slides'][2])
            record=json.loads((out/'visual-assets.json').read_text(encoding='utf-8'))['3']
            self.assertEqual(record['status'],'unavailable')
            self.assertIn('沒有支持的數據',record['reason'])

    def test_overflow_retains_all_venn_classification_text(self):
        deck=sample_deck();page=deck['slides'][2]
        page['visual_model']={'type':'venn_diagram','left':{'title':'甲','items':['左側特徵']},
            'right':{'title':'乙','items':['右側特徵']},'intersection':['共同特徵']}
        result=safe_fallback_page(deck,2)
        text='\n'.join('\n'.join(p['body']['bullets']) for p in result['slides'])
        for value in ('甲：左側特徵','乙：右側特徵','共同點：共同特徵'):self.assertIn(value,text)


if __name__=='__main__':unittest.main()
