"""Student-facing structure and editable diagrams, without a model/network."""
import copy
import tempfile
import unittest
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from workflow_v2 import pipeline,qa
from workflow_v2.contracts import validate_slide
from workflow_v2.design import expand_staged
from workflow_v2.renderer import render
from workflow_v2.source_reader import separate_context
from workflow_v2.text_safety import format_teaching_text
from workflow_v2.visuals import sanitize
from tests.test_design_system import sample_deck


class StudentContentTest(unittest.TestCase):
    def question(self):
        page=sample_deck()['slides'][6]
        page['heading']='Where Is the Panda?'
        page['reading_context']='The panda is under the desk. The gift is on the desk.'
        page['quiz']={'questions':[{'id':3,'question':'Where is the gift?'},
            {'id':1,'question':'Where is the panda?'},{'id':2,'question':'Is the gift under the desk?'}],
            'answers':[{'id':2,'answer':'No, it is on the desk.'},
            {'id':3,'answer':'It is on the desk.'},{'id':1,'answer':'It is under the desk.'}]}
        return pipeline.normalize_slide(page)

    def test_metadata_blocks_are_separate_but_student_strategy_stays(self):
        text='Find Details: Read the sentence.\n\n搭配 R1 Find Details 排序\n教學重點：提醒教師先示範。\n\n學生回答：The gift is on the desk.'
        student,teacher=separate_context(text)
        self.assertIn('Find Details: Read the sentence.',student)
        self.assertIn('學生回答',student)
        self.assertNotIn('搭配 R1',student);self.assertNotIn('教學重點',student)
        self.assertIn('教學重點',teacher)

    def test_layout_blocks_move_sidebar_and_footer_without_losing_paragraph(self):
        blocks=[(30,60,400,160,'Read: The panda is under the desk.',0,0),
            (420,60,540,160,'教學補充\n先讓學生指出桌子的位置。',1,0),
            (260,790,280,800,'12',2,0)]
        student,teacher=separate_context('',blocks,820)
        self.assertEqual(student,'Read: The panda is under the desk.')
        self.assertIn('教學補充',teacher);self.assertIn('12',teacher)

    def test_teacher_leak_and_unclosed_titles_are_hard_errors(self):
        page=sample_deck()['slides'][2];page['body']['bullets']=['教學補充：請教師帶學生閱讀。']
        self.assertTrue(any(f['code']=='teacher_context_leak' for f in qa.check_residue(page)))
        page['heading']='哪句話提到「fruit'
        self.assertTrue(any(f['code']=='truncated' and f['field']=='heading' and f['severity']=='major' for f in qa.check_residue(page)))

    def test_questions_and_answers_pair_by_id_then_expand_only_once(self):
        page=self.question();validate_slide(page,[1])
        self.assertEqual([q['id'] for q in page['quiz']['questions']],[1,2,3])
        pair=expand_staged({'slides':[page]})
        self.assertEqual([p['stage'] for p in pair['slides']],['question','answer'])
        self.assertEqual(expand_staged(pair),pair)
        self.assertIn(page['reading_context'],pair['slides'][0]['body']['question'])
        self.assertNotIn(page['reading_context'],pair['slides'][1]['body']['question'])
        for index in (1,2,3):self.assertIn(f'({index})',pair['slides'][1]['body']['answer'])

    def test_missing_or_duplicate_answers_are_never_silently_accepted(self):
        page=self.question();page['quiz']['answers'].pop()
        with self.assertRaisesRegex(ValueError,'答案缺漏'):validate_slide(page,[1])
        with self.assertRaisesRegex(ValueError,'答案缺漏'):pipeline.normalize_slide(page)
        page=self.question();page['quiz']['questions'][1]['id']=1
        with self.assertRaisesRegex(ValueError,'題號重複'):validate_slide(page,[1])

    def test_full_english_heading_and_all_answers_reach_native_pptx(self):
        page=self.question();page['heading']='Relative Positions: On, In and Under'
        page['key_sentence']='';page.pop('visual_model',None)
        deck=sample_deck();deck['slides'][6]=page
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'english.pptx';render(expand_staged(deck),path)
            slides=Presentation(path).slides
            texts=['\n'.join(s.text for s in p.shapes if s.has_text_frame) for p in slides]
            self.assertEqual(sum(page['reading_context'] in text for text in texts),1)
            answers=next(t for t in texts if '(3) It is on the desk.' in t)
            self.assertIn('(1) It is under the desk.',answers)
            self.assertIn('(2) No, it is on the desk.',answers)
            self.assertIn(page['heading'],answers)

    def test_venn_is_editable_has_three_regions_and_never_invents_membership(self):
        vm={'type':'venn_diagram','title':'兩者異同','left':{'title':'Cat','items':['四隻腳']},
            'right':{'title':'Bird','items':['兩隻腳']},'intersection':['有眼睛']}
        clean,reason=sanitize(vm,'Text_Visual');self.assertEqual(reason,'')
        deck=sample_deck();page=deck['slides'][2];page['heading']='比較動物';page['visual_model']=clean
        page['body']['bullets']=['左圈是貓，右圈是鳥，中間是共同點。']
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'venn.pptx';render(deck,path)
            shapes=Presentation(path).slides[2].shapes
            circles=[s for s in shapes if s.shape_type==1 and s.auto_shape_type==MSO_SHAPE.OVAL]
            self.assertEqual(len(circles),2)
            self.assertLess(circles[1].left,circles[0].left+circles[0].width)
            text=' '.join(s.text for s in shapes if s.has_text_frame)
            self.assertTrue(all(t in text for t in ('四隻腳','兩隻腳','有眼睛','共同點')))
        bad=copy.deepcopy(vm);bad['intersection']=['四隻腳']
        self.assertIsNone(sanitize(bad,'Text_Visual')[0])

    def test_preposition_diagrams_use_fixed_geometry_and_complete_labels(self):
        for relation in ('on','in','under','next_to','behind','in_front_of','between'):
            vm={'type':'spatial_scene','relation':relation,'object_label':'ball',
                'references':['box','box'] if relation=='between' else ['box']}
            clean,reason=sanitize(vm,'Text_Visual');self.assertEqual(reason,'',relation)
            deck=sample_deck();deck['slides'][2]['visual_model']=clean
            deck['slides'][2]['body']['bullets']=['The ball is '+relation.replace('_',' ')+' the box.']
            with tempfile.TemporaryDirectory() as folder:
                path=Path(folder)/'position.pptx';report=render(deck,path)
                self.assertEqual(report[2]['visual_model'],'spatial_scene')
                shapes=Presentation(path).slides[2].shapes
                for s in shapes:
                    self.assertGreaterEqual(s.left,0);self.assertGreaterEqual(s.top,0)
                    self.assertLessEqual(s.left+s.width,12191621)
                self.assertFalse(any(s.shape_type==13 for s in shapes))

    def test_office_residue_is_converted_without_deleting_english_words(self):
        self.assertEqual(format_teaching_text(r'\(B->(->A)\)'), 'B→(→A)')
        self.assertEqual(format_teaching_text('18_x000D_ightarrow k=3'), '18 →  k=3')
        self.assertEqual(format_teaching_text('three times a day'), 'three times a day')


if __name__=='__main__':unittest.main()
