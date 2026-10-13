"""Editing one page of a finished deck by hand (workflow_v2/page_edit.py)."""
import json
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation
from pptx.util import Inches, Pt

from workflow_v2 import page_edit
from workflow_v2.store import Store

REQUEST = {'unit': '因數與倍數', 'grade': 6, 'subject': '數學', 'semester': '上學期', 'slide_count': 2}


def make_deck(path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for title, body in (('認識因數', '12 的因數有 1、2、3、4、6、12'), ('找最大公因數', '先寫出兩數的因數')):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        head = slide.shapes.add_textbox(Inches(1), Inches(.5), Inches(10), Inches(1))
        head.text_frame.text = title
        head.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(10), Inches(3))
        box.text_frame.text = body
        foot = slide.shapes.add_textbox(Inches(1), Inches(7.1), Inches(5), Inches(.3))
        foot.text_frame.text = '圖片：作者未標示'
        slide.notes_slide.notes_text_frame.text = '先問學生 12 可以怎麼分。'
    prs.save(path)


def geometry(path, page):
    slide = Presentation(path).slides[page - 1]
    return [(s.shape_id, s.left, s.top, s.width, s.height) for s in slide.shapes]


class PageEditTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.store = Store(Path(self.temp.name))
        self.jid = 'b' * 31 + 'c'
        with self.store.db() as c:
            c.execute('INSERT INTO jobs VALUES(?,?,?,?,?,?,?,?)', (self.jid, json.dumps(REQUEST, ensure_ascii=False), 'done', 7, 0, '', time.time(), time.time()))
        self.out = self.store.folder(self.jid)
        self.out.mkdir(parents=True)
        make_deck(self.out / '教材.pptx')
        (self.out / '教材.pdf').write_bytes(b'old pdf')
        deck = {'slides': [{'heading': '認識因數', 'body': {'bullets': ['12 的因數有 1、2、3、4、6、12']}, 'speaker_notes': '先問學生 12 可以怎麼分。'},
                           {'heading': '找最大公因數', 'body': {'bullets': ['先寫出兩數的因數']}, 'speaker_notes': 'x'}]}
        (self.out / 'rendered-deck.json').write_text(json.dumps(deck, ensure_ascii=False), encoding='utf-8')
        self.named = {kind: str(self.out / f'六年級_因數與倍數.{ext}') for kind, ext in (('pptx', 'pptx'), ('pdf', 'pdf'), ('json', 'json'))}
        self.named['docx'] = str(self.out / '六年級_因數與倍數_教師備課稿.docx')
        for path in self.named.values():
            Path(path).write_bytes(b'old')
        (self.out / 'result.json').write_text(json.dumps({'artifacts': self.named}), encoding='utf-8')
        self.exported = []
        patcher = patch.object(page_edit, 'log_event')
        patcher.start()
        self.addCleanup(patcher.stop)

    def export(self, request, out):
        self.exported.append(out)
        (out / '教材.pdf').write_bytes(b'new pdf')

    def test_read_lists_content_paragraphs_and_notes_but_not_the_footer(self):
        page = page_edit.read_page(self.out, 1)
        self.assertEqual([t['text'] for t in page['texts']], ['認識因數', '12 的因數有 1、2、3、4、6、12'])
        self.assertEqual(page['notes'], '先問學生 12 可以怎麼分。')
        with self.assertRaises(ValueError):
            page_edit.read_page(self.out, 3)

    def test_save_changes_only_the_words_and_refreshes_every_download(self):
        before = geometry(self.out / '教材.pptx', 1)
        keys = [t['key'] for t in page_edit.read_page(self.out, 1)['texts']]
        result = page_edit.save_page(self.jid, 1, {keys[0]: '因數是什麼？', keys[1]: '12 的因數有 1、2、3、4、6、12'},
                                     '請學生用 12 顆彈珠分組。', self.store, self.export)
        self.assertTrue(result['changed'])
        self.assertIn('1 段文字', result['message'])
        self.assertIn('講者備註', result['message'])
        prs = Presentation(self.out / '教材.pptx')
        self.assertEqual(prs.slides[0].shapes[0].text_frame.text, '因數是什麼？')
        self.assertEqual(prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.size, Pt(40))
        self.assertEqual(geometry(self.out / '教材.pptx', 1), before)
        self.assertEqual(prs.slides[0].notes_slide.notes_text_frame.text, '請學生用 12 顆彈珠分組。')
        self.assertEqual(self.exported, [self.out])
        self.assertEqual(Path(self.named['pdf']).read_bytes(), b'new pdf')
        self.assertEqual(Presentation(self.named['pptx']).slides[0].shapes[0].text_frame.text, '因數是什麼？')
        deck = json.loads(Path(self.named['json']).read_text(encoding='utf-8'))
        self.assertEqual(deck['slides'][0]['heading'], '因數是什麼？')
        self.assertEqual(deck['slides'][0]['speaker_notes'], '請學生用 12 顆彈珠分組。')
        self.assertGreater(Path(self.named['docx']).stat().st_size, 100)
        self.assertEqual(len(list(self.out.glob('page-edit-backup-*'))), 1)

    def test_unchanged_text_is_not_an_edit(self):
        texts = {t['key']: t['text'] + ' ' for t in page_edit.read_page(self.out, 2)['texts']}
        self.assertFalse(page_edit.save_page(self.jid, 2, texts, None, self.store, self.export)['changed'])
        self.assertEqual(self.exported, [])

    def test_empty_text_and_unfinished_decks_are_refused(self):
        key = page_edit.read_page(self.out, 1)['texts'][0]['key']
        with self.assertRaises(ValueError):
            page_edit.save_page(self.jid, 1, {key: '  '}, None, self.store, self.export)
        with self.store.db() as c:
            c.execute("UPDATE jobs SET status='running' WHERE id=?", (self.jid,))
        with self.assertRaises(ValueError):
            page_edit.save_page(self.jid, 1, {key: '新標題'}, None, self.store, self.export)

    def test_a_failed_export_puts_the_old_deck_back(self):
        key = page_edit.read_page(self.out, 1)['texts'][0]['key']

        def broken(request, out):
            raise RuntimeError('LibreOffice 沒有回應')
        with self.assertRaises(ValueError) as caught:
            page_edit.save_page(self.jid, 1, {key: '新標題'}, None, self.store, broken)
        self.assertIn('已還原', str(caught.exception))
        self.assertEqual(Presentation(self.out / '教材.pptx').slides[0].shapes[0].text_frame.text, '認識因數')
        self.assertEqual((self.out / '教材.pdf').read_bytes(), b'old pdf')
        self.assertEqual(Path(self.named['pptx']).read_bytes(), b'old')


if __name__ == '__main__':
    unittest.main()
