
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock

from pptx import Presentation

from tests.test_final_optimizer import FinalOptimizerTest
from workflow_v2 import final_optimizer as final
from workflow_v2 import polish_scope as scope
from workflow_v2.store import atomic


def clean_report(count=8):
    return {'schema_version': 1, 'pages': {str(i): {'history': [], 'findings': [], 'passed': True, 'final': 'passed'}
                                           for i in range(1, count + 1)}}


class LintTest(unittest.TestCase):
    def test_unfinished_or_repeated_text_is_noticed(self):
        for text in ('本頁完整題目與詳解已保留於，請教師重新排版後使用。', '誰或什麼東西？也就是誰或什麼東西？',
                     '請同學一起看這張圖請同學一起看這張圖', '這裡還有 **粗體** 標記', '先做完，，再檢查'):
            self.assertTrue(scope.lint(text), text)

    def test_normal_classroom_text_is_not_flagged(self):
        for text in ('終於，我們找到答案了。', '這是我的，那是你的。', '答案是：對！', '姓名：＿＿＿＿＿＿＿＿＿＿＿＿',
                     '24 和 36 的最大公因數是 12。', '(1) The ball is under the box.'):
            self.assertEqual(scope.lint(text), [], text)


class SelectTest(unittest.TestCase):
    def pages(self, count=8):
        return [{'page': i, 'text': ['正常的一句話。'], 'editable': []} for i in range(1, count + 1)]

    def folder(self, report):
        out = Path(tempfile.mkdtemp())
        if report is not None:
            atomic(out / 'qa-report.json', report)
        return out

    def test_without_a_trustworthy_report_every_page_is_checked(self):
        self.assertIsNone(scope.select(self.folder(None), self.pages()))
        self.assertIsNone(scope.select(self.folder(clean_report(5)), self.pages(8)))
        out = self.folder(None)
        (out / 'qa-report.json').write_text('{broken', encoding='utf-8')
        self.assertIsNone(scope.select(out, self.pages()))

    def test_clean_deck_has_nothing_to_polish(self):
        self.assertEqual(scope.select(self.folder(clean_report()), self.pages()), {})

    def test_each_kind_of_offline_evidence_flags_its_page(self):
        report = clean_report()
        report['pages']['2']['final'] = 'fallback'
        report['pages']['3']['findings'] = [{'code': 'long_sentence', 'severity': 'minor', 'message': '句子過長'}]
        report['pages']['4']['findings'] = [{'code': 'ghost_image', 'severity': 'major', 'message': '文字提到圖，但本頁沒有圖'}]
        report['pages']['5']['findings'] = [{'code': 'hook_names_concept', 'severity': 'minor', 'message': '導入題不該先講出概念名稱'}]
        report['pages']['6']['findings'] = [{'code': 'fixed', 'severity': 'minor', 'message': '已自動修正'}]
        out = self.folder(report)
        atomic(out / 'teacher-checks.json', {'7': {'reason': 'x', 'codes': []}})
        atomic(out / 'deck-validation.json', {'findings': [{'page': 8, 'code': 'too_long', 'severity': 'minor', 'message': '正文過長'}]})
        atomic(out / 'visual-qa.json', {'remaining_pages': [1]})
        pages = self.pages()
        pages[2]['text'] = ['完全正常']
        result = scope.select(out, pages)
        self.assertEqual(sorted(result), [1, 2, 3, 4, 7, 8])
        self.assertTrue(all(result[n] for n in result))

    def test_lint_flags_a_page_even_when_every_report_is_clean(self):
        pages = self.pages()
        pages[4]['text'] = ['本頁完整題目與詳解已保留於，請教師重新排版後使用。']
        self.assertEqual(sorted(scope.select(self.folder(clean_report()), pages)), [5])


class OptimizeScopeTest(unittest.TestCase):
    def setup_job(self, out, report):
        request = FinalOptimizerTest().fixture(out)
        if report is not None:
            atomic(out / 'qa-report.json', report)
        return request

    def test_clean_deck_skips_the_model_and_leaves_the_file_untouched(self):
        with tempfile.TemporaryDirectory() as folder:
            out = Path(folder)
            request = self.setup_job(out, clean_report())
            before = final.sha(out / '教材.pptx')
            model = Mock()
            result = final.optimize(request, out, model)
            model.assert_not_called()
            self.assertEqual(result['status'], 'skipped')
            self.assertEqual(result['model_calls'], 0)
            self.assertFalse(result['changed'])
            self.assertFalse(result['requires_teacher_review'])
            self.assertEqual(final.sha(out / '教材.pptx'), before)
            self.assertFalse((out / 'final-optimization' / 'before.pptx').exists())
            self.assertEqual(final.optimize(request, out, model), result)
            model.assert_not_called()

    def test_only_flagged_pages_reach_the_model_and_their_flags_are_in_the_prompt(self):
        with tempfile.TemporaryDirectory() as folder:
            out = Path(folder)
            report = clean_report()
            report['pages']['3']['findings'] = [{'code': 'long_sentence', 'severity': 'minor', 'message': '句子過長了'}]
            request = self.setup_job(out, report)
            mapping, pages = final.fields(Presentation(out / '教材.pptx'))
            key = next(k for k, v in mapping.items() if v[0].text == '這是一個關於位置的說明。')
            seen = {}

            def propose(model, prompt, schema, *args, **kwargs):
                seen['prompt'] = prompt
                seen['fields'] = schema['properties']['revisions']['items']['properties']['field']['enum']
                return {'revisions': [{'field': key, 'text': '這頁說明位置。', 'reason': '刪除贅詞'}], 'issues': []}
            model = Mock(side_effect=propose)
            result = final.optimize(request, out, model)
            self.assertEqual(model.call_count, 1)
            self.assertTrue(seen['fields'])
            self.assertTrue(all(f.startswith('3:') for f in seen['fields']))
            self.assertIn('句子過長了', seen['prompt'])
            self.assertIn('flags', seen['prompt'])
            self.assertEqual(result['status'], 'completed')
            self.assertEqual(result['flagged_pages'], [3])
            self.assertEqual(result['pages_checked'], 1)
            self.assertEqual(result['pages'], 8)
            self.assertEqual(result['model_reviewed_pages'], 1)
            self.assertTrue(result['changed'])
            self.assertIn('這頁說明位置。', '\n'.join(s.text for s in Presentation(out / '教材.pptx').slides[2].shapes if s.has_text_frame))

    def test_several_flagged_pages_still_share_calls_of_four(self):
        with tempfile.TemporaryDirectory() as folder:
            out = Path(folder)
            report = clean_report()
            for n in (1, 2, 3, 5, 6, 8):
                report['pages'][str(n)]['final'] = 'warning'
            request = self.setup_job(out, report)
            model = Mock(return_value={'revisions': [], 'issues': []})
            result = final.optimize(request, out, model)
            self.assertEqual(model.call_count, 2)
            self.assertEqual(result['flagged_pages'], [1, 2, 3, 5, 6, 8])
            progress = json.loads((out / 'final-optimization-progress.json').read_text(encoding='utf-8'))
            self.assertEqual(progress['total_pages'], 6)

    def test_a_missing_report_keeps_the_old_thorough_behaviour(self):
        with tempfile.TemporaryDirectory() as folder:
            out = Path(folder)
            request = self.setup_job(out, None)
            model = Mock(return_value={'revisions': [], 'issues': []})
            result = final.optimize(request, out, model)
            self.assertEqual(model.call_count, 2)
            self.assertEqual(result['model_reviewed_pages'], 8)
            self.assertEqual(result['flagged_pages'], [])

    def test_a_failed_model_call_still_reports_incomplete_review(self):
        with tempfile.TemporaryDirectory() as folder:
            out = Path(folder)
            report = clean_report()
            report['pages']['2']['final'] = 'warning'
            request = self.setup_job(out, report)
            before = final.sha(out / '教材.pptx')
            result = final.optimize(request, out, Mock(side_effect=RuntimeError('離線')))
            self.assertEqual(result['status'], 'unavailable')
            self.assertTrue(result['requires_teacher_review'])
            self.assertEqual(final.sha(out / '教材.pptx'), before)


if __name__ == '__main__':
    unittest.main()
